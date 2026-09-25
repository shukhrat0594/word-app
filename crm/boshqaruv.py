"""CRM boshqaruv API — video-TZ ("CRM TZ.mp4", 2026-09-23).

Shuhrat qarori: CRM endi ASOSIY manba. Talaba, guruh, xodim shu yerda
yaratiladi; saytda (LMS) faqat mashqlar va o'qituvchi ishi qoladi.

Bog'lanish yo'nalishi o'zgarmadi — `crm -> LMS`: bu modul LMS
jadvallariga (`User`, `Guruh`, `GuruhAzoligi`, `Davomat`) YOZADI, lekin
LMS kodi CRM haqida hech narsa bilmaydi. Talaba maydonlarini yozish uchun
LMS'ning o'z funksiyasi (`accounts.views._talaba_maydonlarini_yoz`)
qayta ishlatiladi — qulf va tekshiruv qoidalari ikki joyda turmasin.

Konvensiya `crm/views.py` bilan bir xil: serializer emas, dict
quruvchilar; har bir o'zgarish `audit.utils.logla()` orqali yoziladi.
"""

import secrets
from datetime import timedelta
from decimal import Decimal

from django.db import transaction
from django.db.models import Count, Prefetch, Q, Sum
from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.response import Response

from academics.models import Davomat, Guruh, GuruhAzoligi
from accounts.excel_import import LOGIN_QOIDASI
from accounts.models import Markaz, User
from accounts.permissions import owner_mi
from audit.models import FaoliyatYozuvi
from audit.utils import logla
from courses.models import KursTugun

from . import mantiq
from .models import (
    AzolikMoliya,
    Chegirma,
    CrmRol,
    DarsBahosi,
    DarsJadvali,
    DarsMavzusi,
    DarsOzgarish,
    DavomatIzoh,
    Eslatma,
    Filial,
    GuruhMoliya,
    GuruhOqituvchi,
    Hisob,
    GuruhdanChiqish,
    Lid,
    LidBolim,
    LidDoska,
    LidTarix,
    TalabaProfil,
    Tolov,
    Xona,
    XodimProfil,
)
from .filial import (
    cheklanganmi, filial_korinadimi, filial_q, filial_tekshir, guruh_q, guruh_tekshir, ruxsat_filiallari,
    talaba_korinadimi, talaba_tekshir, tolov_filiali_q, tolov_q, umumiy_yozuv_taqiq,
)
from .permissions import CrmView
from .ruxsatlar import BARCHA_KALITLAR, RUXSAT_DARAXTI, lms_roli, ruxsatlar, sayt_menyusini_toraytir
from .views import _guruh_dict, _oy, _ruxsatsiz, _sana, _son, _vaqt, _xato, jadvalni_tekshir

NOL = Decimal("0")


def _ism(u):
    return (u.get_full_name() or u.username) if u else None


def _telefon(qiymat):
    """Telefonni bitta ko'rinishga keltiradi: faqat raqamlar va boshida '+'.
    Qidiruv va takror tekshiruvi shu ko'rinishga tayanadi."""
    matn = str(qiymat or "").strip()
    raqamlar = "".join(ch for ch in matn if ch.isdigit())
    if not raqamlar:
        return ""
    return ("+" if matn.startswith("+") or len(raqamlar) == 12 else "") + raqamlar


def _markaz_id():
    """Platforma bitta markaz bilan ishlaydi — LMS'dagi fallback bilan bir xil."""
    return Markaz.objects.order_by("id").values_list("id", flat=True).first()


# ── Login va parol ───────────────────────────────────────────────────


_PAROL_ALIFBOSI = "abcdefghjkmnpqrstuvwxyz23456789"


def parol_yarat():
    """8 belgili tasodifiy parol. Raqam+harf aralash — Django'ning parol
    validatorlaridan o'tadi (faqat raqam bo'lsa `NumericPasswordValidator`
    rad etardi). O'xshash belgilar (0/o, 1/l/i) yo'q — ota-onaga telefonda
    aytib berilganda adashmasin."""
    while True:
        parol = "".join(secrets.choice(_PAROL_ALIFBOSI) for _ in range(8))
        if any(c.isdigit() for c in parol) and any(c.isalpha() for c in parol):
            return parol


def login_yarat(ism, telefon):
    """Talaba/xodim uchun bo'sh login. Telefon bo'lsa — uning oxirgi 9
    raqami (talaba loginini eslab qolishi oson), bo'lmasa ismdan."""
    raqamlar = "".join(ch for ch in str(telefon or "") if ch.isdigit())
    if len(raqamlar) >= 9:
        asos = "u" + raqamlar[-9:]
    else:
        lotin = "".join(ch for ch in str(ism or "").lower() if ch.isalnum() and ch.isascii())
        asos = lotin[:20] or "talaba"
    login = asos
    n = 1
    while User.objects.filter(username=login).exists():
        n += 1
        login = f"{asos}{n}"
    return login


# ── Ruxsatlar va joriy foydalanuvchi ─────────────────────────────────


class MenView(CrmView):
    """Joriy foydalanuvchining CRM ruxsatlari — frontend menyu va
    tugmalarni shunga qarab ko'rsatadi (himoya baribir backendda)."""

    def get(self, request):
        return Response({
            "ruxsatlar": sorted(ruxsatlar(request.user)),
            "is_owner": owner_mi(request.user),
            # Filial cheklovi: `null` — cheklovsiz, aks holda ko'rinadigan filiallar.
            "filiallar": (lambda s: None if s is None else sorted(s))(ruxsat_filiallari(request.user)),
            "daraxt": [
                {"kalit": k, "nomi": n, "bolalar": [{"kalit": kk, "nomi": nn} for kk, nn in b]}
                for k, n, b in RUXSAT_DARAXTI
            ],
        })


def _rol_dict(r):
    return {
        "id": r.id, "nomi": r.nomi, "faol": r.faol, "ruxsatlar": r.ruxsatlar,
        # Tizim roli (lavozim) — o'chirilmaydi, nomi o'zgarmaydi, xodimga
        # "maxsus rol" sifatida berilmaydi (lavozimning o'zi yetadi).
        "lavozim": r.lavozim,
        "xodimlar_soni": getattr(r, "_soni", None),
    }


def _ruxsat_royxati(xom):
    if not isinstance(xom, list):
        raise ValueError("ruxsatlar ro'yxat bo'lishi kerak")
    return sorted({str(k) for k in xom if str(k) in BARCHA_KALITLAR})


def _lavozim_beradi(user):
    """Xodim yaratish va lavozim berish — faqat owner yoki administrator
    (2026-09-23, Shuhrat). Aks holda `xodimlar.qoshish` ruxsati bor kassir
    o'qituvchi yoki kengroq ruxsatli lavozimdagi hisob ochib, parolini
    o'zi qo'yib, o'sha huquqlar bilan kira olardi."""
    return owner_mi(user) or user.role == User.Role.ADMIN


def _rollarni_boshqaradi(user):
    """Rol yaratish/tahrirlash — faqat `sozlamalar.rollar` (yoki owner).
    Bo'lim darajasidagi `sozlamalar` YETMAYDI: aks holda "Kurs narxlari"
    ruxsati bor xodim o'z rolini tahrirlab hamma ruxsatni olardi."""
    return owner_mi(user) or "sozlamalar.rollar" in ruxsatlar(user)


class RollarView(CrmView):
    bolim = "xodimlar"
    oqish_ochiq = True

    def get(self, request):
        from .ruxsatlar import TIZIM_ROLLARI

        qs = list(CrmRol.objects.annotate(_soni=Count("xodimlar")))
        # Tizim rolida "xodimlar soni" — maxsus rolsiz shu lavozimdagilar.
        lavozim_soni = dict(
            XodimProfil.objects.filter(Q(rol__isnull=True) | Q(rol__lavozim__isnull=False), user__is_active=True)
            .values_list("lavozim").annotate(n=Count("id"))
        )
        # Saytda yaratilgan, CRM profili yo'q o'qituvchi va administratorlar
        # ham o'z lavozimi bo'yicha sanaladi (`ruxsatlar._lavozim` bilan bir xil).
        profilsiz = _xodimlar_qs().filter(is_active=True, crm_xodim__isnull=True)
        lavozim_soni["oqituvchi"] = lavozim_soni.get("oqituvchi", 0) + profilsiz.filter(role=User.Role.TEACHER).count()
        lavozim_soni["admin"] = lavozim_soni.get("admin", 0) + profilsiz.filter(role=User.Role.ADMIN).count()
        tartib = {k: i for i, (k, _) in enumerate(TIZIM_ROLLARI)}
        for r in qs:
            if r.lavozim:
                r._soni = lavozim_soni.get(r.lavozim, 0)
        qs.sort(key=lambda r: (0, tartib.get(r.lavozim, 99), "") if r.lavozim else (1, 0, r.nomi.lower()))
        return Response([_rol_dict(r) for r in qs])

    def post(self, request):
        if not _rollarni_boshqaradi(request.user):
            return _xato("Rol yaratishga ruxsat yo'q", kod=403)
        nomi = (request.data.get("nomi") or "").strip()
        if not nomi:
            return _xato("Rol nomi bo'sh bo'lmasin")
        if CrmRol.objects.filter(nomi__iexact=nomi).exists():
            return _xato("Bunday rol bor")
        try:
            kalitlar = _ruxsat_royxati(request.data.get("ruxsatlar") or [])
        except ValueError as e:
            return _xato(str(e))
        rol = CrmRol.objects.create(nomi=nomi[:100], faol=bool(request.data.get("faol", True)), ruxsatlar=kalitlar)
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.YARATISH, obyekt=rol,
              obyekt_turi="CRM Rol", snapshot=_rol_dict(rol))
        return Response(_rol_dict(rol), status=201)


class RolDetailView(CrmView):
    bolim = "sozlamalar"

    def patch(self, request, pk):
        if not _rollarni_boshqaradi(request.user):
            return _xato("Rollarni tahrirlashga ruxsat yo'q", kod=403)
        rol = get_object_or_404(CrmRol, pk=pk)
        eski = _rol_dict(rol)
        if rol.lavozim and ("nomi" in request.data and request.data.get("nomi") != rol.nomi
                            or "faol" in request.data and not request.data["faol"]):
            return _xato("Lavozim rolining nomi o'zgarmaydi va u o'chirilmaydi — faqat ruxsatlari")
        if "nomi" in request.data:
            nomi = (request.data.get("nomi") or "").strip()
            if not nomi:
                return _xato("Rol nomi bo'sh bo'lmasin")
            if CrmRol.objects.filter(nomi__iexact=nomi).exclude(pk=pk).exists():
                return _xato("Bunday rol bor")
            rol.nomi = nomi[:100]
        if "faol" in request.data:
            rol.faol = bool(request.data["faol"])
        if "ruxsatlar" in request.data:
            try:
                rol.ruxsatlar = _ruxsat_royxati(request.data["ruxsatlar"])
            except ValueError as e:
                return _xato(str(e))
        rol.save()
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.OZGARTIRISH, obyekt=rol,
              obyekt_turi="CRM Rol", obyekt_nomi=rol.nomi,
              eski_qiymatlar={k: str(v) for k, v in eski.items()},
              yangi_qiymatlar={k: str(v) for k, v in _rol_dict(rol).items()})
        return Response(_rol_dict(rol))

    def delete(self, request, pk):
        if not _rollarni_boshqaradi(request.user):
            return _xato("Rollarni o'chirishga ruxsat yo'q", kod=403)
        rol = get_object_or_404(CrmRol, pk=pk)
        if rol.lavozim:
            return _xato("Lavozim roli o'chirilmaydi")
        if rol.xodimlar.exists():
            return _xato("Bu rol xodimlarga berilgan — avval ularning rolini almashtiring")
        nomi = rol.nomi
        rol.delete()
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.OCHIRISH, obyekt=request.user,
              obyekt_turi="CRM Rol", obyekt_nomi=nomi)
        return Response(status=204)


# ── Xodimlar ─────────────────────────────────────────────────────────


def _xodim_dict(u, oylik_korsin=True):
    p = getattr(u, "crm_xodim", None)
    filiallar = list(p.filiallar.all()) if p else []
    return {
        "id": u.id,
        "ism": _ism(u),
        "username": u.username,
        "telefon": u.telefon,
        "tugilgan_sana": u.tugilgan_sana,
        "faol": u.is_active,
        "lms_roli": u.role,
        # Saytdagi owner CRM'da CEO (2026-09-23, Shuhrat) — CEO lavozimi
        # boshqa xodimga berilmaydi.
        "lavozim": "ceo" if u.is_superuser else (
            p.lavozim if p else ("oqituvchi" if u.role == User.Role.TEACHER else "admin")
        ),
        "owner": u.is_superuser,
        "rol_id": p.rol_id if p else None,
        "rol": p.rol.nomi if p and p.rol_id else None,
        # Bir nechta filial (2026-09-23). Bo'sh — cheklovsiz (hamma filial).
        "filial_idlar": [f.id for f in filiallar],
        "filial": ", ".join(f.nomi for f in filiallar) or None,
        "jins": p.jins if p else "",
        "ishga_olingan_sana": p.ishga_olingan_sana if p else None,
        # Oylik va ulush — maxfiy: faqat `xodimlar.oylik` ruxsati bo'lsa.
        "oylik": (p.oylik if p else NOL) if oylik_korsin else None,
        "foiz_ulushi": (p.foiz_ulushi if p else NOL) if oylik_korsin else None,
    }


def _xodimlar_qs():
    """CRM'dagi xodimlar: o'qituvchi va administratorlar (LMS'da
    yaratilgan eskilari ham) + CRM profili bor har qanday foydalanuvchi
    (kassir, marketolog...). Owner ro'yxatda KO'RINMAYDI — u xodim emas,
    platforma egasi; uni tahrirlash CRM ishi emas."""
    return (
        User.objects.filter(
            Q(role__in=[User.Role.TEACHER, User.Role.ADMIN]) | Q(crm_xodim__isnull=False)
        )
        .filter(is_superuser=False)
        .exclude(role__in=[User.Role.STUDENT, User.Role.PARENT])
        .select_related("crm_xodim", "crm_xodim__rol")
        .prefetch_related("crm_xodim__filiallar")
        .distinct()
    )


class _Taqiq(Exception):
    """Ruxsat yetmaydi — 403 bo'lib qaytadi (ValueError — 400)."""


def _xodim_maydonlari(user, profil, data, kim):
    """Xodim formasi maydonlarini yozadi (yaratish va tahrir uchun
    umumiy). Xato bo'lsa ValueError, ruxsat yetmasa `_Taqiq`.

    `kim` — so'rov egasi:
      * maxsus rolni faqat rollarni boshqaradigan (owner yoki
        `sozlamalar.rollar`) beradi — aks holda `xodimlar` ruxsati bor
        odam o'ziga yoki sherigiga to'liq huquqli rolni biriktirardi;
      * oylik va foiz ulushi `xodimlar.oylik` ruxsatisiz E'TIBORSIZ
        qoldiriladi: bunday foydalanuvchida ular yashirin (`null`) keladi
        va forma ularni bo'sh qaytarib, oylikni 0 ga tushirib yuborardi.
    """
    kim_ruxsatlari = ruxsatlar(kim)
    if "ism" in data:
        ism = (data.get("ism") or "").strip()
        if not ism:
            raise ValueError("Ism familiya bo'sh bo'lmasin")
        user.first_name = ism[:150]
        user.last_name = ""
    if "telefon" in data:
        user.telefon = _telefon(data.get("telefon"))[:20]
    if "tugilgan_sana" in data:
        user.tugilgan_sana = _sana(data.get("tugilgan_sana"), "tugilgan_sana", majburiy=False)
    if "lavozim" in data:
        lavozim = data.get("lavozim")
        if lavozim not in dict(XodimProfil.Lavozim.choices):
            raise ValueError("Noma'lum lavozim")
        if lavozim != profil.lavozim and not _lavozim_beradi(kim):
            raise _Taqiq("Lavozimni faqat owner yoki administrator o'zgartiradi")
        if lavozim == XodimProfil.Lavozim.CEO:
            raise ValueError("CEO — bu owner, xodimga berilmaydi. Administrator lavozimini tanlang")
        profil.lavozim = lavozim
    if "rol_id" in data:
        rol_id = data.get("rol_id") or None
        if str(rol_id or "") != str(profil.rol_id or ""):
            if not _rollarni_boshqaradi(kim):
                raise _Taqiq("Xodimga rol berishga ruxsat yo'q")
            rol = get_object_or_404(CrmRol, pk=rol_id) if rol_id else None
            if rol is not None and rol.lavozim:
                raise ValueError("Lavozim roli maxsus rol sifatida berilmaydi — lavozimni tanlang")
            profil.rol = rol
    filiallar = None
    if "filial_idlar" in data or "filial_id" in data:
        filiallar = _xodim_filiallari(profil, data, kim)
    if "jins" in data:
        jins = data.get("jins") or ""
        if jins and jins not in dict(XodimProfil.Jins.choices):
            raise ValueError("Jins noto'g'ri")
        profil.jins = jins
    if "ishga_olingan_sana" in data:
        profil.ishga_olingan_sana = _sana(data.get("ishga_olingan_sana"), "ishga_olingan_sana", majburiy=False)
    oylik_ruxsati = "xodimlar.oylik" in kim_ruxsatlari
    if "oylik" in data and oylik_ruxsati:
        oylik = _son(data.get("oylik") or 0, "oylik")
        if oylik < 0:
            raise ValueError("Oylik manfiy bo'lmasin")
        profil.oylik = oylik
    if "foiz_ulushi" in data and oylik_ruxsati:
        foiz = _son(data.get("foiz_ulushi") or 0, "foiz_ulushi")
        if not 0 <= foiz <= 100:
            raise ValueError("Foiz ulushi 0..100 oralig'ida bo'lsin")
        profil.foiz_ulushi = foiz
    # M2M — profil saqlangandan keyin yoziladi (chaqiruvchi `.set()` qiladi).
    return filiallar


def _xodim_filiallari(profil, data, kim):
    """Xodimning yangi filiallari ro'yxati. Filialga bog'langan (cheklangan)
    xodim boshqa xodimga faqat O'Z filiallarini beradi/oladi; tahrirdagi
    xodimning boshqa filiallari saqlanadi; va hech kimni filialsiz qoldira
    olmaydi — filialsiz xodim hamma filialni ko'radi (qaror), ya'ni bu
    cheklovdan chiqarib yuborish bo'lardi."""
    if "filial_idlar" in data:
        xom = data.get("filial_idlar") or []
        if not isinstance(xom, list):
            raise ValueError("filial_idlar ro'yxat bo'lishi kerak")
    else:
        xom = [data.get("filial_id")] if data.get("filial_id") else []
    try:
        idlar = {int(x) for x in xom if x not in (None, "")}
    except (TypeError, ValueError):
        raise ValueError("Filial noto'g'ri") from None
    if Filial.objects.filter(pk__in=idlar).count() != len(idlar):
        raise ValueError("Filial topilmadi")
    ruxsat = ruxsat_filiallari(kim)
    if ruxsat is None:
        return sorted(idlar)
    eski = set(profil.filiallar.values_list("id", flat=True)) if profil.pk else set()
    # Xodimning ALLAQACHON bor boshqa filiali qaytib kelsa — xato emas:
    # forma mavjud ro'yxatni to'liq qaytaradi (ikki filialli xodimni ismini
    # tuzatish uchun ochganda ham). Yangi begona filial qo'shish — taqiq.
    if idlar - ruxsat - eski:
        raise _Taqiq("Faqat o'z filialingizni biriktira olasiz")
    yakuniy = (idlar & ruxsat) | (eski - ruxsat)
    if not yakuniy:
        # Yangi xodim — qo'shuvchining bitta filiali bo'lsa, o'shanisi o'zi qo'yiladi.
        if not profil.pk and len(ruxsat) == 1:
            return sorted(ruxsat)
        raise ValueError("Filialni tanlang")
    return sorted(yakuniy)


class XodimlarView(CrmView):
    """SoffCRM "Sozlamalar -> Xodimlar": ro'yxat (lavozim tablari bilan)
    va "Xodim qo'shish"."""

    bolim = "xodimlar"
    oqish_ochiq = True

    def get(self, request):
        p = request.query_params
        arxiv = bool(p.get("arxiv"))
        qs = _xodimlar_qs().filter(is_active=not arxiv)
        # `?tanlov=1` — guruh/lid formasidagi o'qituvchi tanlovi: HAMMA
        # o'qituvchi (bir nechta filialda dars beradi — qaror). "Xodimlar"
        # sahifasida esa filial xodimi faqat o'z filiali xodimlarini ko'radi.
        ruxsat = None if p.get("tanlov") else ruxsat_filiallari(request.user)
        if ruxsat is not None:
            qs = qs.filter(
                Q(crm_xodim__isnull=True) | Q(crm_xodim__filiallar__isnull=True)
                | Q(crm_xodim__filiallar__in=ruxsat)
            )
        qidiruv = (p.get("q") or "").strip()
        if qidiruv:
            qs = qs.filter(
                Q(first_name__icontains=qidiruv) | Q(username__icontains=qidiruv) | Q(telefon__icontains=qidiruv)
            )
        oylik_korsin = "xodimlar.oylik" in ruxsatlar(request.user)
        users = list(qs.order_by("first_name", "username"))
        if not arxiv and not p.get("tanlov"):
            # Owner — "CEO" qatori (CRM'dan tahrirlanmaydi).
            owners = User.objects.filter(is_superuser=True, is_active=True).order_by("id")
            if qidiruv:
                owners = owners.filter(Q(first_name__icontains=qidiruv) | Q(username__icontains=qidiruv))
            users = list(owners) + users
        royxat = [_xodim_dict(u, oylik_korsin) for u in users]
        lavozim = request.query_params.get("lavozim")
        # Tab hisoblagichlari filtrdan OLDIN — "O'QITUVCHI - 12" har doim ko'rinsin.
        soni = {}
        for x in royxat:
            soni[x["lavozim"]] = soni.get(x["lavozim"], 0) + 1
        if lavozim:
            royxat = [x for x in royxat if x["lavozim"] == lavozim]
        return Response({"xodimlar": royxat, "soni": soni, "jami": sum(soni.values())})

    def post(self, request):
        if "xodimlar.qoshish" not in ruxsatlar(request.user) or not _lavozim_beradi(request.user):
            return _xato("Xodimni faqat owner yoki administrator qo'shadi", kod=403)
        data = request.data
        lavozim = data.get("lavozim") or XodimProfil.Lavozim.OQITUVCHI
        if lavozim not in dict(XodimProfil.Lavozim.choices):
            return _xato("Noma'lum lavozim")
        if lavozim == XodimProfil.Lavozim.CEO:
            return _xato("CEO — bu owner, xodimga berilmaydi. Administrator lavozimini tanlang")
        if lavozim == "admin" and not owner_mi(request.user):
            # Administrator LMS'da ham admin bo'ladi — buni faqat owner beradi.
            return _xato("Administratorni faqat owner qo'sha oladi", kod=403)
        ism = (data.get("ism") or "").strip()
        telefon = _telefon(data.get("telefon"))
        if not ism:
            return _xato("Ism familiya kiritilsin")
        if not telefon:
            return _xato("Telefon raqam kiritilsin")

        login = (data.get("login") or "").strip() or login_yarat(ism, telefon)
        if not LOGIN_QOIDASI.match(login):
            return _xato("Login faqat harf/raqam/./@/+/-/_ dan iborat bo'lsin")
        if User.objects.filter(username=login).exists():
            return _xato(f"«{login}» login band")
        parol = (data.get("parol") or "").strip()
        parol_yaratildi = not parol
        parol = parol or parol_yarat()
        from accounts.views import _parolni_tekshir

        xatolar = _parolni_tekshir(parol)
        if xatolar:
            return _xato(" ".join(xatolar))

        with transaction.atomic():
            user = User(username=login, role=lms_roli(lavozim), markaz_id=_markaz_id())
            sayt_menyusini_toraytir(user)
            user.set_password(parol)
            profil = XodimProfil(user=user, lavozim=lavozim)
            malumot = {**data, "lavozim": lavozim, "ism": ism, "telefon": telefon}
            # Filialga bog'langan xodim qo'shsa — filial majburiy (bittasi bo'lsa o'zi).
            if cheklanganmi(request.user) and "filial_idlar" not in malumot and "filial_id" not in malumot:
                malumot["filial_idlar"] = []
            try:
                filiallar = _xodim_maydonlari(user, profil, malumot, request.user)
            except _Taqiq as e:
                return _xato(str(e), kod=403)
            except ValueError as e:
                return _xato(str(e))
            user.save()
            profil.user = user
            profil.save()
            if filiallar is not None:
                profil.filiallar.set(filiallar)

        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.YARATISH, obyekt=user,
              obyekt_turi="CRM Xodim", snapshot={"username": login, "lavozim": lavozim})
        javob = _xodim_dict(user)
        # Avtomatik parol FAQAT shu javobda bir marta ko'rsatiladi.
        javob["parol"] = parol if parol_yaratildi else None
        return Response(javob, status=201)


_OZI_TAHRIRLAYDI = {"ism", "telefon", "tugilgan_sana", "jins", "parol", "parol_tiklash"}


def _oz_profilidagi_taqiq(user, profil, data):
    """O'z profilida o'zgartirib bo'lmaydigan maydon o'zgargan bo'lsa — xato matni."""
    def farq(maydon, joriy):
        return maydon in data and str(data.get(maydon) or "") != str(joriy or "")

    if farq("lavozim", profil.lavozim) or farq("rol_id", profil.rol_id) or ("faol" in data and not data["faol"]):
        return "O'z lavozimingiz, rolingiz va holatingizni o'zgartira olmaysiz"
    if "filial_idlar" in data or "filial_id" in data:
        xom = data.get("filial_idlar") if "filial_idlar" in data else [data.get("filial_id")]
        try:
            yangi = {int(x) for x in (xom or []) if x not in (None, "")}
        except (TypeError, ValueError):
            return "Filial noto'g'ri"
        if yangi != set(profil.filiallar.values_list("id", flat=True)):
            return "O'z filiallaringizni o'zgartira olmaysiz"
    for maydon in ("oylik", "foiz_ulushi"):
        if maydon in data and data.get(maydon) is not None:
            try:
                if Decimal(str(data.get(maydon) or 0)) != getattr(profil, maydon):
                    return "O'z oylik va ulushingizni o'zgartira olmaysiz"
            except ArithmeticError:
                return f"{maydon}: son bo'lishi kerak"
    if farq("ishga_olingan_sana", profil.ishga_olingan_sana):
        return "Ishga olingan sanangizni o'zgartira olmaysiz"
    return None


class XodimDetailView(CrmView):
    pk_turi = "xodim"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "xodimlar"

    def patch(self, request, pk):
        if "xodimlar.tahrirlash" not in ruxsatlar(request.user):
            return _xato("Xodimni tahrirlashga ruxsat yo'q", kod=403)
        user = get_object_or_404(_xodimlar_qs(), pk=pk)
        profil, _ = XodimProfil.objects.get_or_create(
            user=user,
            defaults={"lavozim": "oqituvchi" if user.role == User.Role.TEACHER else "admin"},
        )
        eski = _xodim_dict(user)
        data = request.data
        owner = owner_mi(request.user)
        ozi = user.pk == request.user.pk
        yangi_lavozim = data.get("lavozim", profil.lavozim)
        if (yangi_lavozim in ("admin", "ceo") or user.role == User.Role.ADMIN) and not owner and not ozi:
            if yangi_lavozim != profil.lavozim or user.role == User.Role.ADMIN:
                return _xato("Administratorni faqat owner tahrirlay oladi", kod=403)
        # O'zini tahrirlash (owner'dan boshqa): ism, telefon, tug'ilgan sana,
        # jins va o'z paroli — mumkin; lavozim, rol, holat, filial, oylik va
        # ishga olingan sana — yo'q. Forma hamma maydonni qaytaradi, shuning
        # uchun QIYMAT o'zgargani tekshiriladi, maydon borligi emas.
        if ozi and not owner:
            if xato := _oz_profilidagi_taqiq(user, profil, data):
                return _xato(xato, kod=403)
            data = {k: v for k, v in data.items() if k in _OZI_TAHRIRLAYDI}
        parol_sorandi = bool(data.get("parol_tiklash") or data.get("parol"))
        # Boshqa xodimning parolini faqat owner yoki administrator tiklaydi —
        # aks holda `xodimlar` ruxsati bor kassir o'qituvchi nomidan saytga kira olardi.
        if parol_sorandi and user.pk != request.user.pk and not (owner or request.user.role == User.Role.ADMIN):
            return _xato("Boshqa xodimning parolini faqat administrator tiklaydi", kod=403)
        try:
            filiallar = _xodim_maydonlari(user, profil, data, request.user)
        except _Taqiq as e:
            return _xato(str(e), kod=403)
        except ValueError as e:
            return _xato(str(e))
        if "lavozim" in data:
            user.role = lms_roli(profil.lavozim)
            sayt_menyusini_toraytir(user)
        if "faol" in data:
            user.is_active = bool(data["faol"])
        yangi_parol = None
        if parol_sorandi:
            from accounts.views import _parolni_tekshir

            yangi_parol = (data.get("parol") or "").strip() or parol_yarat()
            xatolar = _parolni_tekshir(yangi_parol, user=user)
            if xatolar:
                return _xato(" ".join(xatolar))
            user.set_password(yangi_parol)
        user.save()
        profil.save()
        if filiallar is not None:
            profil.filiallar.set(filiallar)
        # `user.crm_xodim` so'rov boshida (select/prefetch bilan) yuklangan —
        # `profil` esa alohida obyekt. Qayta o'qilmasa javob va audit ESKI
        # qiymatlarni (sana, oylik, filiallar) ko'rsatardi.
        user = _xodimlar_qs().get(pk=user.pk)
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.OZGARTIRISH, obyekt=user,
              obyekt_turi="CRM Xodim", obyekt_nomi=_ism(user),
              eski_qiymatlar={k: str(v) for k, v in eski.items()},
              yangi_qiymatlar={k: str(v) for k, v in _xodim_dict(user).items()})
        javob = _xodim_dict(user)
        javob["parol"] = yangi_parol
        return Response(javob)


# ── Lidlar ───────────────────────────────────────────────────────────


def _bolim_dict(b, soni=None):
    return {
        "id": b.id, "nomi": b.nomi, "tartib": b.tartib, "filial_id": b.filial_id, "soni": soni,
        "doska_id": b.doska_id, "guruh_id": b.guruh_id,
    }


def _lid_dict(lid, oxirgi_eslatma=None):
    return {
        "id": lid.id,
        "ism": lid.ism,
        "telefon": lid.telefon,
        "qoshimcha_telefon": lid.qoshimcha_telefon,
        "qoshimcha_ism": lid.qoshimcha_ism,
        "tugilgan_sana": lid.tugilgan_sana,
        "manba": lid.manba,
        "bolim_id": lid.bolim_id,
        "holat": lid.holat,
        "holat_nomi": lid.get_holat_display(),
        "harorat": lid.harorat,
        "filial_id": lid.filial_id,
        "kurs_id": lid.kurs_id,
        "kurs": lid.kurs.nomi if lid.kurs_id else None,
        "oqituvchi_id": lid.oqituvchi_id,
        "oqituvchi": _ism(lid.oqituvchi) if lid.oqituvchi_id else None,
        "qulay_vaqt": lid.qulay_vaqt,
        "kunlar": lid.kunlar,
        "izoh": lid.izoh,
        "arxiv": lid.arxiv,
        "arxiv_sabab": lid.arxiv_sabab,
        "arxiv_sabab_nomi": lid.get_arxiv_sabab_display() if lid.arxiv_sabab else "",
        "arxiv_izoh": lid.arxiv_izoh,
        "qora_royxat": lid.qora_royxat,
        "qora_royxat_izoh": lid.qora_royxat_izoh,
        "talaba_id": lid.talaba_id,
        "yigilayotgan_guruh_id": lid.yigilayotgan_guruh_id,
        "yigilayotgan_guruh": lid.yigilayotgan_guruh.name if lid.yigilayotgan_guruh_id else None,
        "kim_qoshdi": _ism(lid.kim_qoshdi) if lid.kim_qoshdi_id else None,
        "vaqt": lid.created_at,
        # Kartochka ustidagi izoh (SoffCRM'da sichqoncha olib borilganda
        # chiqadigan "23.09 kuni keladi").
        "oxirgi_eslatma": oxirgi_eslatma,
    }


_LID_MATN_MAYDONLARI = {
    "ism": 200, "qoshimcha_ism": 100, "manba": 100, "qulay_vaqt": 50, "kunlar": 20, "izoh": 500,
}


def _lid_maydonlari(lid, data, user=None):
    """Yaratish va tahrir uchun umumiy. O'zgargan maydon nomlarini qaytaradi.
    `user` — filial cheklovi: ustun va yig'ilayotgan guruh o'z filialidan."""
    ozgardi = []
    for maydon, uzunlik in _LID_MATN_MAYDONLARI.items():
        if maydon in data:
            qiymat = (data.get(maydon) or "").strip()[:uzunlik]
            if maydon == "ism" and not qiymat:
                raise ValueError("Ism bo'sh bo'lmasin")
            if getattr(lid, maydon) != qiymat:
                ozgardi.append(maydon)
            setattr(lid, maydon, qiymat)
    for maydon in ("telefon", "qoshimcha_telefon"):
        if maydon in data:
            qiymat = _telefon(data.get(maydon))[:20]
            if maydon == "telefon" and not qiymat:
                raise ValueError("Telefon raqam kiritilsin")
            if getattr(lid, maydon) != qiymat:
                ozgardi.append(maydon)
            setattr(lid, maydon, qiymat)
    if "tugilgan_sana" in data:
        lid.tugilgan_sana = _sana(data.get("tugilgan_sana"), "tugilgan_sana", majburiy=False)
    if "harorat" in data:
        harorat = data.get("harorat") or ""
        if harorat not in ("", "issiq", "iliq", "sovuq"):
            raise ValueError("Harorat noto'g'ri")
        lid.harorat = harorat
    if "holat" in data:
        if data["holat"] not in dict(Lid.Holat.choices):
            raise ValueError("Noma'lum holat")
        if lid.holat != data["holat"]:
            ozgardi.append("holat")
        lid.holat = data["holat"]
    for maydon, model in (("bolim_id", LidBolim), ("filial_id", Filial), ("kurs_id", KursTugun)):
        if maydon in data:
            qiymat = data.get(maydon) or None
            yozuv = model.objects.filter(pk=qiymat).first() if qiymat else None
            if qiymat and yozuv is None:
                raise ValueError(f"{maydon}: topilmadi")
            if (user is not None and model is LidBolim and yozuv is not None
                    and not filial_korinadimi(user, yozuv.filial_id)):
                raise ValueError(f"{maydon}: topilmadi")
            if getattr(lid, maydon) != (int(qiymat) if qiymat else None):
                ozgardi.append(maydon)
            setattr(lid, maydon, int(qiymat) if qiymat else None)
    if "oqituvchi_id" in data:
        qiymat = data.get("oqituvchi_id") or None
        if qiymat and not User.objects.filter(pk=qiymat, role=User.Role.TEACHER).exists():
            raise ValueError("O'qituvchi topilmadi")
        lid.oqituvchi_id = int(qiymat) if qiymat else None
    if "yigilayotgan_guruh_id" in data:
        qiymat = data.get("yigilayotgan_guruh_id") or None
        guruhlar = Guruh.objects.filter(faol=True)
        if user is not None:
            guruhlar = guruhlar.filter(guruh_q(user))
        if qiymat and not guruhlar.filter(pk=qiymat).exists():
            raise ValueError("Guruh topilmadi")
        if (int(qiymat) if qiymat else None) != lid.yigilayotgan_guruh_id:
            ozgardi.append("yigilayotgan_guruh_id")
        lid.yigilayotgan_guruh_id = int(qiymat) if qiymat else None
    # Arxiv va qora ro'yxat — SABAB bilan (video-TZ 2026-09-25): arxivda
    # ro'yxatdan sabab (+ "Boshqa sabab"da izoh majburiy), qora ro'yxatda
    # izoh majburiy. Chiqarilganda sabab tozalanadi.
    if "arxiv" in data:
        qiymat = bool(data["arxiv"])
        if qiymat and not lid.arxiv:
            sabab = data.get("arxiv_sabab") or ""
            izoh = str(data.get("arxiv_izoh") or "").strip()[:300]
            if sabab not in dict(Lid._meta.get_field("arxiv_sabab").choices):
                raise ValueError("Arxivlash sababini tanlang")
            if sabab == "boshqa" and not izoh:
                raise ValueError("Sababni yozing")
            lid.arxiv_sabab, lid.arxiv_izoh = sabab, izoh
            ozgardi.append("arxiv")
        elif not qiymat and lid.arxiv:
            lid.arxiv_sabab, lid.arxiv_izoh = "", ""
            ozgardi.append("arxiv")
        lid.arxiv = qiymat
    if "qora_royxat" in data:
        qiymat = bool(data["qora_royxat"])
        if qiymat and not lid.qora_royxat:
            izoh = str(data.get("qora_royxat_izoh") or "").strip()[:300]
            if not izoh:
                raise ValueError("Qora ro'yxatga olish sababini yozing")
            lid.qora_royxat_izoh = izoh
            ozgardi.append("qora_royxat")
        elif not qiymat and lid.qora_royxat:
            lid.qora_royxat_izoh = ""
            ozgardi.append("qora_royxat")
        lid.qora_royxat = qiymat
    return ozgardi


def _lid_tarix_matni(lid, ozgardi, eski_bolim):
    """Lid tarixi odam tilida (video-TZ 2026-09-25: "O'zgartirildi: arxiv"
    arxivlandimi yoki chiqarildimi — bilinmasdi)."""
    qismlar = []
    if "arxiv" in ozgardi:
        if lid.arxiv:
            sabab = lid.get_arxiv_sabab_display()
            qismlar.append(f"Arxivlandi — {sabab}" + (f": {lid.arxiv_izoh}" if lid.arxiv_izoh else ""))
        else:
            qismlar.append("Arxivdan chiqarildi")
    if "qora_royxat" in ozgardi:
        qismlar.append(f"Qora ro'yxatga olindi: {lid.qora_royxat_izoh}" if lid.qora_royxat
                       else "Qora ro'yxatdan chiqarildi")
    if "bolim_id" in ozgardi:
        yangi = LidBolim.objects.filter(pk=lid.bolim_id).values_list("nomi", flat=True).first() or "—"
        qismlar.append(f"Bo'lim: {eski_bolim} → {yangi}")
    qolgan = [m for m in ozgardi if m not in ("arxiv", "qora_royxat", "bolim_id")]
    if qolgan:
        qismlar.append("O'zgartirildi: " + ", ".join(_MAYDON_NOMLARI.get(m, m) for m in qolgan))
    return "; ".join(qismlar)


_MAYDON_NOMLARI = {
    "ism": "ism", "telefon": "telefon", "holat": "holat", "bolim_id": "bo'lim", "arxiv": "arxiv",
    "qora_royxat": "qora ro'yxat", "filial_id": "filial", "kurs_id": "kurs", "manba": "manba",
    "izoh": "izoh", "qoshimcha_telefon": "qo'shimcha raqam", "qoshimcha_ism": "qo'shimcha ism",
    "qulay_vaqt": "qulay vaqt", "kunlar": "kunlar", "yigilayotgan_guruh_id": "yig'ilayotgan guruh",
}


class LidBolimlarView(CrmView):
    bolim = "lidlar"

    def get(self, request):
        qs = LidBolim.objects.all()
        # `?doska=ID` — shu doskaning ustunlari; `?doska=0` — doskasizlar
        # (video-TZ'dan oldin yaratilgan ustunlar "Umumiy" doskada).
        doska = request.query_params.get("doska")
        if doska == "0":
            qs = qs.filter(doska__isnull=True)
        elif doska:
            qs = qs.filter(doska_id=doska)
        filial = request.query_params.get("filial")
        if filial:
            qs = qs.filter(Q(filial_id=filial) | Q(filial__isnull=True))
        qs = qs.filter(filial_q(request.user, "filial")).annotate(_soni=Count(
            "lidlar",
            filter=Q(lidlar__arxiv=False, lidlar__qora_royxat=False) & filial_q(request.user, "lidlar__filial"),
        ))
        return Response([_bolim_dict(b, b._soni) for b in qs])

    def post(self, request):
        if xato := _ruxsatsiz(request, "lidlar.bolim"):
            return xato
        nomi = (request.data.get("nomi") or "").strip()
        if not nomi:
            return _xato("Bo'lim nomi bo'sh bo'lmasin")
        oxirgi = LidBolim.objects.order_by("-tartib").values_list("tartib", flat=True).first() or 0
        doska = None
        doska_id = request.data.get("doska_id") or None
        if doska_id:
            doska = LidDoska.objects.filter(pk=doska_id).first()
            if doska is None or not filial_korinadimi(request.user, doska.filial_id):
                return _xato("Doska topilmadi")
        # Filial doskasidagi ustun — o'sha filialniki. Filial xodimi
        # filialsiz ("hammaniki") ustun ocha olmaydi: bittasi bo'lsa o'zi qo'yiladi.
        filial_id = request.data.get("filial_id") or (doska.filial_id if doska else None)
        ruxsat = ruxsat_filiallari(request.user)
        if ruxsat is not None and not filial_id and len(ruxsat) == 1:
            filial_id = next(iter(ruxsat))
        if filial_id or ruxsat is not None:
            try:
                filial_tekshir(request.user, filial_id)
            except ValueError as e:
                return _xato(str(e), kod=403)
            if not str(filial_id).isdigit() or not Filial.objects.filter(pk=filial_id).exists():
                return _xato("Filial topilmadi")
        if doska and doska.filial_id and str(filial_id) != str(doska.filial_id):
            return _xato("Ustun filiali doska filiali bilan bir xil bo'lsin")
        guruh_id = request.data.get("guruh_id") or None
        if guruh_id:
            guruh = Guruh.objects.select_related("moliya").filter(pk=guruh_id).first()
            if guruh is None or not filial_korinadimi(
                request.user, getattr(getattr(guruh, "moliya", None), "filial_id", None)
            ):
                return _xato("Guruh topilmadi")
        bolim = LidBolim.objects.create(
            nomi=nomi[:100], tartib=oxirgi + 1, filial_id=filial_id, doska_id=doska_id, guruh_id=guruh_id,
        )
        return Response(_bolim_dict(bolim, 0), status=201)


class LidBolimDetailView(CrmView):
    pk_turi = "lid_bolim"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "lidlar"

    def patch(self, request, pk):
        if xato := _ruxsatsiz(request, "lidlar.bolim"):
            return xato
        bolim = get_object_or_404(LidBolim, pk=pk)
        try:
            umumiy_yozuv_taqiq(request.user, bolim.filial_id, "ustun")
        except ValueError as e:
            return _xato(str(e), kod=403)
        if "nomi" in request.data:
            nomi = (request.data.get("nomi") or "").strip()
            if not nomi:
                return _xato("Bo'lim nomi bo'sh bo'lmasin")
            bolim.nomi = nomi[:100]
        if "tartib" in request.data:
            try:
                bolim.tartib = int(request.data["tartib"])
            except (TypeError, ValueError):
                return _xato("tartib son bo'lsin")
        bolim.save()
        return Response(_bolim_dict(bolim))

    def delete(self, request, pk):
        if xato := _ruxsatsiz(request, "lidlar.bolim"):
            return xato
        bolim = get_object_or_404(LidBolim, pk=pk)
        try:
            umumiy_yozuv_taqiq(request.user, bolim.filial_id, "ustun")
        except ValueError as e:
            return _xato(str(e), kod=403)
        # Ichidagi lidlar YO'QOLMAYDI — bo'limsiz qoladi ("Yangi lidlar").
        bolim.delete()
        return Response(status=204)


def lidlar_qs(p, user=None):
    """Kanban va Excel eksport uchun BITTA filtr — eksport ekrandagi
    ro'yxatning aynan o'zi bo'lsin."""
    qs = Lid.objects.select_related("kurs", "oqituvchi", "kim_qoshdi", "bolim", "filial", "yigilayotgan_guruh")
    # Filial cheklovi. Qora ro'yxat — ISTISNO: u hamma filialga ko'rinadi
    # (Shuhrat, 2026-09-23), toki bir filialda "yomon" mijoz boshqasida
    # yangidan yozilmasin.
    if user is not None and not p.get("qora_royxat"):
        qs = qs.filter(filial_q(user, "filial"))
    qs = qs.filter(arxiv=bool(p.get("arxiv")))
    qs = qs.filter(qora_royxat=bool(p.get("qora_royxat")))
    if p.get("filial"):
        qs = qs.filter(Q(filial_id=p["filial"]) | Q(filial__isnull=True))
    if p.get("bolim"):
        qs = qs.filter(bolim_id=p["bolim"])
    # Doska: lid ustuni orqali. `doska=0` — ustunsiz yoki doskasiz ustundagi.
    if p.get("doska") == "0":
        qs = qs.filter(Q(bolim__isnull=True) | Q(bolim__doska__isnull=True))
    elif p.get("doska"):
        qs = qs.filter(bolim__doska_id=p["doska"])
    if p.get("manba"):
        qs = qs.filter(manba=p["manba"])
    if p.get("oqituvchi"):
        qs = qs.filter(oqituvchi_id=p["oqituvchi"])
    if p.get("kunlar"):
        qs = qs.filter(kunlar=p["kunlar"])
    if p.get("q"):
        q = p["q"].strip()
        qs = qs.filter(Q(ism__icontains=q) | Q(telefon__icontains=q) | Q(qoshimcha_telefon__icontains=q))
    return qs


class LidlarView(CrmView):
    """Lidlar kanbani: bo'limlar bo'yicha guruhlangan ro'yxat."""

    bolim = "lidlar"

    def get(self, request):
        lidlar = list(lidlar_qs(request.query_params, request.user)[:1000])
        # Har lidning OXIRGI eslatmasi bitta so'rovda (N+1 emas).
        eslatmalar = {}
        for e in Eslatma.objects.filter(lid__in=lidlar).order_by("lid_id", "-created_at").values(
            "lid_id", "matn", "eslatish_vaqti", "created_at"
        ):
            eslatmalar.setdefault(e["lid_id"], e)
        return Response([_lid_dict(lid, eslatmalar.get(lid.id)) for lid in lidlar])

    def post(self, request):
        if xato := _ruxsatsiz(request, "lidlar.qoshish", "Lid qo'shishga ruxsat yo'q"):
            return xato
        lid = Lid(kim_qoshdi=request.user)
        data = dict(request.data)
        data.setdefault("ism", "")
        data.setdefault("telefon", "")
        # Filial xodimi qo'shgan lid — uning filialida (bitta bo'lsa o'zi
        # qo'yiladi); filialsiz lid hammaga ko'rinardi.
        ruxsat = ruxsat_filiallari(request.user)
        if ruxsat is not None and not data.get("filial_id") and len(ruxsat) == 1:
            data["filial_id"] = next(iter(ruxsat))
        try:
            if ruxsat is not None:
                filial_tekshir(request.user, data.get("filial_id"))
            _lid_maydonlari(lid, data, request.user)
        except ValueError as e:
            return _xato(str(e))
        # Telefon bo'yicha takror — ogohlantiramiz, lekin to'xtatmaymiz
        # (bitta raqamdan aka-uka yozilishi mumkin), faqat qora ro'yxatdagi
        # raqamni kiritib bo'lmaydi.
        if Lid.objects.filter(telefon=lid.telefon, qora_royxat=True).exists() or TalabaProfil.objects.filter(
            qora_royxat=True, user__telefon=lid.telefon
        ).exists():
            if not request.data.get("majburan"):
                return _xato("Bu raqam qora ro'yxatda", kod=409)
        lid.save()
        LidTarix.objects.create(lid=lid, matn="Lid yaratildi", kim=request.user)
        # "Harakatlar tarixi"da o'chirish bor edi, yaratish yo'q edi.
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.YARATISH, obyekt=lid,
              obyekt_turi="CRM Lid", obyekt_nomi=str(lid), snapshot={"manba": lid.manba, "bolim": lid.bolim_id})
        if request.data.get("eslatma"):
            Eslatma.objects.create(lid=lid, matn=str(request.data["eslatma"])[:2000], kim=request.user)
        return Response(_lid_dict(lid), status=201)


class LidlarEksportView(CrmView):
    """Lidlar ro'yxati Excel'da — kanbandagi filtrlar bilan bir xil."""

    bolim = "lidlar"

    def get(self, request):
        if "lidlar.excel" not in ruxsatlar(request.user):
            return _xato("Excel eksportga ruxsat yo'q", kod=403)
        from openpyxl import Workbook

        from .eksport import _varaq_yoz, javob_qil

        eslatmalar = {}
        lidlar = list(lidlar_qs(request.query_params, request.user)[:5000])
        for e in Eslatma.objects.filter(lid__in=lidlar).order_by("lid_id", "-created_at").values("lid_id", "matn"):
            eslatmalar.setdefault(e["lid_id"], e["matn"])
        kitob = Workbook()
        ws = kitob.active
        ws.title = "Lidlar"
        _varaq_yoz(
            ws,
            ["ID", "Ism familiya", "Telefon", "Qo'shimcha raqam", "Kimning raqami", "Tug'ilgan sana",
             "Bo'lim", "Holat", "Manba", "Filial", "Kurs", "O'qituvchi", "Qulay vaqt", "Kunlar",
             "Izoh", "Oxirgi eslatma", "Kim qo'shdi", "Qo'shilgan vaqt"],
            [
                [
                    l.id, l.ism, l.telefon, l.qoshimcha_telefon, l.qoshimcha_ism,
                    l.tugilgan_sana.strftime("%d.%m.%Y") if l.tugilgan_sana else "",
                    l.bolim.nomi if l.bolim_id else "Yangi lidlar", l.get_holat_display(), l.manba,
                    l.filial.nomi if l.filial_id else "", l.kurs.nomi if l.kurs_id else "",
                    _ism(l.oqituvchi) if l.oqituvchi_id else "", l.qulay_vaqt, l.kunlar, l.izoh,
                    eslatmalar.get(l.id, ""), _ism(l.kim_qoshdi) if l.kim_qoshdi_id else "",
                    timezone.localtime(l.created_at).strftime("%d.%m.%Y %H:%M"),
                ]
                for l in lidlar
            ],
        )
        return javob_qil(kitob, f"lidlar_{timezone.localdate():%Y-%m-%d}.xlsx")


class LidDetailView(CrmView):
    pk_turi = "lid"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "lidlar"

    def get(self, request, pk):
        lid = get_object_or_404(Lid.objects.select_related("kurs", "oqituvchi", "kim_qoshdi"), pk=pk)
        # Boshqa filial lidi takrorlarda ko'rinmaydi — qora ro'yxatdagisidan tashqari.
        takrorlar = (
            Lid.objects.filter(telefon=lid.telefon).exclude(pk=lid.pk)
            .filter(filial_q(request.user, "filial") | Q(qora_royxat=True))
            .values("id", "ism", "arxiv")
        )
        return Response({
            **_lid_dict(lid),
            "tarix": [
                {"matn": t.matn, "kim": _ism(t.kim), "vaqt": t.created_at}
                for t in lid.tarix.select_related("kim")[:100]
            ],
            "takrorlar": list(takrorlar),
        })

    # Maydon -> kerakli amal ruxsati. Qolgan maydonlar — `lidlar.tahrirlash`.
    _MAYDON_RUXSATI = {
        "arxiv": "lidlar.arxiv", "arxiv_sabab": "lidlar.arxiv", "arxiv_izoh": "lidlar.arxiv",
        "qora_royxat": "lidlar.qora_royxat", "qora_royxat_izoh": "lidlar.qora_royxat",
        "yigilayotgan_guruh_id": "lidlar.guruhga",
    }

    def patch(self, request, pk):
        berilgan = ruxsatlar(request.user)
        for maydon in request.data:
            kerak = self._MAYDON_RUXSATI.get(maydon, "lidlar.tahrirlash")
            if kerak not in berilgan:
                return _xato("Bu o'zgarishga ruxsat yo'q", kod=403)
        lid = get_object_or_404(Lid, pk=pk)
        eski_bolim = lid.bolim.nomi if lid.bolim_id else "—"
        try:
            if "filial_id" in request.data and ruxsat_filiallari(request.user) is not None:
                filial_tekshir(request.user, request.data.get("filial_id"))
            ozgardi = _lid_maydonlari(lid, request.data, request.user)
        except ValueError as e:
            return _xato(str(e))
        lid.save()
        if ozgardi:
            LidTarix.objects.create(lid=lid, matn=_lid_tarix_matni(lid, ozgardi, eski_bolim)[:300], kim=request.user)
        return Response(_lid_dict(lid))

    def delete(self, request, pk):
        if "lidlar.ochirish" not in ruxsatlar(request.user):
            return _xato("O'chirishga ruxsat yo'q", kod=403)
        lid = get_object_or_404(Lid, pk=pk)
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.OCHIRISH, obyekt=lid,
              obyekt_turi="CRM Lid", obyekt_nomi=str(lid))
        lid.delete()
        return Response(status=204)


# ── Talaba yaratish va guruhga qo'shish ──────────────────────────────


def _birinchi_unit(guruh):
    """LMS qoidasi (`academics.views._azolarni_saqla`): yangi a'zo guruh
    darajasining BIRINCHI Unit'idan boshlaydi."""
    if not guruh.daraja_id:
        return None
    return (
        KursTugun.objects.filter(parent_id=guruh.daraja_id, unit_darsi=True)
        .order_by("tartib", "id")
        .first()
    )


def guruhga_qosh(talaba, guruh, *, sana=None, narx=None, holat=AzolikMoliya.Holat.FAOL):
    """Talabani guruhga qo'shadi: LMS a'zoligi (`GuruhAzoligi`) + CRM
    moliyasi (`AzolikMoliya`). Allaqachon a'zo bo'lsa — mavjudini
    qaytaradi (qayta qo'shish xato emas)."""
    azolik, yaratildi = GuruhAzoligi.objects.get_or_create(
        guruh=guruh, talaba=talaba, defaults={"boshlanish_unit": _birinchi_unit(guruh)}
    )
    am, _ = AzolikMoliya.objects.get_or_create(
        azolik=azolik,
        defaults={"boshlanish_sana": sana or timezone.localdate(), "narx": narx, "holat": holat},
    )
    return am, yaratildi


def talaba_yarat(data, kim):
    """CRM'dan yangi talaba (LMS foydalanuvchisi, role=student).

    Qaytaradi: `(user, parol, None)` yoki `(None, None, xato)`.
    Login/parol berilmasa avtomatik yaratiladi — admin har safar o'ylab
    topmasin, talabaga esa SMS/qog'ozda beriladi.
    """
    from accounts.views import _parolni_tekshir, _talaba_maydonlarini_yoz

    ism = (data.get("ism") or "").strip()
    if not ism:
        return None, None, "Ism familiya kiritilsin"
    telefon = _telefon(data.get("telefon"))
    login = (data.get("login") or "").strip() or login_yarat(ism, telefon)
    if not LOGIN_QOIDASI.match(login):
        return None, None, "Login faqat harf/raqam/./@/+/-/_ dan iborat bo'lsin"
    if User.objects.filter(username=login).exists():
        return None, None, f"«{login}» login band"
    parol = (data.get("parol") or "").strip() or parol_yarat()
    xatolar = _parolni_tekshir(parol)
    if xatolar:
        return None, None, " ".join(xatolar)

    user = User(username=login, first_name=ism[:150], role=User.Role.STUDENT)
    user.set_password(parol)
    user.save()
    maydonlar = {
        k: data[k] for k in ("telefon", "ota_ona_telefon", "ota_ona_ismi", "tugilgan_sana", "manba", "izoh")
        if k in data
    }
    if "telefon" in maydonlar:
        maydonlar["telefon"] = telefon
    if "ota_ona_telefon" in maydonlar:
        maydonlar["ota_ona_telefon"] = _telefon(maydonlar["ota_ona_telefon"])
    xato = _talaba_maydonlarini_yoz(user, maydonlar)
    if xato:
        raise ValueError(xato)
    jins = data.get("jins") or ""
    if jins and jins not in dict(TalabaProfil.Jins.choices):
        raise ValueError("Jins noto'g'ri")
    TalabaProfil.objects.create(user=user, jins=jins, maktab=(data.get("maktab") or "").strip()[:100])
    logla(foydalanuvchi=kim, harakat=FaoliyatYozuvi.Harakat.YARATISH, obyekt=user,
          obyekt_turi="Foydalanuvchi", snapshot={"username": login, "role": "student", "manba": "crm"})
    return user, parol, None


class TalabaYaratishView(CrmView):
    """"Yangi o'quvchi qo'shish" — talaba yaratiladi va (ixtiyoriy)
    darhol guruhga qo'shiladi."""

    bolim = "talabalar"

    def post(self, request):
        if xato := _ruxsatsiz(request, "talabalar.qoshish", "O'quvchi qo'shishga ruxsat yo'q"):
            return xato
        data = request.data
        guruh = None
        if data.get("guruh_id"):
            guruh = get_object_or_404(Guruh, pk=data["guruh_id"])
            guruh_tekshir(request.user, guruh)
        try:
            sana = _sana(data.get("boshlanish_sana"), "boshlanish_sana", majburiy=False)
            narx = None if data.get("narx") in (None, "") else _son(data["narx"], "narx")
        except ValueError as e:
            return _xato(str(e))
        holat = data.get("holat") or AzolikMoliya.Holat.FAOL
        if holat not in (AzolikMoliya.Holat.FAOL, AzolikMoliya.Holat.SINOV):
            return _xato("Holat: faol yoki sinov")

        try:
            with transaction.atomic():
                user, parol, xato = talaba_yarat(data, request.user)
                if xato:
                    return _xato(xato)
                if guruh is not None:
                    guruhga_qosh(user, guruh, sana=sana, narx=narx, holat=holat)
        except ValueError as e:
            return _xato(str(e))
        return Response(
            {"id": user.id, "ism": _ism(user), "username": user.username, "parol": parol},
            status=201,
        )


class TalabaCrmView(CrmView):
    """Talabaning CRM maydonlari: jins, maktab, qora ro'yxat, arxiv."""

    pk_turi = "talaba"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "talabalar"

    def patch(self, request, pk):
        data = request.data
        # Qora ro'yxatdan boshqa hamma narsa (jins, maktab, arxiv, parol) — tahrirlash.
        if set(data) - {"qora_royxat", "sabab"}:
            if xato := _ruxsatsiz(request, "talabalar.tahrirlash"):
                return xato
        talaba = get_object_or_404(User, pk=pk, role=User.Role.STUDENT)
        profil, _ = TalabaProfil.objects.get_or_create(user=talaba)
        if "qora_royxat" in data:
            if "talabalar.qora_royxat" not in ruxsatlar(request.user):
                return _xato("Qora ro'yxatga ruxsat yo'q", kod=403)
            profil.qora_royxat = bool(data["qora_royxat"])
            profil.qora_royxat_sabab = (data.get("sabab") or "").strip()[:300] if profil.qora_royxat else ""
        if "jins" in data:
            jins = data.get("jins") or ""
            if jins and jins not in dict(TalabaProfil.Jins.choices):
                return _xato("Jins noto'g'ri")
            profil.jins = jins
        if "maktab" in data:
            profil.maktab = (data.get("maktab") or "").strip()[:100]
        profil.save()
        if "faol" in data:
            faol = bool(data["faol"])
            # Arxivlash — SABAB bilan (video-TZ 2026-09-25), guruhdan
            # chiqarishdagi kabi: ro'yxatdan sabab + izoh (majburiy).
            if not faol and talaba.is_active:
                sabab = data.get("arxiv_sabab") or ""
                izoh = str(data.get("arxiv_izoh") or "").strip()[:300]
                if sabab not in KETISH_SABABLARI:
                    return _xato("Arxivlash sababini tanlang")
                if not izoh:
                    return _xato("Izoh yozilsin")
                profil.arxiv_sabab, profil.arxiv_izoh = sabab, izoh
            elif faol:
                profil.arxiv_sabab, profil.arxiv_izoh = "", ""
            profil.save(update_fields=["arxiv_sabab", "arxiv_izoh"])
            talaba.is_active = faol
            talaba.save(update_fields=["is_active"])
        yangi_parol = None
        if data.get("parol_tiklash"):
            yangi_parol = parol_yarat()
            talaba.set_password(yangi_parol)
            talaba.save(update_fields=["password"])
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.OZGARTIRISH, obyekt=talaba,
              obyekt_turi="Talaba", obyekt_nomi=_ism(talaba),
              ozgarishlar={k: {"eski": "—", "yangi": str(data[k])} for k in data if k != "parol_tiklash"})
        return Response({
            "id": talaba.id, "faol": talaba.is_active, "qora_royxat": profil.qora_royxat,
            "qora_royxat_sabab": profil.qora_royxat_sabab, "jins": profil.jins, "maktab": profil.maktab,
            "arxiv_sabab": profil.arxiv_sabab, "arxiv_izoh": profil.arxiv_izoh,
            "parol": yangi_parol,
        })


class LidGuruhgaView(CrmView):
    """Lid(lar)ni guruhga qo'shish — har biridan talaba yaratiladi.

    SoffCRM'dagi "Guruhga qo'shish" (bitta lid) va "Lidlarni guruhga
    qo'shish" (bir nechta). Lid arxivga o'tadi, holati "qo'shildi",
    `talaba` maydoni yangi talabaga ishora qiladi.
    """

    bolim = "lidlar"

    def post(self, request):
        if xato := _ruxsatsiz(request, "lidlar.guruhga"):
            return xato
        idlar = request.data.get("lid_idlar") or []
        if request.data.get("lid_id"):
            idlar = [request.data["lid_id"]]
        if not idlar:
            return _xato("Lid tanlanmagan")
        guruh = get_object_or_404(Guruh, pk=request.data.get("guruh_id"), faol=True)
        guruh_tekshir(request.user, guruh)
        holat = request.data.get("holat") or AzolikMoliya.Holat.SINOV
        if holat not in (AzolikMoliya.Holat.FAOL, AzolikMoliya.Holat.SINOV):
            return _xato("Holat: faol yoki sinov")
        try:
            sana = _sana(request.data.get("boshlanish_sana"), "boshlanish_sana", majburiy=False)
        except ValueError as e:
            return _xato(str(e))

        natija = []
        try:
            with transaction.atomic():
                for lid in Lid.objects.filter(pk__in=idlar).filter(filial_q(request.user, "filial")).select_for_update():
                    if lid.qora_royxat:
                        raise ValueError(f"{lid.ism} qora ro'yxatda")
                    parol = None
                    if lid.talaba_id:
                        talaba = lid.talaba
                    else:
                        talaba, parol, xato = talaba_yarat(
                            {
                                "ism": lid.ism, "telefon": lid.telefon,
                                "ota_ona_telefon": lid.qoshimcha_telefon,
                                "ota_ona_ismi": lid.qoshimcha_ism,
                                "tugilgan_sana": lid.tugilgan_sana.isoformat() if lid.tugilgan_sana else "",
                                "manba": lid.manba, "izoh": lid.izoh,
                            },
                            request.user,
                        )
                        if xato:
                            raise ValueError(f"{lid.ism}: {xato}")
                    guruhga_qosh(talaba, guruh, sana=sana, holat=holat)
                    lid.talaba = talaba
                    lid.holat = Lid.Holat.QOSHILDI
                    lid.arxiv = True
                    lid.yigilayotgan_guruh = None
                    lid.save(update_fields=["talaba", "holat", "arxiv", "yigilayotgan_guruh", "updated_at"])
                    LidTarix.objects.create(lid=lid, matn=f"«{guruh.name}» guruhiga qo'shildi", kim=request.user)
                    natija.append({
                        "lid_id": lid.id, "talaba_id": talaba.id, "ism": _ism(talaba),
                        "username": talaba.username, "parol": parol,
                    })
        except ValueError as e:
            return _xato(str(e))
        return Response({"qoshildi": natija}, status=201)


# ── Guruh yaratish va boshqarish ─────────────────────────────────────


def _oqituvchilarni_saqla(guruh, oqituvchilar):
    """`[{oqituvchi_id, turi, foiz}]` (ko'pi bilan 3). Asosiy o'qituvchi
    LMS'dagi `Guruh.oqituvchi`ga ham yoziladi — davomat va Kurslar unga
    tayanadi."""
    if not isinstance(oqituvchilar, list):
        raise ValueError("oqituvchilar ro'yxat bo'lishi kerak")
    oqituvchilar = [o for o in oqituvchilar if o.get("oqituvchi_id")]
    if len(oqituvchilar) > 3:
        raise ValueError("Ko'pi bilan 3 ta o'qituvchi")
    korilgan = set()
    yozuvlar = []
    for o in oqituvchilar:
        oid = int(o["oqituvchi_id"])
        if oid in korilgan:
            raise ValueError("Bir o'qituvchi ikki marta tanlangan")
        korilgan.add(oid)
        if not User.objects.filter(pk=oid, role=User.Role.TEACHER).exists():
            raise ValueError("O'qituvchi topilmadi")
        turi = o.get("turi") or GuruhOqituvchi.Turi.ASOSIY
        if turi not in dict(GuruhOqituvchi.Turi.choices):
            raise ValueError("O'qituvchi turi noto'g'ri")
        foiz = None if o.get("foiz") in (None, "") else _son(o["foiz"], "foiz")
        if foiz is not None and not 0 <= foiz <= 100:
            raise ValueError("Foiz 0..100 oralig'ida bo'lsin")
        ulush_turi = o.get("ulush_turi") or GuruhOqituvchi.UlushTuri.FOIZ
        if ulush_turi not in dict(GuruhOqituvchi.UlushTuri.choices):
            raise ValueError("Ulush turi noto'g'ri")
        dars_haqi = None if o.get("dars_haqi") in (None, "") else _son(o["dars_haqi"], "dars_haqi")
        if dars_haqi is not None and dars_haqi < 0:
            raise ValueError("Dars haqi manfiy bo'lmasin")
        if ulush_turi == GuruhOqituvchi.UlushTuri.DARS and dars_haqi is None:
            raise ValueError("Har dars uchun summa kiritilsin")
        yozuvlar.append(GuruhOqituvchi(guruh=guruh, oqituvchi_id=oid, turi=turi, foiz=foiz,
                                       ulush_turi=ulush_turi, dars_haqi=dars_haqi))
    asosiylar = [y for y in yozuvlar if y.turi == GuruhOqituvchi.Turi.ASOSIY]
    if len(asosiylar) > 1:
        raise ValueError("Asosiy o'qituvchi bitta bo'lsin")
    GuruhOqituvchi.objects.filter(guruh=guruh).delete()
    GuruhOqituvchi.objects.bulk_create(yozuvlar)
    asosiy = asosiylar[0] if asosiylar else (yozuvlar[0] if yozuvlar else None)
    guruh.oqituvchi_id = asosiy.oqituvchi_id if asosiy else None
    guruh.save(update_fields=["oqituvchi"])


def _fan_daraja(data):
    """(fan, daraja) — daraja berilsa, fan uning ota tuguni."""
    daraja = None
    fan = None
    if data.get("daraja_id"):
        daraja = get_object_or_404(KursTugun, pk=data["daraja_id"])
        fan = daraja.parent
    elif data.get("fan_id"):
        fan = get_object_or_404(KursTugun, pk=data["fan_id"])
    return fan, daraja


def _guruh_toliq(g):
    d = _guruh_dict(g)
    d.update({
        "fan_id": g.fan_id,
        "oqituvchilar": [
            {"oqituvchi_id": o.oqituvchi_id, "ism": _ism(o.oqituvchi), "turi": o.turi, "foiz": o.foiz,
             "ulush_turi": o.ulush_turi, "dars_haqi": o.dars_haqi}
            for o in g.crm_oqituvchilar.select_related("oqituvchi")
        ],
    })
    return d


class GuruhYaratishView(CrmView):
    """SoffCRM "Guruh qo'shish": nom, kurs, dars kunlari (har kunga
    alohida vaqt va xona), 3 tagacha o'qituvchi ulushi bilan, sanalar,
    filial, narx. Hammasi BITTA tranzaksiyada: jadval xonasi to'qnashsa,
    guruh ham yaratilmaydi."""

    bolim = "guruhlar"

    def post(self, request):
        if xato := _ruxsatsiz(request, "guruhlar.qoshish", "Guruh qo'shishga ruxsat yo'q"):
            return xato
        data = request.data
        nomi = (data.get("nomi") or "").strip()
        if not nomi:
            return _xato("Guruh nomi kiritilsin")
        try:
            fan, daraja = _fan_daraja(data)
            boshlanish = _sana(data.get("boshlanish_sana"), "boshlanish_sana")
            tugash = _sana(data.get("tugash_sana"), "tugash_sana", majburiy=False)
            narx = None if data.get("narx") in (None, "") else _son(data["narx"], "narx")
        except ValueError as e:
            return _xato(str(e))
        if tugash and tugash < boshlanish:
            return _xato("Tugash sanasi boshlanish sanasidan oldin bo'la olmaydi")
        filial_id = data.get("filial_id") or None
        ruxsat = ruxsat_filiallari(request.user)
        if ruxsat is not None and not filial_id and len(ruxsat) == 1:
            filial_id = next(iter(ruxsat))
        try:
            filial_tekshir(request.user, filial_id)
        except ValueError as e:
            return _xato(str(e), kod=403)
        filial = get_object_or_404(Filial, pk=filial_id) if filial_id else None
        baholash = data.get("baholash_tizimi") or ""
        if baholash not in ("", "5", "10", "100"):
            return _xato("Baholash tizimi noto'g'ri")

        try:
            with transaction.atomic():
                guruh = Guruh.objects.create(name=nomi[:200], markaz_id=_markaz_id(), fan=fan, daraja=daraja)
                GuruhMoliya.objects.create(
                    guruh=guruh, filial=filial, narx=narx, boshlanish_sana=boshlanish, tugash_sana=tugash,
                    baholash_tizimi=baholash,
                )
                yangilar, xato = jadvalni_tekshir(guruh, data.get("jadval") or [], request.user)
                if xato:
                    raise ValueError(xato)
                DarsJadvali.objects.bulk_create(yangilar)
                _oqituvchilarni_saqla(guruh, data.get("oqituvchilar") or [])
        except ValueError as e:
            return _xato(str(e))

        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.YARATISH, obyekt=guruh,
              obyekt_turi="Guruh", snapshot={"name": guruh.name, "manba": "crm"})
        guruh.refresh_from_db()
        return Response(_guruh_toliq(guruh), status=201)


class GuruhBoshqaruvView(CrmView):
    """Guruhni tahrirlash (nom, kurs, o'qituvchilar) va arxivlash."""

    pk_turi = "guruh"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"

    def get(self, request, pk):
        return Response(_guruh_toliq(get_object_or_404(Guruh, pk=pk)))

    def patch(self, request, pk):
        if xato := _ruxsatsiz(request, "guruhlar.tahrirlash", "Guruhni tahrirlashga ruxsat yo'q"):
            return xato
        guruh = get_object_or_404(Guruh, pk=pk)
        data = request.data
        eski = {"nomi": guruh.name, "faol": guruh.faol, "oqituvchi": _ism(guruh.oqituvchi)}
        try:
            with transaction.atomic():
                if "nomi" in data:
                    nomi = (data.get("nomi") or "").strip()
                    if not nomi:
                        raise ValueError("Guruh nomi bo'sh bo'lmasin")
                    guruh.name = nomi[:200]
                if "daraja_id" in data or "fan_id" in data:
                    guruh.fan, guruh.daraja = _fan_daraja(data)
                if "faol" in data:
                    guruh.faol = bool(data["faol"])
                guruh.save()
                if "oqituvchilar" in data:
                    _oqituvchilarni_saqla(guruh, data["oqituvchilar"])
        except ValueError as e:
            return _xato(str(e))
        guruh.refresh_from_db()
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.OZGARTIRISH, obyekt=guruh,
              obyekt_turi="Guruh", obyekt_nomi=guruh.name, eski_qiymatlar={k: str(v) for k, v in eski.items()},
              yangi_qiymatlar={"nomi": guruh.name, "faol": str(guruh.faol), "oqituvchi": str(_ism(guruh.oqituvchi))})
        return Response(_guruh_toliq(guruh))


class GuruhTalabalariView(CrmView):
    """"Guruhga o'quvchi qo'shish" (qo'lda yoki Excel) va chiqarish."""

    pk_turi = "guruh"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"
    parser_classes = [JSONParser, MultiPartParser, FormParser]

    def post(self, request, pk):
        if xato := _ruxsatsiz(request, "guruhlar.talaba_qoshish"):
            return xato
        guruh = get_object_or_404(Guruh, pk=pk)
        try:
            sana = _sana(request.data.get("sana"), "sana", majburiy=False) or timezone.localdate()
            narx = None if request.data.get("narx") in (None, "") else _son(request.data["narx"], "narx")
        except ValueError as e:
            return _xato(str(e))
        holat = request.data.get("holat") or AzolikMoliya.Holat.FAOL
        if holat not in (AzolikMoliya.Holat.FAOL, AzolikMoliya.Holat.SINOV):
            return _xato("Holat: faol yoki sinov")

        fayl = request.FILES.get("excel_fayl")
        if fayl is not None:
            return self._excel(request, guruh, fayl, sana, holat)

        idlar = request.data.get("talaba_idlar") or []
        if request.data.get("talaba_id"):
            idlar = [request.data["talaba_id"]]
        talabalar = list(User.objects.filter(pk__in=idlar, role=User.Role.STUDENT))
        if not talabalar:
            return _xato("O'quvchi tanlanmagan")
        for t in talabalar:
            talaba_tekshir(request.user, t.id)
        qoshildi = []
        with transaction.atomic():
            for t in talabalar:
                am, yangi = guruhga_qosh(t, guruh, sana=sana, narx=narx, holat=holat)
                if yangi:
                    qoshildi.append(t.id)
                    if "izoh" in request.data and request.data["izoh"]:
                        Eslatma.objects.create(guruh=guruh, talaba=t, matn=str(request.data["izoh"])[:2000],
                                               kim=request.user)
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.OZGARTIRISH, obyekt=guruh,
              obyekt_turi="Guruh", obyekt_nomi=guruh.name,
              ozgarishlar={"qoshildi": {"eski": "—", "yangi": qoshildi}})
        return Response({"qoshildi": qoshildi, "allaqachon": [t.id for t in talabalar if t.id not in qoshildi]},
                        status=201)

    def _excel(self, request, guruh, fayl, sana, holat):
        """Excel: A=ism, B=telefon, C=ota-ona telefoni (1-qator sarlavha).
        Telefon bo'yicha mavjud talaba topilsa — o'shasi qo'shiladi,
        yangisi yaratilmaydi."""
        import openpyxl

        try:
            varaq = openpyxl.load_workbook(fayl, read_only=True, data_only=True).active
            qatorlar = [r for i, r in enumerate(varaq.iter_rows(values_only=True)) if i > 0 and r and any(r)]
        except Exception:
            return _xato("Excel fayl noto'g'ri formatda")
        natija, xatolar = [], []
        for n, r in enumerate(qatorlar, start=2):
            ism = str(r[0] or "").strip() if len(r) > 0 else ""
            telefon = _telefon(r[1]) if len(r) > 1 else ""
            ota_ona = _telefon(r[2]) if len(r) > 2 else ""
            if not ism:
                xatolar.append({"qator": n, "xato": "ism bo'sh"})
                continue
            try:
                with transaction.atomic():
                    talaba = (
                        User.objects.filter(role=User.Role.STUDENT, telefon=telefon).first() if telefon else None
                    )
                    # Qo'lda qo'shishdagi kabi filial cheklovi: boshqa filial
                    # talabasini telefon orqali o'z guruhiga tortib bo'lmaydi.
                    if talaba is not None and not talaba_korinadimi(request.user, talaba.id):
                        raise ValueError("Bu telefon boshqa filial o'quvchisiga tegishli")
                    parol = None
                    if talaba is None:
                        talaba, parol, xato = talaba_yarat(
                            {"ism": ism, "telefon": telefon, "ota_ona_telefon": ota_ona}, request.user
                        )
                        if xato:
                            raise ValueError(xato)
                    guruhga_qosh(talaba, guruh, sana=sana, holat=holat)
                natija.append({"ism": _ism(talaba), "username": talaba.username, "parol": parol})
            except ValueError as e:
                xatolar.append({"qator": n, "xato": str(e)})
        return Response({"qoshildi": natija, "xatolar": xatolar}, status=201)

    def delete(self, request, pk):
        """Guruhdan chiqarish (`?talaba=&sana=&sabab_turi=&sabab=`) —
        `guruhdan_chiqar` izohiga qara."""
        if xato := _ruxsatsiz(request, "guruhlar.talaba_qoshish"):
            return xato
        guruh = get_object_or_404(Guruh, pk=pk)
        talaba = get_object_or_404(User, pk=request.query_params.get("talaba"))
        azolik = get_object_or_404(GuruhAzoligi, guruh=guruh, talaba=talaba)
        sabab_turi = request.query_params.get("sabab_turi") or GuruhdanChiqish.SababTuri.BOSHQA
        # Bitirdi / ko'chirildi / lidga — o'z amallari orqali (`GuruhTalabaAmaliView`).
        if sabab_turi not in KETISH_SABABLARI:
            return _xato("Ketish sababi noto'g'ri")
        try:
            sana = _sana(request.query_params.get("sana"), "sana", majburiy=False) or timezone.localdate()
        except ValueError as e:
            return _xato(str(e))
        with transaction.atomic():
            guruhdan_chiqar(request.user, azolik, sana, sabab_turi, request.query_params.get("sabab"))
        return Response(status=204)


# Guruhdan "ketish" sabablari — chiqarish oynasidagi ro'yxat.
KETISH_SABABLARI = (
    GuruhdanChiqish.SababTuri.JOYLASHUV, GuruhdanChiqish.SababTuri.NARX, GuruhdanChiqish.SababTuri.NATIJA,
    GuruhdanChiqish.SababTuri.DARS_JADVALI, GuruhdanChiqish.SababTuri.OQITUVCHI,
    GuruhdanChiqish.SababTuri.BOSHQA,
)


def guruhdan_chiqar(kim, azolik, sana, sabab_turi, izoh="", *, tugash_sana=None):
    """Talabani guruhdan chiqaradi va `GuruhdanChiqish` yozuvini qaytaradi.

    Joriy oy chiqish sanasigacha qayta hisoblanadi, keyin LMS a'zoligi
    o'chadi (LMS'dagi kabi). Pul tarixi (`Hisob`/`Tolov`) qoladi — ular
    a'zolikka emas, talaba va guruhga bog'langan. Chaqiruvchi
    `transaction.atomic()` ichida chaqiradi.

    `tugash_sana` — hisob-kitob uchun oxirgi kun (standart: `sana`).
    Ko'chirishda u bir kun oldin: shu kundan yangi guruh hisoblaydi.
    """
    guruh, talaba = azolik.guruh, azolik.talaba
    tugash_sana = tugash_sana or sana
    am = getattr(azolik, "moliya", None)
    moliya = getattr(guruh, "moliya", None)
    oylik_narx = chegirma_bor = None
    if am is not None:
        oylik_narx = mantiq.amaldagi_narx(am, mantiq.oy_boshi(sana))
        chegirma_bor = bool(am.narx is not None or mantiq.chegirma_narxi(am, mantiq.oy_boshi(sana)) is not None)
        am.tugash_sana = tugash_sana
        am.save(update_fields=["tugash_sana"])
        mantiq.azolikni_qayta_hisobla(am, mantiq.oy_boshi(tugash_sana))
    chiqish = GuruhdanChiqish.objects.create(
        talaba=talaba, talaba_ism=_ism(talaba), guruh=guruh, guruh_nomi=guruh.name,
        filial=moliya.filial if moliya else None,
        boshlagan_sana=am.boshlanish_sana if am else azolik.created_at.date(),
        sana=sana, sabab_turi=sabab_turi, sabab=(izoh or "").strip()[:300],
        oylik_narx=oylik_narx, chegirma_bor=bool(chegirma_bor), kim=kim,
    )
    azolik.delete()
    logla(foydalanuvchi=kim, harakat=FaoliyatYozuvi.Harakat.OZGARTIRISH, obyekt=guruh,
          obyekt_turi="Guruh", obyekt_nomi=guruh.name,
          ozgarishlar={"chiqarildi": {"eski": _ism(talaba), "yangi": f"{sana} · {chiqish.get_sabab_turi_display()}"}})
    return chiqish


class GuruhTalabaAmaliView(CrmView):
    """Guruhdagi o'quvchi ustidagi ⋮ harakatlar (video-TZ 2026-09-25,
    SoffCRM guruh sahifasi): `kochirish` (boshqa guruhga), `bitirdi`,
    `lidga` (lidlarga qaytarish).

    POST {amal, talaba_id, sana?, izoh?, yangi_guruh_id?}
    """

    pk_turi = "guruh"
    bolim = "guruhlar"

    def post(self, request, pk):
        if xato := _ruxsatsiz(request, "guruhlar.talaba_qoshish"):
            return xato
        guruh = get_object_or_404(Guruh, pk=pk)
        azolik = get_object_or_404(
            GuruhAzoligi.objects.select_related("guruh", "talaba"), guruh=guruh, talaba_id=request.data.get("talaba_id")
        )
        amal = request.data.get("amal")
        izoh = request.data.get("izoh") or ""
        try:
            sana = _sana(request.data.get("sana"), "sana", majburiy=False) or timezone.localdate()
        except ValueError as e:
            return _xato(str(e))

        if amal == "kochirish":
            yangi = Guruh.objects.filter(pk=request.data.get("yangi_guruh_id"), faol=True).first()
            if yangi is None or yangi.pk == guruh.pk:
                return _xato("Yangi guruh tanlanmagan")
            guruh_tekshir(request.user, yangi)
            if GuruhAzoligi.objects.filter(guruh=yangi, talaba=azolik.talaba).exists():
                return _xato("O'quvchi bu guruhda allaqachon bor")
            am = getattr(azolik, "moliya", None)
            # Individual narx va holat (sinov/faol/muzlatilgan) yangi guruhga
            # ham o'tadi — ko'chirish o'quvchining shartini o'zgartirmaydi.
            narx = am.narx if am else None
            holat = am.holat if am and am.holat != AzolikMoliya.Holat.ARXIV else AzolikMoliya.Holat.FAOL
            talaba = azolik.talaba
            with transaction.atomic():
                guruhdan_chiqar(request.user, azolik, sana, GuruhdanChiqish.SababTuri.KOCHIRILDI,
                                f"→ {yangi.name}" + (f". {izoh}" if izoh else ""),
                                tugash_sana=sana - timedelta(days=1))
                guruhga_qosh(talaba, yangi, sana=sana, narx=narx, holat=holat)
            return Response({"yangi_guruh_id": yangi.pk})

        if amal == "bitirdi":
            with transaction.atomic():
                guruhdan_chiqar(request.user, azolik, sana, GuruhdanChiqish.SababTuri.BITIRDI, izoh)
            return Response({})

        if amal == "lidga":
            talaba = azolik.talaba
            moliya = getattr(guruh, "moliya", None)
            with transaction.atomic():
                guruhdan_chiqar(request.user, azolik, sana, GuruhdanChiqish.SababTuri.LIDGA, izoh)
                # Shu talabadan kelgan lid bo'lsa — o'sha qayta ochiladi
                # (tarixi bilan), bo'lmasa yangisi. `talaba` bog'lanadi:
                # keyin "Guruhga qo'shish" yangi hisob OCHMAYDI, shu talabani
                # qaytaradi (`LidGuruhgaView`).
                lid = Lid.objects.filter(talaba=talaba).order_by("-id").first()
                if lid is None:
                    lid = Lid.objects.create(
                        ism=_ism(talaba), telefon=talaba.telefon or "", talaba=talaba,
                        filial=moliya.filial if moliya else None, kim_qoshdi=request.user,
                    )
                lid.arxiv = False
                lid.arxiv_sabab = lid.arxiv_izoh = ""
                lid.holat = Lid.Holat.YANGI
                lid.bolim = None
                lid.save()
                LidTarix.objects.create(
                    lid=lid, kim=request.user,
                    matn=(f"«{guruh.name}» guruhidan lidlarga qaytarildi" + (f": {izoh}" if izoh else ""))[:300],
                )
            return Response({"lid_id": lid.id})

        return _xato("amal: kochirish | bitirdi | lidga")


# ── Davomat (CRM'dan belgilash) ──────────────────────────────────────


def guruh_dars_sanalari(guruh, oy):
    """Oydagi haqiqiy dars sanalari: haftalik jadval + qo'shimcha /
    ko'chirilgan darslar − bekor qilingan / ko'chirib ketilganlari."""
    sanalar = set(mantiq.oylik_dars_kunlari(guruh, oy))
    oxiri = mantiq.oy_oxiri(oy)
    for oz in guruh.crm_dars_ozgarishlari.all():
        if oz.asl_sana and oy <= oz.asl_sana <= oxiri and oz.turi in ("kochirish", "bekor"):
            sanalar.discard(oz.asl_sana)
        if oz.yangi_sana and oy <= oz.yangi_sana <= oxiri and oz.turi in ("kochirish", "qoshimcha"):
            sanalar.add(oz.yangi_sana)
    return sanalar


class GuruhDavomatView(CrmView):
    """Guruh davomati — oy bo'yicha jadval, CRM'dan BELGILANADI.

    Avval (2026-09-15) davomat faqat LMS'da belgilanardi va CRM uni
    o'qirdi. Video-TZ (2026-09-23): SoffCRM'dagidek admin ham belgilaydi.
    Ikki xil raqam muammosi YO'Q — yozuv AYNAN o'sha `academics.Davomat`
    jadvaliga tushadi, o'qituvchi LMS'da ham, admin CRM'da ham bitta
    yozuvni ko'radi va o'zgartiradi.

    Ustunlar — oyning dars sanalari (jadval + ko'chirishlar) VA
    jadvaldan tashqari belgilangan sanalar (eski yozuvlar yo'qolmasin).
    """

    pk_turi = "guruh"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"

    def get(self, request, pk):
        guruh = get_object_or_404(Guruh, pk=pk)
        try:
            oy = _oy(request.query_params.get("oy") or timezone.localdate().strftime("%Y-%m"))
        except ValueError as e:
            return _xato(str(e))

        yozuvlar = list(
            Davomat.objects.filter(guruh=guruh, sana__gte=oy, sana__lte=mantiq.oy_oxiri(oy))
            .select_related("crm_izoh")
        )
        sanalar = sorted(guruh_dars_sanalari(guruh, oy) | {y.sana for y in yozuvlar})
        katak = {}
        for y in yozuvlar:
            izoh = getattr(y, "crm_izoh", None)
            holat = y.holat
            if holat == Davomat.Holat.KELMADI and izoh and izoh.sababli:
                holat = "sababli"
            katak[(y.talaba_id, y.sana)] = {"holat": holat, "izoh": izoh.izoh if izoh else ""}

        ozgarishlar = [
            _ozgarish_dict(o)
            for o in guruh.crm_dars_ozgarishlari.filter(
                Q(asl_sana__range=(oy, mantiq.oy_oxiri(oy))) | Q(yangi_sana__range=(oy, mantiq.oy_oxiri(oy)))
            )
        ]
        azoliklar = GuruhAzoligi.objects.filter(guruh=guruh).select_related("talaba", "moliya")
        bugun = timezone.localdate()
        talabalar = []
        for a in sorted(azoliklar, key=lambda a: (_ism(a.talaba) or "").lower()):
            am = getattr(a, "moliya", None)
            boshi = am.boshlanish_sana if am else None
            kunlar = []
            for s in sanalar:
                k = katak.get((a.talaba_id, s))
                kunlar.append({
                    "sana": s,
                    "holat": k["holat"] if k else None,
                    "izoh": k["izoh"] if k else "",
                    # Talaba hali qo'shilmagan kun — qulf (SoffCRM'dagi 🔒).
                    # Yozuvi BOR katak qulflanmaydi: eski ma'lumot
                    # yashirinib, lekin "keldi" soniga kirib qolmasin.
                    "qulf": bool(boshi and s < boshi and not k),
                    "kelajak": s > bugun,
                })
            talabalar.append({
                "id": a.talaba_id,
                "ism": _ism(a.talaba),
                "holat": am.holat if am else None,
                "kunlar": kunlar,
                "keldi": sum(1 for k in kunlar if k["holat"] == "keldi"),
                "kelmadi": sum(1 for k in kunlar if k["holat"] == "kelmadi"),
                "sababli": sum(1 for k in kunlar if k["holat"] == "sababli"),
            })
        return Response({"oy": oy, "sanalar": sanalar, "talabalar": talabalar, "ozgarishlar": ozgarishlar})

    def post(self, request, pk):
        """Bitta katak: `{talaba_id, sana, holat: keldi|kelmadi|sababli|null, izoh}`.
        `holat: null` — belgini olib tashlash."""
        if "guruhlar.davomat" not in ruxsatlar(request.user):
            return _xato("Davomat belgilashga ruxsat yo'q", kod=403)
        guruh = get_object_or_404(Guruh, pk=pk)
        talaba = get_object_or_404(User, pk=request.data.get("talaba_id"))
        if not GuruhAzoligi.objects.filter(guruh=guruh, talaba=talaba).exists():
            return _xato("Talaba bu guruhda emas")
        try:
            sana = _sana(request.data.get("sana"), "sana")
        except ValueError as e:
            return _xato(str(e))
        if sana > timezone.localdate():
            return _xato("Kelajakdagi darsga davomat qo'yib bo'lmaydi")
        holat = request.data.get("holat")
        if holat not in (None, "", "keldi", "kelmadi", "sababli"):
            return _xato("Holat noto'g'ri")

        yozuv = Davomat.objects.filter(guruh=guruh, talaba=talaba, sana=sana).first()
        if not holat:
            if yozuv:
                yozuv.delete()
            return Response({"holat": None})

        with transaction.atomic():
            if yozuv is None:
                yozuv = Davomat(guruh=guruh, talaba=talaba, sana=sana)
            yozuv.holat = Davomat.Holat.KELDI if holat == "keldi" else Davomat.Holat.KELMADI
            yozuv.belgilagan = request.user
            yozuv.save()
            izoh_matn = (request.data.get("izoh") or "").strip()[:300]
            if holat == "sababli" or izoh_matn:
                DavomatIzoh.objects.update_or_create(
                    davomat=yozuv, defaults={"sababli": holat == "sababli", "izoh": izoh_matn}
                )
            else:
                DavomatIzoh.objects.filter(davomat=yozuv).delete()
        return Response({"holat": holat, "izoh": izoh_matn})


# ── Darsni ko'chirish / qo'shimcha dars ──────────────────────────────


def _ozgarish_dict(o):
    return {
        "id": o.id, "turi": o.turi, "turi_nomi": o.get_turi_display(),
        "asl_sana": o.asl_sana, "yangi_sana": o.yangi_sana,
        "boshlanish_vaqti": o.boshlanish_vaqti.strftime("%H:%M") if o.boshlanish_vaqti else None,
        "tugash_vaqti": o.tugash_vaqti.strftime("%H:%M") if o.tugash_vaqti else None,
        "xona_id": o.xona_id, "mavzu": o.mavzu, "izoh": o.izoh,
        "kim": _ism(o.kim) if o.kim_id else None, "vaqt": o.created_at,
    }


def _dars_yozuvlarini_kochir(guruh, dan, ga):
    """Ko'chirilgan dars bilan birga uning davomati (izohi bilan), mavzusi
    va baholari ham yangi sanaga o'tadi (SoffCRM: "davomat belgilari,
    mavzu ... dars bilan birga ko'chadi"). Yangi sanada o'z yozuvi bo'lsa —
    ustidan yozilmaydi, xato."""
    modellar = (Davomat, DarsMavzusi, DarsBahosi)
    for model in modellar:
        if model.objects.filter(guruh=guruh, sana=ga).exists():
            raise ValueError(f"{ga:%d.%m.%Y} kuni bu guruhda davomat/mavzu/baho allaqachon bor")
    for model in modellar:
        model.objects.filter(guruh=guruh, sana=dan).update(sana=ga)


class DarsOzgarishlariView(CrmView):
    pk_turi = "guruh"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"

    def get(self, request, pk):
        guruh = get_object_or_404(Guruh, pk=pk)
        return Response([_ozgarish_dict(o) for o in guruh.crm_dars_ozgarishlari.select_related("kim")[:200]])

    def post(self, request, pk):
        if "guruhlar.dars_kochirish" not in ruxsatlar(request.user):
            return _xato("Darsni ko'chirishga ruxsat yo'q", kod=403)
        guruh = get_object_or_404(Guruh, pk=pk)
        turi = request.data.get("turi")
        if turi not in dict(DarsOzgarish.Turi.choices):
            return _xato("Noma'lum tur")
        try:
            asl = _sana(request.data.get("asl_sana"), "asl_sana", majburiy=turi in ("kochirish", "bekor"))
            yangi = _sana(request.data.get("yangi_sana"), "yangi_sana", majburiy=turi in ("kochirish", "qoshimcha"))
            boshlanish = _vaqt(request.data["boshlanish_vaqti"], "boshlanish_vaqti") if request.data.get("boshlanish_vaqti") else None
            tugash = _vaqt(request.data["tugash_vaqti"], "tugash_vaqti") if request.data.get("tugash_vaqti") else None
        except ValueError as e:
            return _xato(str(e))
        if turi in ("kochirish", "bekor") and asl not in guruh_dars_sanalari(guruh, mantiq.oy_boshi(asl)):
            return _xato("Tanlangan kunda bu guruhning darsi yo'q")
        if boshlanish and tugash and tugash <= boshlanish:
            return _xato("Tugash vaqti boshlanishdan keyin bo'lsin")
        if yangi:
            # Video (25:15): o'tgan kunlar, darsi bor kunlar va guruh
            # tugaganidan keyingi kunlar tanlanmaydi. Qo'shimcha dars
            # o'tgan kunga yozilishi mumkin (bo'lib o'tgan darsni kiritish).
            if turi == "kochirish" and yangi < timezone.localdate():
                return _xato("Darsni o'tgan kunga ko'chirib bo'lmaydi")
            if yangi in guruh_dars_sanalari(guruh, mantiq.oy_boshi(yangi)):
                return _xato(f"{yangi:%d.%m.%Y} kuni bu guruhning darsi allaqachon bor")
            moliya = getattr(guruh, "moliya", None)
            if moliya and moliya.tugash_sana and yangi > moliya.tugash_sana:
                return _xato("Yangi sana guruh tugash sanasidan keyin")
            if moliya and moliya.boshlanish_sana and yangi < moliya.boshlanish_sana:
                return _xato("Yangi sana guruh boshlanishidan oldin")
        xona_id = request.data.get("xona_id") or None
        if xona_id and not Xona.objects.filter(filial_q(request.user, "filial", filialsiz_ham=False),
                                               pk=xona_id).exists():
            return _xato("Xona topilmadi")
        try:
            with transaction.atomic():
                if turi == "kochirish":
                    _dars_yozuvlarini_kochir(guruh, asl, yangi)
                oz = DarsOzgarish.objects.create(
                    guruh=guruh, turi=turi, asl_sana=asl, yangi_sana=yangi, boshlanish_vaqti=boshlanish,
                    tugash_vaqti=tugash, xona_id=xona_id, mavzu=(request.data.get("mavzu") or "").strip()[:200],
                    izoh=(request.data.get("izoh") or "").strip()[:300], kim=request.user,
                )
        except ValueError as e:
            return _xato(str(e))
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.YARATISH, obyekt=guruh,
              obyekt_turi="CRM Dars o'zgarishi", obyekt_nomi=guruh.name,
              snapshot={k: str(v) for k, v in _ozgarish_dict(oz).items()})
        return Response(_ozgarish_dict(oz), status=201)


class DarsOzgarishDetailView(CrmView):
    pk_turi = "dars_ozgarish"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"

    def delete(self, request, pk):
        if xato := _ruxsatsiz(request, "guruhlar.dars_kochirish"):
            return xato
        oz = get_object_or_404(DarsOzgarish, pk=pk)
        try:
            with transaction.atomic():
                # Ko'chirish bekor qilinsa — yozuvlar asl sanaga qaytadi.
                if oz.turi == DarsOzgarish.Turi.KOCHIRISH and oz.asl_sana and oz.yangi_sana:
                    _dars_yozuvlarini_kochir(oz.guruh, oz.yangi_sana, oz.asl_sana)
                oz.delete()
        except ValueError as e:
            return _xato(str(e))
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.OCHIRISH, obyekt=oz.guruh,
              obyekt_turi="CRM Dars o'zgarishi", obyekt_nomi=oz.guruh.name,
              snapshot={"turi": oz.turi, "asl_sana": str(oz.asl_sana), "yangi_sana": str(oz.yangi_sana)})
        return Response(status=204)


# ── Chegirmalar ──────────────────────────────────────────────────────


def _chegirma_dict(c):
    oxirgi_oy = c.boshlanish_oy
    for _ in range(c.oylar_soni - 1):
        oxirgi_oy = mantiq.keyingi_oy(oxirgi_oy)
    bugun_oy = mantiq.oy_boshi(timezone.localdate())
    otgan = (bugun_oy.year - c.boshlanish_oy.year) * 12 + (bugun_oy.month - c.boshlanish_oy.month)
    doimiy = c.oylar_soni == 0
    return {
        "id": c.id, "azolik_moliya_id": c.azolik_id, "narx": c.narx, "boshlanish_oy": c.boshlanish_oy,
        "oxirgi_oy": None if doimiy else oxirgi_oy, "oylar_soni": c.oylar_soni, "doimiy": doimiy,
        "qolgan_oylar": None if doimiy else max(0, min(c.oylar_soni, c.oylar_soni - max(0, otgan))),
        "izoh": c.izoh, "kim": _ism(c.kim) if c.kim_id else None, "vaqt": c.created_at,
    }


def _chegirma_oylarini_qayta_hisobla(am, boshlanish_oy, oylar_soni, dan=None):
    """Chegirma oynasidagi ALLAQACHON OCHILGAN, to'lanmagan hisoblar
    yangi narx bilan qayta hisoblanadi (kelgusilari generatsiyada o'zi
    oladi). To'langan va `qolda` oyga tegilmaydi — `azolikni_qayta_hisobla`
    qoidasi. `dan` — shu oydan oldingilari o'tkazib yuboriladi (owner
    bo'lmaganda o'tgan oylar qarzi o'zgarmasin)."""
    oy = boshlanish_oy
    # Doimiy chegirma (0 oy) — joriy oygacha ochilgan hamma oylar.
    oxirgi = mantiq.oy_boshi(timezone.localdate())
    soni = oylar_soni or max(1, (oxirgi.year - oy.year) * 12 + oxirgi.month - oy.month + 1)
    otkazildi = []
    for _ in range(soni):
        if dan is None or oy >= dan:
            hisob = mantiq.azolikni_qayta_hisobla(am, oy)
            if hisob is not None and hisob.qolda and hisob.holat != Hisob.Holat.TOLANDI:
                otkazildi.append(oy)
        oy = mantiq.keyingi_oy(oy)
    return otkazildi


def _qolda_ogohlantirish(oylar):
    """Chegirma `qolda` hisobga TEGMAYDI (qaror) — lekin jim ham qolmasin."""
    if not oylar:
        return None
    return (", ".join(f"{o:%Y-%m}" for o in oylar)
            + " hisobi qo'lda belgilangan — chegirma unga tegmadi. Kerak bo'lsa, summani owner tuzatadi.")


class GuruhChegirmalariView(CrmView):
    """SoffCRM guruhdagi "Chegirmalar" tabi: har talaba qatori, uning
    chegirmalari va "+" tugmasi."""

    pk_turi = "guruh"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"

    def get(self, request, pk):
        guruh = get_object_or_404(Guruh, pk=pk)
        azoliklar = (
            AzolikMoliya.objects.filter(azolik__guruh=guruh)
            .select_related("azolik__talaba", "azolik__guruh__moliya", "azolik__guruh__daraja__crm_narxi")
            .prefetch_related("chegirmalar__kim")
        )
        return Response([
            {
                "azolik_moliya_id": am.id,
                "talaba_id": am.azolik.talaba_id,
                "talaba": _ism(am.azolik.talaba),
                "narx": mantiq.guruh_narxi(guruh)[0],
                "narx_talabaga": am.narx,
                "chegirmalar": [_chegirma_dict(c) for c in am.chegirmalar.all()],
            }
            for am in sorted(azoliklar, key=lambda a: (_ism(a.azolik.talaba) or "").lower())
        ])

    def post(self, request, pk):
        if "guruhlar.chegirma" not in ruxsatlar(request.user):
            return _xato("Chegirma berishga ruxsat yo'q", kod=403)
        am = get_object_or_404(AzolikMoliya, pk=request.data.get("azolik_moliya_id"), azolik__guruh_id=pk)
        try:
            oy = _oy(request.data.get("boshlanish_oy") or timezone.localdate().strftime("%Y-%m"), "boshlanish_oy")
            xom_oylar = request.data.get("oylar_soni")
            oylar = 1 if xom_oylar in (None, "") else int(xom_oylar)
            # Video (28:05): chegirma SUMMADA (narxdan ayiriladi) yoki
            # FOIZDA beriladi; yoki to'g'ridan-to'g'ri yangi narx.
            if request.data.get("foiz") not in (None, ""):
                foiz = _son(request.data["foiz"], "foiz")
                if not 0 <= foiz <= 100:
                    raise ValueError("Foiz 0..100 oralig'ida bo'lsin")
                asl = am.narx if am.narx is not None else mantiq.guruh_narxi(am.azolik.guruh)[0]
                if asl is None:
                    raise ValueError("Guruh narxi belgilanmagan")
                narx = (asl * (Decimal(100) - foiz) / Decimal(100)).quantize(Decimal("1"))
            elif request.data.get("chegirma_summasi") not in (None, ""):
                asl = am.narx if am.narx is not None else mantiq.guruh_narxi(am.azolik.guruh)[0]
                if asl is None:
                    raise ValueError("Guruh narxi belgilanmagan")
                narx = asl - _son(request.data["chegirma_summasi"], "chegirma_summasi")
            else:
                narx = _son(request.data.get("narx"), "narx")
        except (ValueError, TypeError) as e:
            return _xato(str(e))
        if narx < 0:
            return _xato("Chegirma narxdan katta bo'lmasin")
        if not 0 <= oylar <= 24:
            return _xato("Oylar soni 0..24 (0 — doimiy)")
        # O'tgan oydan boshlangan chegirma o'sha oylarning qarzini kamaytiradi —
        # bu qarzdorlikni tuzatish bilan bir xil, u esa FAQAT owner'da
        # (`HisobDetailView`). Admin chegirmani joriy oydan beradi.
        if oy < mantiq.oy_boshi(timezone.localdate()) and not owner_mi(request.user):
            return _xato("O'tgan oy uchun chegirmani faqat owner beradi — joriy yoki keyingi oyni tanlang", kod=403)
        with transaction.atomic():
            c = Chegirma.objects.create(azolik=am, narx=narx, boshlanish_oy=oy, oylar_soni=oylar,
                                        izoh=(request.data.get("izoh") or "").strip()[:300], kim=request.user)
            otkazildi = _chegirma_oylarini_qayta_hisobla(am, oy, oylar)
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.YARATISH, obyekt=am,
              obyekt_turi="CRM Chegirma", obyekt_nomi=_ism(am.azolik.talaba),
              snapshot={k: str(v) for k, v in _chegirma_dict(c).items()})
        return Response({**_chegirma_dict(c), "ogohlantirish": _qolda_ogohlantirish(otkazildi)}, status=201)


class ChegirmaDetailView(CrmView):
    pk_turi = "chegirma"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"

    def delete(self, request, pk):
        if "guruhlar.chegirma" not in ruxsatlar(request.user):
            return _xato("Chegirmani o'chirishga ruxsat yo'q", kod=403)
        c = get_object_or_404(Chegirma, pk=pk)
        am, oy, oylar = c.azolik, c.boshlanish_oy, c.oylar_soni
        snapshot = {k: str(v) for k, v in _chegirma_dict(c).items()}
        with transaction.atomic():
            c.delete()
            otkazildi = _chegirma_oylarini_qayta_hisobla(
                am, oy, oylar, dan=None if owner_mi(request.user) else mantiq.oy_boshi(timezone.localdate()),
            )
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.OCHIRISH, obyekt=am,
              obyekt_turi="CRM Chegirma", obyekt_nomi=_ism(am.azolik.talaba), snapshot=snapshot)
        ogohlantirish = _qolda_ogohlantirish(otkazildi)
        if ogohlantirish:
            return Response({"ogohlantirish": ogohlantirish})
        return Response(status=204)


# ── Bosh sahifa ko'rsatkichlari ──────────────────────────────────────


class KorsatkichlarView(CrmView):
    """SoffCRM bosh sahifasidagi 12 ta kartochka + "Markaz foydaliligi".

    Har kartochka o'z ruxsatiga bo'ysunadi — ruxsati yo'q ko'rsatkich
    `None` bo'lib keladi va frontendda ko'rinmaydi.
    """

    bolim = "bosh_sahifa"

    def get(self, request):
        r = ruxsatlar(request.user)
        filial = request.query_params.get("filial")
        bugun = timezone.localdate()
        oy = mantiq.oy_boshi(bugun)

        # Filial cheklovi (2026-09-23): filial xodimi — faqat o'z filiallari.
        u = request.user
        guruhlar = Guruh.objects.filter(faol=True).filter(guruh_q(u))
        azoliklar = AzolikMoliya.objects.filter(azolik__guruh__faol=True).filter(guruh_q(u, "azolik__guruh__"))
        lidlar = Lid.objects.filter(arxiv=False, qora_royxat=False).filter(filial_q(u, "filial"))
        hisoblar = Hisob.objects.filter(filial_q(u, "filial"))
        if filial:
            guruhlar = guruhlar.filter(moliya__filial_id=filial)
            azoliklar = azoliklar.filter(azolik__guruh__moliya__filial_id=filial)
            lidlar = lidlar.filter(Q(filial_id=filial) | Q(filial__isnull=True))
            hisoblar = hisoblar.filter(filial_id=filial)

        faol_talaba_idlar = set(
            azoliklar.filter(holat=AzolikMoliya.Holat.FAOL).values_list("azolik__talaba_id", flat=True)
        )
        qarz_qs = mantiq_qarzdorlar(hisoblar)
        tolov_yaqin = 0
        if "moliya" in r:
            # "To'lovi yaqin" — keyingi 3 kunda to'lov sanasi keladigan
            # faol talabalar (joriy oy hisobi hali to'lanmagan).
            # `mantiq.keyingi_tolov_sanasi` mantig'i, lekin har a'zolikka
            # alohida so'rov emas (N+1) — hamma hisoblar bitta so'rovda.
            chegara = bugun + timedelta(days=3)
            juftlar = set(
                azoliklar.filter(holat=AzolikMoliya.Holat.FAOL).values_list("azolik__talaba_id", "azolik__guruh_id")
            )
            tolanmagan, oxirgi = {}, {}
            for tid, gid, h_oy, h_holat in Hisob.objects.filter(
                talaba_id__in={t for t, _ in juftlar}, guruh_id__in={g for _, g in juftlar}
            ).values_list("talaba_id", "guruh_id", "oy", "holat"):
                k = (tid, gid)
                if k not in juftlar:
                    continue
                if h_holat != Hisob.Holat.TOLANDI:
                    tolanmagan[k] = min(tolanmagan.get(k, h_oy), h_oy)
                oxirgi[k] = max(oxirgi.get(k, h_oy), h_oy)
            for k in juftlar:
                sana = tolanmagan.get(k) or (mantiq.keyingi_oy(oxirgi[k]) if k in oxirgi else None)
                if sana and bugun <= sana <= chegara:
                    tolov_yaqin += 1

        # Bitirgan va boshqa guruhga o'tgan "ketgan" emas (markazda qolgan).
        oy_ketgan = GuruhdanChiqish.objects.filter(sana__gte=oy, sana__lte=bugun).filter(
            filial_q(u, "filial")).exclude(sabab_turi__in=GuruhdanChiqish.KETMAGAN_TURLAR)
        if filial:
            oy_ketgan = oy_ketgan.filter(filial_id=filial)

        oqituvchilar = User.objects.filter(role=User.Role.TEACHER, is_active=True)
        filiallari = ruxsat_filiallari(u)
        if filiallari is not None:
            # Xodimlar ro'yxatidagi qoida: o'z filiali yoki filialsiz o'qituvchi.
            oqituvchilar = oqituvchilar.filter(
                Q(crm_xodim__isnull=True) | Q(crm_xodim__filiallar__isnull=True)
                | Q(crm_xodim__filiallar__in=filiallari)
            ).distinct()

        # Markaz foydaliligi = (shu oy kassaga tushgan pul − xodimlar
        # oyligi − o'qituvchi ulushlari) / tushum. Ulush guruh tushumidan
        # hisoblanadi (guruhdagi o'qituvchi foizi, bo'lmasa profildagi).
        tushum = foydalilik = None
        if "bosh_sahifa.markaz_foydaliligi" in r:
            tolovlar = Tolov.objects.filter(turi=Tolov.Turi.TOLOV, sana__gte=oy, sana__lte=bugun).filter(tolov_q(u))
            qaytarishlar = Tolov.objects.filter(turi=Tolov.Turi.QAYTARISH, sana__gte=oy, sana__lte=bugun).filter(
                tolov_q(u))
            if filial:
                tolovlar = tolovlar.filter(tolov_filiali_q(filial))
                qaytarishlar = qaytarishlar.filter(tolov_filiali_q(filial))
            tushum = (tolovlar.aggregate(s=Sum("summa"))["s"] or NOL) - (
                qaytarishlar.aggregate(s=Sum("summa"))["s"] or NOL
            )
            xarajat = NOL
            profillar = XodimProfil.objects.filter(user__is_active=True)
            # M2M bo'yicha filtr — subquery orqali: to'g'ridan-to'g'ri join
            # ikki filialli xodimning oyligini ikki marta qo'shardi.
            if filiallari is not None:
                profillar = profillar.filter(pk__in=XodimProfil.objects.filter(filiallar__in=filiallari).values("pk"))
            if filial:
                profillar = profillar.filter(pk__in=XodimProfil.objects.filter(filiallar=filial).values("pk"))
            xarajat += profillar.aggregate(s=Sum("oylik"))["s"] or NOL
            guruh_tushumi = {
                x["guruh_id"]: x["s"] for x in tolovlar.values("guruh_id").annotate(s=Sum("summa"))
            }
            umumiy_foiz = dict(XodimProfil.objects.values_list("user_id", "foiz_ulushi"))
            ulushlar = GuruhOqituvchi.objects.filter(guruh__in=guruhlar).select_related("guruh")
            for go in ulushlar:
                if go.ulush_turi == GuruhOqituvchi.UlushTuri.DARS:
                    # O'tilgan darslar (bugungacha) × bitta dars haqi.
                    otilgan = sum(1 for s in guruh_dars_sanalari(go.guruh, oy) if s <= bugun)
                    xarajat += (go.dars_haqi or NOL) * otilgan
                elif go.guruh_id in guruh_tushumi:
                    foiz = go.foiz if go.foiz is not None else umumiy_foiz.get(go.oqituvchi_id, NOL)
                    xarajat += guruh_tushumi[go.guruh_id] * (foiz or NOL) / Decimal(100)
            foydalilik = round(float((tushum - xarajat) / tushum * 100), 1) if tushum > 0 else 0.0

        def ruxsat(kalit, qiymat):
            return qiymat if kalit in r else None

        return Response({
            "faol_lidlar": ruxsat("bosh_sahifa.faol_lidlar", lidlar.count()),
            "guruhlar": guruhlar.count(),
            "qolgan_qarz": ruxsat("bosh_sahifa.qarzdorlar", qarz_qs["summa"]),
            "qarzdorlar": ruxsat("bosh_sahifa.qarzdorlar", qarz_qs["soni"]),
            "tolovi_yaqin": ruxsat("moliya", tolov_yaqin),
            "faol_talabalar": ruxsat("bosh_sahifa.faol_talabalar", len(faol_talaba_idlar)),
            "jami_guruhdagi": azoliklar.values("azolik__talaba_id").distinct().count(),
            "sinov_darsida": azoliklar.filter(holat=AzolikMoliya.Holat.SINOV).count(),
            "ketganlar": oy_ketgan.count(),
            "oqituvchilar": oqituvchilar.count(),
            "yangi_lidlar_bugun": ruxsat("bosh_sahifa.faol_lidlar", lidlar.filter(created_at__date=bugun).count()),
            # SoffCRM "Yangi guruhga qabul" — yig'ilayotgan guruhlarga navbatdagi lidlar.
            "yangi_guruhga_qabul": ruxsat(
                "bosh_sahifa.faol_lidlar", lidlar.filter(yigilayotgan_guruh__isnull=False).count()
            ),
            "muzlatilgan": azoliklar.filter(holat=AzolikMoliya.Holat.MUZLATILGAN).count(),
            "tushum": tushum,
            "markaz_foydaliligi": foydalilik,
        })


def mantiq_qarzdorlar(hisoblar):
    """To'lanmagan hisoblar bo'yicha qarz summasi va qarzdor talabalar soni."""
    ochiq = hisoblar.exclude(holat=Hisob.Holat.TOLANDI).annotate(
        tolangan=Sum("tolovlar__summa", filter=Q(tolovlar__turi__in=Tolov.YOPUVCHI_TURLAR), default=NOL)
    )
    summa = NOL
    talabalar = set()
    for h in ochiq:
        qoldiq = h.summa - h.tolangan
        if qoldiq > 0:
            summa += qoldiq
            talabalar.add(h.talaba_id)
    return {"summa": summa, "soni": len(talabalar)}


# ── Talaba qidiruvi (guruhga qo'shish oynasi uchun) ──────────────────


class TalabaQidiruvView(CrmView):
    """Barcha talabalar (guruhsizlari ham) — ism/telefon/login bo'yicha."""

    bolim = None

    def get(self, request):
        q = (request.query_params.get("q") or "").strip()
        qs = User.objects.filter(role=User.Role.STUDENT, is_active=True)
        if cheklanganmi(request.user):
            # Filial xodimi: o'z filiali guruhlaridagi yoki hech qaysi
            # guruhda bo'lmagan talabalar (`crm.filial.talaba_korinadimi`).
            qs = qs.filter(
                Q(pk__in=GuruhAzoligi.objects.filter(guruh_q(request.user, "guruh__")).values("talaba_id"))
                | ~Q(pk__in=GuruhAzoligi.objects.values("talaba_id"))
            )
        if q:
            qs = qs.filter(Q(first_name__icontains=q) | Q(username__icontains=q) | Q(telefon__icontains=q))
        # `?faqat_crm=1` — CRM'da yaratilgan yoki guruhga CRM orqali
        # yozilgan talabalar (saytdagi "biriktirish" tanlovi uchun).
        if request.query_params.get("faqat_crm"):
            qs = qs.filter(Q(crm_talaba__isnull=False) | Q(guruhazoligi__moliya__isnull=False)).distinct()
        return Response([
            {
                "id": u.id, "ism": _ism(u), "telefon": u.telefon, "username": u.username,
                # Saytga hech kirmagan — CRM avtomatik bergan login hali
                # ishlatilmagan, uni almashtirish xavfsiz.
                "saytga_kirgan": u.last_login is not None,
                "guruhlar": [g.name for g in u.talaba_guruhlari.all() if g.faol],
            }
            # Boshqa filial guruhlarining nomlari ham chiqmasin.
            for u in qs.prefetch_related(
                Prefetch("talaba_guruhlari", queryset=Guruh.objects.filter(guruh_q(request.user)))
            ).order_by("first_name", "username")[:30]
        ])


class TalabaSaytHisobiView(CrmView):
    """Saytdagi "Yangi foydalanuvchi" formasi (rol = talaba) uchun: yangi
    hisob OCHILMAYDI — kiritilgan login/parol/ism CRM'dagi mavjud
    talabaga yoziladi (biriktiriladi). Aks holda bitta o'quvchining
    ikkita hisobi bo'lib qolardi: CRM'dagisi (guruh, to'lov) va saytdagisi
    (mashqlar) — va ular hech qachon uchrashmasdi.

    Talabaning pul tarixi, guruhlari va natijalari o'zgarmaydi — faqat
    kirish ma'lumoti."""

    pk_turi = "talaba"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "talabalar"

    def post(self, request, pk):
        from accounts.views import _parolni_tekshir

        if xato := _ruxsatsiz(request, "talabalar.tahrirlash"):
            return xato
        talaba = get_object_or_404(User, pk=pk, role=User.Role.STUDENT)
        username = (request.data.get("username") or "").strip()
        parol = request.data.get("parol") or ""
        ism = (request.data.get("ism") or "").strip()
        if not username or not parol:
            return _xato("Login va parol majburiy")
        if not LOGIN_QOIDASI.match(username):
            return _xato("Login faqat harf/raqam/./@/+/-/_ dan iborat bo'lsin")
        if User.objects.filter(username=username).exclude(pk=talaba.pk).exists():
            return _xato("Bu login band")
        xatolar = _parolni_tekshir(parol, user=talaba)
        if xatolar:
            return _xato(" ".join(xatolar))
        eski_login = talaba.username
        talaba.username = username
        talaba.set_password(parol)
        maydonlar = ["username", "password"]
        if ism:
            talaba.first_name = ism[:150]
            talaba.last_name = ""
            maydonlar += ["first_name", "last_name"]
        talaba.save(update_fields=maydonlar)
        logla(foydalanuvchi=request.user, harakat=FaoliyatYozuvi.Harakat.OZGARTIRISH, obyekt=talaba,
              obyekt_turi="Talaba", obyekt_nomi=_ism(talaba),
              ozgarishlar={"username": {"eski": eski_login, "yangi": username},
                           "parol": {"eski": "***", "yangi": "yangilandi"}})
        return Response({"id": talaba.id, "username": talaba.username, "ism": _ism(talaba), "role": talaba.role})


# ── Baholar (SoffCRM "BAHO" tabi) ────────────────────────────────────


class GuruhBaholariView(CrmView):
    """Darslar bo'yicha baholar jadvali. Shkala — guruh sozlamasidagi
    `baholash_tizimi` (1-5 / 1-10 / 100); u bo'sh bo'lsa baho qo'yilmaydi."""

    pk_turi = "guruh"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"

    def get(self, request, pk):
        guruh = get_object_or_404(Guruh, pk=pk)
        try:
            oy = _oy(request.query_params.get("oy") or timezone.localdate().strftime("%Y-%m"))
        except ValueError as e:
            return _xato(str(e))
        baholar = list(DarsBahosi.objects.filter(guruh=guruh, sana__gte=oy, sana__lte=mantiq.oy_oxiri(oy)))
        sanalar = sorted(guruh_dars_sanalari(guruh, oy) | {b.sana for b in baholar})
        katak = {(b.talaba_id, b.sana): b for b in baholar}
        moliya = getattr(guruh, "moliya", None)
        talabalar = []
        for a in sorted(GuruhAzoligi.objects.filter(guruh=guruh).select_related("talaba"),
                        key=lambda a: (_ism(a.talaba) or "").lower()):
            qiymatlar = [katak.get((a.talaba_id, s)) for s in sanalar]
            bor = [b.ball for b in qiymatlar if b]
            talabalar.append({
                "id": a.talaba_id,
                "ism": _ism(a.talaba),
                "baholar": [b.ball if b else None for b in qiymatlar],
                "ortacha": round(sum(bor) / len(bor), 1) if bor else None,
            })
        return Response({
            "oy": oy, "sanalar": sanalar, "talabalar": talabalar,
            "tizim": moliya.baholash_tizimi if moliya else "",
        })

    def post(self, request, pk):
        if "guruhlar.davomat" not in ruxsatlar(request.user):
            return _xato("Baho qo'yishga ruxsat yo'q", kod=403)
        guruh = get_object_or_404(Guruh, pk=pk)
        moliya = getattr(guruh, "moliya", None)
        tizim = moliya.baholash_tizimi if moliya else ""
        if not tizim:
            return _xato("Guruhda baholash tizimi tanlanmagan")
        talaba = get_object_or_404(User, pk=request.data.get("talaba_id"))
        if not GuruhAzoligi.objects.filter(guruh=guruh, talaba=talaba).exists():
            return _xato("Talaba bu guruhda emas")
        try:
            sana = _sana(request.data.get("sana"), "sana")
        except ValueError as e:
            return _xato(str(e))
        xom = request.data.get("ball")
        if xom in (None, ""):
            DarsBahosi.objects.filter(guruh=guruh, talaba=talaba, sana=sana).delete()
            return Response({"ball": None})
        try:
            ball = _son(xom, "ball")
        except ValueError as e:
            return _xato(str(e))
        eng_kichik = Decimal(0) if tizim == "100" else Decimal(1)
        if not eng_kichik <= ball <= Decimal(tizim):
            return _xato(f"Baho {eng_kichik}..{tizim} oralig'ida bo'lsin")
        DarsBahosi.objects.update_or_create(
            guruh=guruh, talaba=talaba, sana=sana,
            defaults={"ball": ball, "izoh": (request.data.get("izoh") or "").strip()[:300], "kim": request.user},
        )
        return Response({"ball": ball})


class MuddatliEslatmalarView(CrmView):
    """Vaqti kelgan eslatmalar (video 09:05-09:26: "28-sanada soat 2 da
    eslatsin") — bosh sahifada bugungi va muddati o'tganlari chiqadi.
    Faqat shu foydalanuvchi YOZGANLARI: har kim o'z eslatmalarini ko'radi."""

    bolim = None

    def get(self, request):
        oxiri = timezone.localtime().replace(hour=23, minute=59, second=59)
        qs = (
            Eslatma.objects.filter(kim=request.user, eslatish_vaqti__isnull=False, eslatish_vaqti__lte=oxiri)
            .select_related("lid", "talaba", "guruh")
            .order_by("eslatish_vaqti")[:50]
        )
        return Response([
            {
                "id": e.id, "matn": e.matn, "eslatish_vaqti": e.eslatish_vaqti,
                "lid_id": e.lid_id, "lid": e.lid.ism if e.lid_id else None,
                "talaba_id": e.talaba_id, "talaba": _ism(e.talaba) if e.talaba_id else None,
                "guruh_id": e.guruh_id, "guruh": e.guruh.name if e.guruh_id else None,
                "otgan": e.eslatish_vaqti < timezone.now(),
            }
            for e in qs
        ])
