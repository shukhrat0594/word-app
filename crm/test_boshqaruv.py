"""Video-TZ (2026-09-23) funksiyalari testlari: ruxsatlar, xodimlar,
lidlar, talaba/guruh yaratish, chegirma, dars ko'chirish."""

from datetime import date
from decimal import Decimal

from academics.models import Guruh, GuruhAzoligi
from accounts.models import User
from crm import mantiq
from crm.models import (
    AzolikMoliya, CrmRol, DarsJadvali, Hisob, Lid, LidBolim, TalabaProfil, Xona, XodimProfil,
)
from crm.test_api import ApiAsos
from crm.tests import NARX, SENTABR, bugun_qilib


class RollarRuxsatTest(ApiAsos):
    def xodim(self, lavozim, rol=None):
        u = User.objects.create_user(username=f"x_{lavozim}", password="x", role=User.Role.ODDIY)
        XodimProfil.objects.create(user=u, lavozim=lavozim, rol=rol)
        return u

    def test_kassir_moliyaga_kiradi_lidlarga_yoq(self):
        m = self.mijoz(self.xodim("kassir"))
        self.assertEqual(m.get("/api/crm/hisoblar/").status_code, 200)
        self.assertEqual(m.get("/api/crm/lidlar/").status_code, 403)
        self.assertEqual(m.get("/api/crm/xodimlar/").status_code, 200)  # o'qish ochiq
        self.assertEqual(m.post("/api/crm/xodimlar/", {"ism": "A"}, format="json").status_code, 403)

    def test_marketolog_faqat_lidlar(self):
        m = self.mijoz(self.xodim("marketolog"))
        self.assertEqual(m.get("/api/crm/lidlar/").status_code, 200)
        self.assertEqual(m.get("/api/crm/hisoblar/").status_code, 403)

    def test_maxsus_rol_ruxsatlari(self):
        rol = CrmRol.objects.create(nomi="Administrator 1", ruxsatlar=["hisobotlar.moliya", "lidlar"])
        m = self.mijoz(self.xodim("boshqa", rol))
        self.assertEqual(m.get("/api/crm/hisobot/").status_code, 200)
        self.assertEqual(m.get("/api/crm/lidlar/").status_code, 200)
        self.assertEqual(m.get("/api/crm/guruhlar/").status_code, 403)
        men = m.get("/api/crm/men/").data["ruxsatlar"]
        self.assertIn("lidlar.excel", men)  # bo'lim berilsa — amallari ham
        self.assertNotIn("hisobotlar.excel", men)  # faqat bitta amal berilgan

    def test_oqituvchi_crmga_kirmaydi(self):
        self.assertEqual(self.mijoz(self.xodim("oqituvchi")).get("/api/crm/men/").status_code, 403)


class XodimTest(ApiAsos):
    def test_oqituvchi_yaratish_avtomatik_login_parol(self):
        javob = self.mijoz(self.admin).post("/api/crm/xodimlar/", {
            "ism": "Isroil Zohidjonov", "telefon": "+998 90 123 45 67", "lavozim": "oqituvchi",
            "oylik": "3000000", "foiz_ulushi": "25", "filial_id": self.filial.id,
        }, format="json")
        self.assertEqual(javob.status_code, 201, javob.data)
        u = User.objects.get(pk=javob.data["id"])
        self.assertEqual(u.role, User.Role.TEACHER)
        self.assertEqual(u.username, "u901234567")
        self.assertTrue(u.check_password(javob.data["parol"]))
        self.assertEqual(u.crm_xodim.foiz_ulushi, Decimal("25"))

    def test_admin_administrator_qosha_olmaydi(self):
        javob = self.mijoz(self.admin).post("/api/crm/xodimlar/", {
            "ism": "B", "telefon": "901111111", "lavozim": "admin"}, format="json")
        self.assertEqual(javob.status_code, 403)

    def test_kassir_lms_roli_oddiy(self):
        javob = self.mijoz(self.owner).post("/api/crm/xodimlar/", {
            "ism": "Kassir", "telefon": "901111112", "lavozim": "kassir"}, format="json")
        self.assertEqual(User.objects.get(pk=javob.data["id"]).role, User.Role.ODDIY)

    def test_ruxsatsiz_oylik_yashirin(self):
        rol = CrmRol.objects.create(nomi="Ko'ruvchi", ruxsatlar=["xodimlar"])
        rol.ruxsatlar = ["xodimlar.qoshish"]
        rol.save()
        u = User.objects.create_user(username="k", password="x", role=User.Role.ODDIY)
        XodimProfil.objects.create(user=u, lavozim="boshqa", rol=rol)
        xodimlar = self.mijoz(u).get("/api/crm/xodimlar/").data["xodimlar"]
        self.assertTrue(all(x["oylik"] is None for x in xodimlar))


class TalabaVaLidTest(ApiAsos):
    def test_talaba_yaratib_guruhga_qoshish(self):
        javob = self.mijoz(self.admin).post("/api/crm/talaba-yaratish/", {
            "ism": "Ali Valiyev", "telefon": "+998901234568", "ota_ona_telefon": "901234569",
            "guruh_id": self.guruh.id, "boshlanish_sana": "2026-09-10", "holat": "faol", "jins": "erkak",
        }, format="json")
        self.assertEqual(javob.status_code, 201, javob.data)
        u = User.objects.get(pk=javob.data["id"])
        self.assertEqual(u.role, User.Role.STUDENT)
        self.assertTrue(u.check_password(javob.data["parol"]))
        am = AzolikMoliya.objects.get(azolik__talaba=u)
        self.assertEqual(am.boshlanish_sana, date(2026, 9, 10))
        self.assertEqual(u.crm_talaba.jins, "erkak")
        # Guruhsiz talabalar ham ro'yxatda ko'rinadi
        ismlar = [x["ism"] for x in self.mijoz(self.admin).get("/api/crm/talabalar/").data]
        self.assertIn("Ali Valiyev", ismlar)

    def test_lid_guruhga_talabaga_aylanadi(self):
        bolim = LidBolim.objects.create(nomi="Beginner")
        m = self.mijoz(self.admin)
        lid = m.post("/api/crm/lidlar/", {"ism": "Laura", "telefon": "+998 94 991 49 60",
                                          "bolim_id": bolim.id, "eslatma": "23.09 kuni keladi"},
                     format="json").data
        self.assertEqual(lid["telefon"], "+998949914960")
        javob = m.post("/api/crm/lidlar/guruhga/", {"lid_idlar": [lid["id"]], "guruh_id": self.guruh.id},
                       format="json")
        self.assertEqual(javob.status_code, 201, javob.data)
        l = Lid.objects.get(pk=lid["id"])
        self.assertTrue(l.arxiv)
        self.assertEqual(l.holat, Lid.Holat.QOSHILDI)
        self.assertTrue(GuruhAzoligi.objects.filter(guruh=self.guruh, talaba=l.talaba).exists())
        self.assertEqual(AzolikMoliya.objects.get(azolik__talaba=l.talaba).holat, "sinov")
        # Kanbanda endi ko'rinmaydi, arxivda bor
        self.assertEqual(m.get("/api/crm/lidlar/").data, [])
        self.assertEqual(len(m.get("/api/crm/lidlar/?arxiv=1").data), 1)

    def test_qora_royxatdagi_raqam_409(self):
        m = self.mijoz(self.admin)
        TalabaProfil.objects.create(user=self.talaba, qora_royxat=True)
        User.objects.filter(pk=self.talaba.pk).update(telefon="+998900000001")
        javob = m.post("/api/crm/lidlar/", {"ism": "X", "telefon": "+998900000001"}, format="json")
        self.assertEqual(javob.status_code, 409)

    def test_guruhdan_chiqarish_pul_tarixi_qoladi(self):
        self.azolik_qosh(boshlanish=SENTABR)
        with bugun_qilib(date(2026, 9, 5)):
            mantiq.hisoblarni_generatsiya_qil()
        self.assertTrue(Hisob.objects.filter(talaba=self.talaba).exists())
        with bugun_qilib(date(2026, 9, 12)):
            javob = self.mijoz(self.admin).delete(
                f"/api/crm/guruhlar/{self.guruh.id}/talabalar/?talaba={self.talaba.id}&sana=2026-09-12"
            )
        self.assertEqual(javob.status_code, 204)
        self.assertFalse(GuruhAzoligi.objects.filter(guruh=self.guruh, talaba=self.talaba).exists())
        hisob = Hisob.objects.get(talaba=self.talaba, oy=SENTABR)
        self.assertLess(hisob.summa, NARX)  # 12-sentabrgacha proporsional


class GuruhTest(ApiAsos):
    def test_guruh_yaratish_jadval_va_oqituvchilar(self):
        xona = Xona.objects.create(filial=self.filial, nomi="1-xona")
        javob = self.mijoz(self.admin).post("/api/crm/guruh-yaratish/", {
            "nomi": "Test guruhi", "daraja_id": self.daraja.id, "filial_id": self.filial.id,
            "boshlanish_sana": "2026-09-23",
            "jadval": [{"hafta_kuni": 0, "boshlanish_vaqti": "18:30", "tugash_vaqti": "20:00", "xona_id": xona.id}],
            "oqituvchilar": [{"oqituvchi_id": self.oqituvchi.id, "turi": "asosiy", "foiz": "25"}],
        }, format="json")
        self.assertEqual(javob.status_code, 201, javob.data)
        g = Guruh.objects.get(pk=javob.data["id"])
        self.assertEqual(g.oqituvchi_id, self.oqituvchi.id)
        self.assertEqual(g.fan_id, self.daraja.parent_id)
        self.assertEqual(g.crm_jadval.count(), 1)

    def test_xona_toqnashsa_guruh_yaratilmaydi(self):
        xona = Xona.objects.create(filial=self.filial, nomi="2-xona")
        DarsJadvali.objects.filter(guruh=self.guruh).update(xona=xona)
        soni = Guruh.objects.count()
        javob = self.mijoz(self.admin).post("/api/crm/guruh-yaratish/", {
            "nomi": "To'qnash", "boshlanish_sana": "2026-09-23",
            "jadval": [{"hafta_kuni": 3, "boshlanish_vaqti": "14:30", "tugash_vaqti": "16:00", "xona_id": xona.id}],
        }, format="json")
        self.assertEqual(javob.status_code, 400)
        self.assertEqual(Guruh.objects.count(), soni)  # tranzaksiya qaytarildi

    def test_chegirma_hisob_summasini_kamaytiradi(self):
        am = self.azolik_qosh(boshlanish=SENTABR)
        with bugun_qilib(date(2026, 9, 5)):
            mantiq.hisoblarni_generatsiya_qil()
        self.assertEqual(Hisob.objects.get(talaba=self.talaba, oy=SENTABR).summa, NARX)
        javob = self.mijoz(self.admin).post(f"/api/crm/guruhlar/{self.guruh.id}/chegirmalar/", {
            "azolik_moliya_id": am.id, "narx": "400000", "boshlanish_oy": "2026-09", "oylar_soni": 2,
        }, format="json")
        self.assertEqual(javob.status_code, 201, javob.data)
        self.assertEqual(Hisob.objects.get(talaba=self.talaba, oy=SENTABR).summa, Decimal("400000"))
        self.assertEqual(mantiq.amaldagi_narx(am, date(2026, 10, 1)), Decimal("400000"))
        self.assertEqual(mantiq.amaldagi_narx(am, date(2026, 11, 1)), NARX)  # muddati tugadi
        # O'chirilsa — qaytadi
        self.mijoz(self.admin).delete(f"/api/crm/chegirmalar/{javob.data['id']}/")
        self.assertEqual(Hisob.objects.get(talaba=self.talaba, oy=SENTABR).summa, NARX)

    def test_dars_kochirish_davomat_ustunlarini_ozgartiradi(self):
        m = self.mijoz(self.admin)
        javob = m.post(f"/api/crm/guruhlar/{self.guruh.id}/dars-ozgarishlari/", {
            "turi": "kochirish", "asl_sana": "2026-09-10", "yangi_sana": "2026-09-14"}, format="json")
        self.assertEqual(javob.status_code, 201, javob.data)
        sanalar = m.get(f"/api/crm/guruhlar/{self.guruh.id}/davomat/?oy=2026-09").data["sanalar"]
        self.assertNotIn(date(2026, 9, 10), sanalar)
        self.assertIn(date(2026, 9, 14), sanalar)
        # Darsi yo'q kunni ko'chirib bo'lmaydi
        xato = m.post(f"/api/crm/guruhlar/{self.guruh.id}/dars-ozgarishlari/", {
            "turi": "bekor", "asl_sana": "2026-09-07"}, format="json")
        self.assertEqual(xato.status_code, 400)

    def test_korsatkichlar(self):
        self.azolik_qosh(boshlanish=SENTABR)
        javob = self.mijoz(self.admin).get("/api/crm/korsatkichlar/")
        self.assertEqual(javob.status_code, 200, javob.data)
        self.assertEqual(javob.data["faol_talabalar"], 1)
        self.assertEqual(javob.data["guruhlar"], 1)


class BahoVaUlushTest(ApiAsos):
    def test_baho_shkala_boyicha(self):
        self.azolik_qosh(boshlanish=SENTABR)
        yol = f"/api/crm/guruhlar/{self.guruh.id}/baholar/"
        m = self.mijoz(self.admin)
        # Tizim tanlanmagan — baho qo'yilmaydi
        self.assertEqual(m.post(yol, {"talaba_id": self.talaba.id, "sana": "2026-09-03", "ball": 4},
                                format="json").status_code, 400)
        self.guruh.moliya.baholash_tizimi = "5"
        self.guruh.moliya.save()
        self.assertEqual(m.post(yol, {"talaba_id": self.talaba.id, "sana": "2026-09-03", "ball": 7},
                                format="json").status_code, 400)
        self.assertEqual(m.post(yol, {"talaba_id": self.talaba.id, "sana": "2026-09-03", "ball": 4},
                                format="json").status_code, 200)
        qator = m.get(yol + "?oy=2026-09").data["talabalar"][0]
        self.assertEqual(qator["ortacha"], 4)

    def test_dars_uchun_ulush_foydalilikka_kiradi(self):
        from crm.models import GuruhOqituvchi
        GuruhOqituvchi.objects.create(guruh=self.guruh, oqituvchi=self.oqituvchi,
                                      ulush_turi="dars", dars_haqi=Decimal("50000"))
        self.azolik_qosh(boshlanish=SENTABR)
        from crm.models import Tolov
        from django.utils import timezone
        Tolov.objects.create(talaba=self.talaba, talaba_ism="t", guruh=self.guruh, guruh_nomi="g",
                             sana=timezone.localdate(), summa=Decimal("1000000"), turi="tolov")
        javob = self.mijoz(self.admin).get("/api/crm/korsatkichlar/")
        self.assertEqual(javob.data["tushum"], Decimal("1000000"))
        self.assertLess(javob.data["markaz_foydaliligi"], 100)

    def test_ulush_dars_summasiz_rad(self):
        javob = self.mijoz(self.admin).post("/api/crm/guruh-yaratish/", {
            "nomi": "X", "boshlanish_sana": "2026-09-23", "jadval": [],
            "oqituvchilar": [{"oqituvchi_id": self.oqituvchi.id, "ulush_turi": "dars"}],
        }, format="json")
        self.assertEqual(javob.status_code, 400)


class LidEksportTest(ApiAsos):
    def test_excel_filtr_bilan(self):
        from io import BytesIO

        from openpyxl import load_workbook

        m = self.mijoz(self.admin)
        m.post("/api/crm/lidlar/", {"ism": "Laura", "telefon": "901112233", "manba": "Instagram"}, format="json")
        m.post("/api/crm/lidlar/", {"ism": "Diana", "telefon": "901112244", "manba": "Telegram"}, format="json")
        javob = m.get("/api/crm/lidlar/eksport/?manba=Instagram")
        self.assertEqual(javob.status_code, 200)
        varaq = load_workbook(BytesIO(javob.content)).active
        ismlar = [r[1] for r in varaq.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(ismlar, ["Laura"])

    def test_ruxsatsiz_403(self):
        rol = CrmRol.objects.create(nomi="Faqat lid", ruxsatlar=["lidlar.qoshish"])
        u = User.objects.create_user(username="m", password="x", role=User.Role.ODDIY)
        XodimProfil.objects.create(user=u, lavozim="boshqa", rol=rol)
        self.assertEqual(self.mijoz(u).get("/api/crm/lidlar/eksport/").status_code, 403)
