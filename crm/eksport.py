"""Excel eksport (openpyxl).

Loyihada openpyxl bor, lekin faqat IMPORT uchun ishlatilgan
(`accounts/excel_import.py`, `courses/excel_import.py`) — eksport noldan
yozildi.

Bu — CRM ma'lumotining CHIQISH ESHIGI: CRM bir kun o'chirilsa, moliya
tarixi shu fayl orqali saqlanib qoladi (TZ 8-band, 1-qadam). Shuning
uchun u 1-bosqichda qilinadi, keyinga qoldirilmaydi.
"""

from io import BytesIO
from urllib.parse import quote

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SARLAVHA_FON = PatternFill("solid", fgColor="4F46E5")
SARLAVHA_SHRIFT = Font(bold=True, color="FFFFFF")
PUL_FORMAT = "#,##0"


def _varaq_yoz(ws, sarlavhalar, qatorlar, pul_ustunlari=()):
    ws.append(sarlavhalar)
    for katak in ws[1]:
        katak.fill = SARLAVHA_FON
        katak.font = SARLAVHA_SHRIFT
        katak.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for qator in qatorlar:
        ws.append(qator)

    for raqam, nomi in enumerate(sarlavhalar, start=1):
        harf = get_column_letter(raqam)
        # Ustun kengligi — sarlavha va qiymatlarning eng uzuni bo'yicha,
        # 50 belgidan oshmasin (aks holda izoh ustuni ekranni egallaydi).
        eng_uzun = max(
            [len(str(nomi))] + [len(str(q[raqam - 1])) for q in qatorlar if raqam <= len(q)]
        )
        ws.column_dimensions[harf].width = min(max(eng_uzun + 2, 10), 50)
        if raqam in pul_ustunlari:
            for katak in ws[harf][1:]:
                katak.number_format = PUL_FORMAT


def javob_qil(kitob: Workbook, fayl_nomi: str) -> HttpResponse:
    oqim = BytesIO()
    kitob.save(oqim)
    oqim.seek(0)
    javob = HttpResponse(
        oqim.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    # RFC 5987 — o'zbekcha fayl nomi pastki chiziqqa aylanib ketmasligi
    # uchun (frontend `apiFayluniYuklab` aynan shu shaklni o'qiydi).
    javob["Content-Disposition"] = (
        f"attachment; filename=\"crm.xlsx\"; filename*=UTF-8''{quote(fayl_nomi)}"
    )
    return javob


def qarzdorlar_kitobi(hisoblar, hisobot_qatorlari, tolovlar):
    """Uch varaqli kitob: qarzdorlar, to'lovlar, hisobot."""
    kitob = Workbook()

    ws = kitob.active
    ws.title = "Qarzdorlar"
    _varaq_yoz(
        ws,
        ["Talaba", "Guruh", "Filial", "Oy", "Hisoblangan", "To'langan", "Qoldiq", "Holat"],
        [
            [
                h["talaba"], h["guruh"], h["filial"] or "—", f"{h['oy']:%Y-%m}",
                float(h["summa"]), float(h["tolangan"]), float(h["qoldiq"]), h["holat"],
            ]
            for h in hisoblar
        ],
        pul_ustunlari=(5, 6, 7),
    )

    ws = kitob.create_sheet("To'lovlar")
    _varaq_yoz(
        ws,
        ["ID", "Sana", "Qaysi oy uchun", "Turi", "Summa", "Talaba", "Guruh", "Izoh", "Kim"],
        [
            [
                str(t["id"]), f"{t['sana']:%Y-%m-%d}",
                f"{t['oy']:%Y-%m}" if t.get("oy") else "—",
                t["turi_nomi"], float(t["summa"]), t["talaba"], t["guruh"],
                t["izoh"], t["kim"] or "—",
            ]
            for t in tolovlar
        ],
        pul_ustunlari=(5,),
    )

    ws = kitob.create_sheet("Hisobot")
    _varaq_yoz(
        ws,
        ["Filial", "Guruh", "Talaba", "Hisoblangan", "Olingan pul", "Chegirma",
         "Bonus", "Qaytarilgan", "Qarz", "Yig'ilish %"],
        [
            [
                q["filial"], q["guruh"], q["talaba_soni"],
                float(q["hisoblangan"]), float(q["olingan"]), float(q["chegirma"]),
                float(q["bonus"]), float(q["qaytarilgan"]), float(q["qarz"]),
                q["yigilish_foizi"],
            ]
            for q in hisobot_qatorlari
        ],
        pul_ustunlari=(4, 5, 6, 7, 8, 9),
    )

    return kitob
