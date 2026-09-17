"""CRM API testlari — ruxsatlar va asosiy oqimlar.

Ruxsat tekshiruvi HAQIQIY JWT bilan o'tkaziladi (`force_authenticate`
emas): aks holda `accounts.authentication.KorishRejimliJWTAuthentication`
qatlami chetlab o'tilardi va owner "Ko'rish rejimi"dagi holat umuman
sinovdan o'tmasdi.
"""

from datetime import date
from decimal import Decimal

from django.test import override_settings
from rest_framework.test import APIClient
from rest_framework_simplejwt.tokens import RefreshToken

from accounts.models import User
from assessment.models import WritingTekshiruv
from crm import mantiq
from academics.models import Davomat, Guruh
from crm.models import (
    AzolikMoliya, Filial, GuruhMoliya, Hisob, KursNarxi, Sozlama, Tolov, Xona,
)
from crm.tests import AVGUST, NARX, SENTABR, CrmAsos, bugun_qilib


class ApiAsos(CrmAsos):
    def setUp(self):
        super().setUp()
        self.admin = User.objects.create_user(
            username="crm_admin", password="x", role=User.Role.ADMIN, markaz=self.markaz
        )
        self.owner = User.objects.create_user(
            username="crm_owner", password="x", role=User.Role.ADMIN,
            markaz=self.markaz, is_superuser=True, is_staff=True,
        )
        self.oqituvchi = User.objects.create_user(
            username="crm_oqituvchi", password="x", role=User.Role.TEACHER, markaz=self.markaz
        )
        self.ota_ona = User.objects.create_user(
            username="crm_otaona", password="x", role=User.Role.PARENT, markaz=self.markaz
        )

    def mijoz(self, user=None):
        client = APIClient()
        if user is not None:
            token = str(RefreshToken.for_user(user).access_token)
            client.credentials(HTTP_AUTHORIZATION=f"Bearer {token}")
        return client


class RuxsatTest(ApiAsos):
    YOLLAR = [
        "/api/crm/hisoblar/",
        "/api/crm/guruhlar/",
        "/api/crm/filiallar/",
        "/api/crm/hisobot/",
        "/api/crm/kurs-narxlari/",
    ]

    def test_autentifikatsiyasiz_401(self):
        for yol in self.YOLLAR:
            self.assertEqual(self.mijoz().get(yol).status_code, 401, yol)

    def test_talaba_403(self):
        """CRM manzilini bilib olgan talaba API'ga to'g'ridan-to'g'ri
        kira olmasligi kerak — frontendda yashirish himoya emas."""
        mijoz = self.mijoz(self.talaba)
        for yol in self.YOLLAR:
            self.assertEqual(mijoz.get(yol).status_code, 403, yol)

    def test_oqituvchi_403(self):
        mijoz = self.mijoz(self.oqituvchi)
        for yol in self.YOLLAR:
            self.assertEqual(mijoz.get(yol).status_code, 403, yol)

    def test_ota_ona_403(self):
        self.assertEqual(self.mijoz(self.ota_ona).get("/api/crm/hisoblar/").status_code, 403)

    def test_admin_kiradi(self):
        self.assertEqual(self.mijoz(self.admin).get("/api/crm/hisoblar/").status_code, 200)

    def test_owner_kiradi(self):
        self.assertEqual(self.mijoz(self.owner).get("/api/crm/hisoblar/").status_code, 200)

    def test_korish_rejimidagi_owner_403(self):
        """Owner "Ko'rish rejimi"ni Talabaga qo'ysa — CRM yopiladi.

        Bu ATAYLAB shunday (rejim aynan shuni sinash uchun). Frontend
        buni alohida xabar bilan tushuntiradi — `src-crm/App.jsx`."""
        User.objects.filter(pk=self.owner.pk).update(
            korish_rejimi=User.KorishRejimi.STUDENT
        )
        self.assertEqual(self.mijoz(self.owner).get("/api/crm/hisoblar/").status_code, 403)

        User.objects.filter(pk=self.owner.pk).update(
            korish_rejimi=User.KorishRejimi.OWNER
        )
        self.assertEqual(self.mijoz(self.owner).get("/api/crm/hisoblar/").status_code, 200)

    def test_hisob_summasini_admin_tuzata_olmaydi(self):
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        with bugun_qilib(date(2026, 9, 30)):
            mantiq.hisoblarni_generatsiya_qil()
        hisob = Hisob.objects.get()

        javob = self.mijoz(self.admin).patch(
            f"/api/crm/hisoblar/{hisob.id}/", {"summa": "100000"}, format="json"
        )
        self.assertEqual(javob.status_code, 403)

        javob = self.mijoz(self.owner).patch(
            f"/api/crm/hisoblar/{hisob.id}/", {"summa": "100000"}, format="json"
        )
        self.assertEqual(javob.status_code, 200)
        hisob.refresh_from_db()
        self.assertEqual(hisob.summa, Decimal("100000"))

    def test_tolovni_faqat_owner_ochiradi(self):
        tolov = Tolov.objects.create(
            talaba=self.talaba, talaba_ism="x", guruh=self.guruh, guruh_nomi="y",
            sana=date(2026, 9, 5), summa=NARX, turi=Tolov.Turi.TOLOV,
        )
        self.assertEqual(
            self.mijoz(self.admin).delete(f"/api/crm/tolov/{tolov.id}/").status_code, 403
        )
        self.assertEqual(
            self.mijoz(self.owner).delete(f"/api/crm/tolov/{tolov.id}/").status_code, 200
        )
        self.assertFalse(Tolov.objects.exists())

    def test_tolovni_faqat_owner_tahrirlaydi(self):
        """SoffCRM'dagi qalam: summa/sana/izoh tuzatiladi, oy holati
        qayta hisoblanadi. Admin — 403."""
        hisob = Hisob.objects.create(
            talaba=self.talaba, talaba_ism="x", guruh=self.guruh, guruh_nomi="y",
            oy=date(2026, 9, 1), summa=NARX,
        )
        tolov = Tolov.objects.create(
            talaba=self.talaba, talaba_ism="x", guruh=self.guruh, guruh_nomi="y",
            hisob=hisob, sana=date(2026, 9, 5), summa=NARX, turi=Tolov.Turi.TOLOV,
        )
        mantiq.hisobni_yangila(hisob)
        self.assertEqual(hisob.holat, Hisob.Holat.TOLANDI)

        self.assertEqual(
            self.mijoz(self.admin).patch(
                f"/api/crm/tolov/{tolov.id}/", {"summa": "100000"}, format="json"
            ).status_code, 403,
        )
        javob = self.mijoz(self.owner).patch(
            f"/api/crm/tolov/{tolov.id}/",
            {"summa": "100000", "sana": "2026-09-07", "izoh": "tuzatildi"},
            format="json",
        )
        self.assertEqual(javob.status_code, 200, javob.content)
        tolov.refresh_from_db()
        hisob.refresh_from_db()
        self.assertEqual(tolov.summa, Decimal("100000"))
        self.assertEqual(tolov.sana, date(2026, 9, 7))
        self.assertEqual(tolov.izoh, "tuzatildi")
        # Summa kamaygach oy qayta "qisman" bo'ladi — holat yagona joyda hisoblanadi.
        self.assertEqual(hisob.holat, Hisob.Holat.QISMAN)


class OqimTest(ApiAsos):
    def test_get_sorovi_hisobni_generatsiya_qiladi(self):
        """"Dangasa" generatsiya: cron yo'q, birinchi GET so'rovida
        ochiladi (TZ 4.1)."""
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        self.assertEqual(Hisob.objects.count(), 0)

        with bugun_qilib(date(2026, 9, 30)):
            javob = self.mijoz(self.admin).get("/api/crm/hisoblar/")

        self.assertEqual(javob.status_code, 200)
        self.assertEqual(Hisob.objects.count(), 1)
        self.assertEqual(len(javob.data), 1)
        self.assertEqual(javob.data[0]["holat"], "qarzdor")
        self.assertEqual(javob.data[0]["qoldiq"], NARX)

    def test_qisman_tolov_keyin_chegirma(self):
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        mijoz = self.mijoz(self.admin)
        with bugun_qilib(date(2026, 9, 30)):
            mijoz.get("/api/crm/hisoblar/")
        hisob = Hisob.objects.get()

        javob = mijoz.post(
            "/api/crm/tolov/",
            {"talaba_id": self.talaba.id, "guruh_id": self.guruh.id,
             "oy": "2026-09", "summa": "400000", "sana": "2026-09-11"},
            format="json",
        )
        self.assertEqual(javob.status_code, 201)
        hisob.refresh_from_db()
        self.assertEqual(hisob.holat, Hisob.Holat.QISMAN)

        mijoz.post(
            "/api/crm/tolov/",
            {"talaba_id": self.talaba.id, "guruh_id": self.guruh.id,
             "oy": "2026-09", "summa": "260000", "turi": "chegirma"},
            format="json",
        )
        hisob.refresh_from_db()
        self.assertEqual(hisob.holat, Hisob.Holat.TOLANDI)

        # Hisobotda "olingan pul" 400 000, chegirma 260 000 — aralashmaydi.
        with bugun_qilib(date(2026, 9, 30)):
            hisobot = mijoz.get("/api/crm/hisobot/?oy=2026-09").data
        self.assertEqual(hisobot["jami"]["olingan"], Decimal("400000"))
        self.assertEqual(hisobot["jami"]["chegirma"], Decimal("260000"))
        self.assertEqual(hisobot["jami"]["qarz"], Decimal("0"))
        self.assertAlmostEqual(hisobot["jami"]["yigilish_foizi"], 60.6, places=1)

    def test_boshqa_talabaning_hisobiga_tolov_rad_etiladi(self):
        """hisob_id boshqa talabaniki bo'lsa — 400. Aks holda to'lov bir
        talabaga yozilib, boshqasining oyi to'langan bo'lib qolardi."""
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        with bugun_qilib(date(2026, 9, 30)):
            mantiq.hisoblarni_generatsiya_qil()
        hisob = Hisob.objects.get()
        boshqa = User.objects.create_user(username="boshqa_t", password="x", role=User.Role.STUDENT)
        javob = self.mijoz(self.admin).post(
            "/api/crm/tolov/",
            {"talaba_id": boshqa.id, "guruh_id": self.guruh.id, "hisob_id": hisob.id,
             "summa": "1000", "turi": "tolov", "sana": "2026-09-11"},
            format="json",
        )
        self.assertEqual(javob.status_code, 400)
        self.assertFalse(Tolov.objects.exists())
        hisob.refresh_from_db()
        self.assertEqual(hisob.holat, Hisob.Holat.QARZDOR)

    def test_qaytarish_oyga_boglanmaydi(self):
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        mijoz = self.mijoz(self.admin)
        with bugun_qilib(date(2026, 9, 30)):
            mijoz.get("/api/crm/hisoblar/")
        hisob = Hisob.objects.get()

        mijoz.post(
            "/api/crm/tolov/",
            {"talaba_id": self.talaba.id, "guruh_id": self.guruh.id,
             "oy": "2026-09", "summa": "660000"},
            format="json",
        )
        mijoz.post(
            "/api/crm/tolov/",
            {"talaba_id": self.talaba.id, "guruh_id": self.guruh.id,
             "oy": "2026-09", "summa": "300000", "turi": "qaytarish"},
            format="json",
        )

        hisob.refresh_from_db()
        self.assertEqual(hisob.holat, Hisob.Holat.TOLANDI)
        self.assertIsNone(Tolov.objects.get(turi="qaytarish").hisob_id)
        self.assertEqual(mantiq.balans(self.talaba), Decimal("-300000"))

    def test_qolda_boshlangich_qarz(self):
        """Tizim yoqilgunga qadar bo'lgan qarz qo'lda kiritiladi (TZ 4.7)."""
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        mijoz = self.mijoz(self.admin)

        javob = mijoz.post(
            "/api/crm/hisoblar/",
            {"talaba_id": self.talaba.id, "guruh_id": self.guruh.id,
             "oy": "2026-08", "summa": "600000"},
            format="json",
        )
        self.assertEqual(javob.status_code, 201)
        self.assertTrue(javob.data["qolda"])

        with bugun_qilib(date(2026, 9, 30)):
            qatorlar = mijoz.get("/api/crm/hisoblar/").data
        self.assertEqual(len(qatorlar), 2)  # avgust (qo'lda) + sentabr (avtomatik)

        # Takror kiritishga yo'l qo'yilmaydi
        javob = mijoz.post(
            "/api/crm/hisoblar/",
            {"talaba_id": self.talaba.id, "guruh_id": self.guruh.id,
             "oy": "2026-08", "summa": "600000"},
            format="json",
        )
        self.assertEqual(javob.status_code, 400)

    def test_kurs_narxi_saqlanadi(self):
        KursNarxi.objects.all().delete()
        mijoz = self.mijoz(self.admin)

        javob = mijoz.put(
            "/api/crm/kurs-narxlari/",
            {"daraja_id": self.daraja.id, "narx": "450000"},
            format="json",
        )
        self.assertEqual(javob.status_code, 200)
        self.assertEqual(KursNarxi.objects.get().narx, Decimal("450000"))

        guruhlar = mijoz.get("/api/crm/guruhlar/").data
        self.assertEqual(guruhlar[0]["narx"], Decimal("450000"))
        self.assertEqual(guruhlar[0]["narx_manbasi"], "kurs")

    def test_guruh_narxi_kurs_narxini_bekor_qiladi(self):
        mijoz = self.mijoz(self.admin)
        mijoz.patch(
            f"/api/crm/guruhlar/{self.guruh.id}/moliya/", {"narx": "600000"}, format="json"
        )
        guruhlar = mijoz.get("/api/crm/guruhlar/").data
        self.assertEqual(guruhlar[0]["narx"], Decimal("600000"))
        self.assertEqual(guruhlar[0]["narx_manbasi"], "guruh")

    def test_jadval_almashtiriladi(self):
        mijoz = self.mijoz(self.admin)
        javob = mijoz.put(
            f"/api/crm/guruhlar/{self.guruh.id}/jadval/",
            {"jadval": [
                {"hafta_kuni": 0, "boshlanish_vaqti": "09:00", "tugash_vaqti": "10:30"},
                {"hafta_kuni": 2, "boshlanish_vaqti": "09:00", "tugash_vaqti": "10:30"},
            ]},
            format="json",
        )
        self.assertEqual(javob.status_code, 200)
        self.assertEqual(len(javob.data["jadval"]), 2)
        self.assertEqual(self.guruh.crm_jadval.count(), 2)

    def test_jadval_notogri_vaqt_rad_etiladi(self):
        javob = self.mijoz(self.admin).put(
            f"/api/crm/guruhlar/{self.guruh.id}/jadval/",
            {"jadval": [{"hafta_kuni": 0, "boshlanish_vaqti": "12:00", "tugash_vaqti": "10:00"}]},
            format="json",
        )
        self.assertEqual(javob.status_code, 400)
        self.assertEqual(self.guruh.crm_jadval.count(), 3)  # eskisi saqlanib qoldi

    def test_azolik_chiqarilsa_qayta_hisoblanadi(self):
        am = self.azolik_qosh(boshlanish=date(2026, 9, 1))
        mijoz = self.mijoz(self.admin)
        with bugun_qilib(date(2026, 9, 30)):
            mijoz.get("/api/crm/hisoblar/")
        self.assertEqual(Hisob.objects.get().summa, NARX)

        # Arxivlash CRM'dan EMAS (2026-09-16) — talaba saytda guruhdan
        # chiqariladi. CRM'da faqat muzlatish qoladi.
        javob = mijoz.patch(
            f"/api/crm/azoliklar/{am.id}/", {"holat": "arxiv"}, format="json"
        )
        self.assertEqual(javob.status_code, 400)

        javob = mijoz.patch(
            f"/api/crm/azoliklar/{am.id}/",
            {"holat": "muzlatilgan", "tugash_sana": "2026-09-12"},
            format="json",
        )
        self.assertEqual(javob.status_code, 200)
        # Joriy sana sentabr emas, shuning uchun qayta hisob joriy oyga
        # tegishli — testda sentabrni to'g'ridan-to'g'ri chaqiramiz.
        am.refresh_from_db()
        mantiq.azolikni_qayta_hisobla(am, SENTABR)
        self.assertEqual(Hisob.objects.get().summa, Decimal("330000"))

    def test_saytdan_chiqarilgan_talaba_ogohlantirishda(self):
        """Talaba saytda guruhdan chiqarilsa a'zolik (va AzolikMoliya)
        o'chadi, joriy oy hisobi esa to'liq summada qoladi — bu
        ogohlantirishda ko'rinishi shart, aks holda jimgina ortiqcha
        qarz turadi."""
        am = self.azolik_qosh(boshlanish=date(2026, 9, 1))
        mijoz = self.mijoz(self.admin)
        with bugun_qilib(date(2026, 9, 20)):
            mijoz.get("/api/crm/hisoblar/")
            self.assertEqual(Hisob.objects.count(), 1)

            am.azolik.delete()  # saytdagi "guruhdan chiqarish"
            self.assertEqual(Hisob.objects.count(), 1)  # pul tarixi qoladi

            javob = mijoz.get("/api/crm/ogohlantirishlar/")
        self.assertEqual(javob.status_code, 200)
        chiqqan = [x for x in javob.data if x.get("hisob_id")]
        self.assertEqual(len(chiqqan), 1)
        self.assertIn("guruhdan chiqarilgan", chiqqan[0]["sabablar"][0])

    def test_ogohlantirishlar_sozlanmagan_guruhni_korsatadi(self):
        KursNarxi.objects.all().delete()
        self.azolik_qosh()
        javob = self.mijoz(self.admin).get("/api/crm/ogohlantirishlar/")
        self.assertEqual(javob.status_code, 200)
        self.assertEqual(len(javob.data), 1)
        self.assertIn("narx yo'q", javob.data[0]["sabablar"])

    def test_talaba_kartasi(self):
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        mijoz = self.mijoz(self.admin)
        with bugun_qilib(date(2026, 9, 30)):
            mijoz.get("/api/crm/hisoblar/")
        javob = mijoz.get(f"/api/crm/talaba/{self.talaba.id}/")

        self.assertEqual(javob.status_code, 200)
        self.assertEqual(javob.data["balans_jami"], -NARX)
        self.assertEqual(len(javob.data["guruhlar"]), 1)
        self.assertEqual(len(javob.data["hisoblar"]), 1)

    def test_eksport_xlsx(self):
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        mijoz = self.mijoz(self.admin)
        with bugun_qilib(date(2026, 9, 30)):
            mijoz.get("/api/crm/hisoblar/")
            javob = mijoz.get("/api/crm/eksport/?oy=2026-09")

        self.assertEqual(javob.status_code, 200)
        self.assertIn("spreadsheetml", javob["Content-Type"])
        self.assertIn("filename*=UTF-8", javob["Content-Disposition"])
        # Haqiqiy xlsx ekanini tekshiramiz (ZIP imzosi) va varaqlarni o'qiymiz
        from io import BytesIO

        from openpyxl import load_workbook

        kitob = load_workbook(BytesIO(javob.content))
        self.assertEqual(kitob.sheetnames, ["Qarzdorlar", "To'lovlar", "Hisobot"])
        self.assertEqual(kitob["Qarzdorlar"]["E2"].value, float(NARX))

    def test_azoliklar_royxati_moliya_yozuvini_yaratadi(self):
        """`AzolikMoliya` signal bilan emas, birinchi murojaatda paydo
        bo'ladi (TZ 3.0, 3-qoida)."""
        from academics.models import Davomat, GuruhAzoligi

        GuruhAzoligi.objects.create(guruh=self.guruh, talaba=self.talaba)
        self.assertEqual(AzolikMoliya.objects.count(), 0)

        javob = self.mijoz(self.admin).get(f"/api/crm/guruhlar/{self.guruh.id}/azoliklar/")
        self.assertEqual(javob.status_code, 200)
        self.assertEqual(AzolikMoliya.objects.count(), 1)
        self.assertEqual(javob.data[0]["holat"], "faol")


class YangiRoyxatlarTest(ApiAsos):
    """`/talabalar/` va `/hisobot/dinamika/` — 2026-09-14 da qo'shilgan."""

    def test_talabalar_royxati_hisobsizlarni_ham_korsatadi(self):
        """Sinov va muzlatilgan talabaga hisob ochilmaydi, lekin ular
        ro'yxatda ko'rinishi SHART — aks holda admin ularni topa olmaydi
        va holatini o'zgartira olmaydi."""
        sinovchi = User.objects.create_user(
            username="sinovchi2", password="x", role=User.Role.STUDENT, markaz=self.markaz
        )
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        self.azolik_qosh(talaba=sinovchi, holat=AzolikMoliya.Holat.SINOV)

        mijoz = self.mijoz(self.admin)
        with bugun_qilib(date(2026, 9, 30)):
            hisoblar = mijoz.get("/api/crm/hisoblar/").data
            talabalar = mijoz.get("/api/crm/talabalar/").data

        # Hisob faqat faol talabaga ochiladi...
        self.assertEqual(len(hisoblar), 1)
        # ...lekin ro'yxatda ikkalasi ham bor.
        self.assertEqual(len(talabalar), 2)
        ismlar = {x["ism"] for x in talabalar}
        self.assertIn("sinovchi2", ismlar)

    def test_talabalar_holat_boyicha_filtrlanadi(self):
        muzlatilgan = User.objects.create_user(
            username="muz2", password="x", role=User.Role.STUDENT, markaz=self.markaz
        )
        self.azolik_qosh()
        self.azolik_qosh(talaba=muzlatilgan, holat=AzolikMoliya.Holat.MUZLATILGAN)

        javob = self.mijoz(self.admin).get("/api/crm/talabalar/?holat=muzlatilgan")
        self.assertEqual(len(javob.data), 1)
        self.assertEqual(javob.data[0]["ism"], "muz2")

    def test_dinamika_oylar_boyicha(self):
        Sozlama.objects.update(boshlangich_oy=date(2026, 7, 1))
        self.azolik_qosh(boshlanish=date(2026, 7, 1))
        mijoz = self.mijoz(self.admin)
        with bugun_qilib(date(2026, 9, 15)):
            mijoz.get("/api/crm/hisoblar/")
            javob = mijoz.get("/api/crm/hisobot/dinamika/?oylar=3")

        self.assertEqual(javob.status_code, 200)
        self.assertEqual(len(javob.data), 3)
        self.assertEqual([q["oy"] for q in javob.data],
                         [date(2026, 7, 1), AVGUST, SENTABR])
        # Uchala oyda ham to'liq narx hisoblangan, to'lov yo'q.
        self.assertTrue(all(q["hisoblangan"] == NARX for q in javob.data))
        self.assertTrue(all(q["yigilish_foizi"] == 0.0 for q in javob.data))

    def test_dinamika_eng_kop_24_oy(self):
        """Cheksiz oraliq so'ralsa ham so'rov cheklanadi."""
        Sozlama.objects.update(boshlangich_oy=date(2020, 1, 1))
        javob = self.mijoz(self.admin).get("/api/crm/hisobot/dinamika/?oylar=999")
        self.assertEqual(javob.status_code, 200)
        self.assertEqual(len(javob.data), 24)

    def test_dinamika_tizim_yoqilishidan_oldingi_oylarni_korsatmaydi(self):
        """Nol qatorlar "o'sha oyda hech kim to'lamagan" degan yolg'on
        taassurot beradi — tizim yoqilgan oygacha bo'lgani kesiladi."""
        Sozlama.objects.update(boshlangich_oy=SENTABR)
        with bugun_qilib(date(2026, 9, 15)):
            javob = self.mijoz(self.admin).get("/api/crm/hisobot/dinamika/?oylar=12")
        self.assertEqual(len(javob.data), 1)
        self.assertEqual(javob.data[0]["oy"], SENTABR)

    def test_filial_qoshiladi_va_arxivlanadi(self):
        mijoz = self.mijoz(self.admin)
        javob = mijoz.post("/api/crm/filiallar/", {"nomi": "Yangi filial"}, format="json")
        self.assertEqual(javob.status_code, 201)
        yangi_id = javob.data["id"]

        # Filial O'CHIRILMAYDI, arxivlanadi — hisoblarda u snapshot
        # sifatida turibdi va o'tgan oylar hisoboti buzilmasligi kerak.
        javob = mijoz.patch(f"/api/crm/filiallar/{yangi_id}/", {"faol": False}, format="json")
        self.assertEqual(javob.status_code, 200)
        self.assertFalse(javob.data["faol"])
        self.assertTrue(Filial.objects.filter(pk=yangi_id).exists())


class XonaVaSetkaTest(ApiAsos):
    """Xonalar va haftalik setka (2026-09-14, foydalanuvchi talabi bilan
    2-bosqichdan oldinga ko'chirildi)."""

    def setUp(self):
        super().setUp()
        self.xona = Xona.objects.create(filial=self.filial, nomi="2-xona", tartib=2)
        self.ikkinchi_guruh = Guruh.objects.create(
            name="Ikkinchi guruh", markaz=self.markaz, daraja=self.daraja
        )
        GuruhMoliya.objects.create(guruh=self.ikkinchi_guruh, filial=self.filial)

    def jadval_yubor(self, guruh, kun, dan, gacha, xona_id):
        return self.mijoz(self.admin).put(
            f"/api/crm/guruhlar/{guruh.id}/jadval/",
            {"jadval": [{"hafta_kuni": kun, "boshlanish_vaqti": dan,
                         "tugash_vaqti": gacha, "xona_id": xona_id}]},
            format="json",
        )

    def test_xona_qoshiladi(self):
        javob = self.mijoz(self.admin).post(
            "/api/crm/xonalar/",
            {"filial_id": self.filial.id, "nomi": "5-xona", "sigimi": 12},
            format="json",
        )
        self.assertEqual(javob.status_code, 201)
        self.assertEqual(javob.data["nomi"], "5-xona")

    def test_bir_filialda_takror_nom_rad_etiladi(self):
        javob = self.mijoz(self.admin).post(
            "/api/crm/xonalar/", {"filial_id": self.filial.id, "nomi": "2-xona"}, format="json"
        )
        self.assertEqual(javob.status_code, 400)

    def test_bir_xonada_ikki_guruh_toqnashadi(self):
        """Asosiy qoida: bitta xonada bir vaqtda ikkita guruh dars
        qila olmaydi."""
        self.assertEqual(
            self.jadval_yubor(self.guruh, 0, "14:00", "15:30", self.xona.id).status_code, 200
        )
        javob = self.jadval_yubor(self.ikkinchi_guruh, 0, "15:00", "16:30", self.xona.id)
        self.assertEqual(javob.status_code, 400)
        self.assertIn("2-xona", javob.data["detail"])
        # Rad etilgach, ikkinchi guruhda jadval YARATILMAYDI
        self.assertEqual(self.ikkinchi_guruh.crm_jadval.count(), 0)

    def test_chegara_teginishi_toqnashuv_emas(self):
        """10:30 da tugadi — 10:30 da boshlandi: bu to'qnashuv EMAS."""
        self.jadval_yubor(self.guruh, 0, "09:00", "10:30", self.xona.id)
        javob = self.jadval_yubor(self.ikkinchi_guruh, 0, "10:30", "12:00", self.xona.id)
        self.assertEqual(javob.status_code, 200)

    def test_boshqa_kunda_toqnashuv_yoq(self):
        self.jadval_yubor(self.guruh, 0, "14:00", "15:30", self.xona.id)
        javob = self.jadval_yubor(self.ikkinchi_guruh, 1, "14:00", "15:30", self.xona.id)
        self.assertEqual(javob.status_code, 200)

    def test_xonasiz_darslar_toqnashmaydi(self):
        """Xonasi belgilanmagan darslar joyi noma'lum — to'qnashuv
        tekshirilmaydi."""
        self.jadval_yubor(self.guruh, 0, "14:00", "15:30", None)
        javob = self.jadval_yubor(self.ikkinchi_guruh, 0, "14:00", "15:30", None)
        self.assertEqual(javob.status_code, 200)

    def test_oz_jadvalini_qayta_saqlash_toqnashmaydi(self):
        """Guruh o'z jadvalini qayta saqlasa, o'zi bilan to'qnashmasligi
        kerak — eski yozuvlar baribir almashtiriladi."""
        self.jadval_yubor(self.guruh, 0, "14:00", "15:30", self.xona.id)
        javob = self.jadval_yubor(self.guruh, 0, "14:00", "16:00", self.xona.id)
        self.assertEqual(javob.status_code, 200)

    def test_bitta_sorov_ichida_toqnashuv(self):
        javob = self.mijoz(self.admin).put(
            f"/api/crm/guruhlar/{self.guruh.id}/jadval/",
            {"jadval": [
                {"hafta_kuni": 0, "boshlanish_vaqti": "14:00",
                 "tugash_vaqti": "15:30", "xona_id": self.xona.id},
                {"hafta_kuni": 0, "boshlanish_vaqti": "15:00",
                 "tugash_vaqti": "16:30", "xona_id": self.xona.id},
            ]},
            format="json",
        )
        self.assertEqual(javob.status_code, 400)

    def test_setka_malumoti(self):
        self.jadval_yubor(self.guruh, 3, "14:00", "15:30", self.xona.id)
        javob = self.mijoz(self.admin).get("/api/crm/jadval/")
        self.assertEqual(javob.status_code, 200)
        self.assertEqual(len(javob.data["xonalar"]), 1)
        dars = next(d for d in javob.data["darslar"] if d["guruh_id"] == self.guruh.id)
        self.assertEqual(dars["hafta_kuni"], 3)
        self.assertEqual(dars["boshlanish_vaqti"], "14:00")
        self.assertEqual(dars["xona_id"], self.xona.id)

    def test_xona_ochirilsa_dars_qoladi(self):
        """Xona `SET_NULL` bilan bog'langan: o'chsa dars yozuvi
        yo'qolmaydi, faqat "xonasiz" bo'lib qoladi."""
        self.jadval_yubor(self.guruh, 0, "14:00", "15:30", self.xona.id)
        self.xona.delete()
        dars = self.guruh.crm_jadval.get()
        self.assertIsNone(dars.xona_id)
        self.assertEqual(dars.boshlanish_vaqti.strftime("%H:%M"), "14:00")


class DavomatVaNatijaTest(ApiAsos):
    """Davomat va natijalar — FAQAT O'QISH (2026-09-14).

    Ular LMS'da hosil bo'ladi; CRM faqat ko'rsatadi. Shuning uchun
    testlar ham LMS modellariga yozib, CRM API'sidan o'qiydi.
    """

    def setUp(self):
        super().setUp()
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        self.ikkinchi = User.objects.create_user(
            username="talaba2", password="x", role=User.Role.STUDENT, markaz=self.markaz
        )
        self.azolik_qosh(talaba=self.ikkinchi, boshlanish=date(2026, 9, 1))

    def davomat_yoz(self, talaba, kun, holat):
        return Davomat.objects.create(
            sana=date(2026, 9, kun), guruh=self.guruh, talaba=talaba, holat=holat
        )

    def test_davomat_matritsasi(self):
        self.davomat_yoz(self.talaba, 3, Davomat.Holat.KELDI)
        self.davomat_yoz(self.talaba, 4, Davomat.Holat.KELMADI)
        self.davomat_yoz(self.ikkinchi, 3, Davomat.Holat.KELDI)
        # 4-sentabrda ikkinchi talabaga yozuv YO'Q — katak bo'sh qolishi kerak

        javob = self.mijoz(self.admin).get(
            f"/api/crm/guruhlar/{self.guruh.id}/davomat/?oy=2026-09"
        )
        self.assertEqual(javob.status_code, 200)
        # Ustunlar — faqat dars bo'lgan sanalar
        self.assertEqual(javob.data["sanalar"], [date(2026, 9, 3), date(2026, 9, 4)])

        qatorlar = {x["ism"]: x for x in javob.data["talabalar"]}
        self.assertEqual(qatorlar["talaba1"]["kunlar"], ["keldi", "kelmadi"])
        self.assertEqual(qatorlar["talaba1"]["keldi"], 1)
        self.assertEqual(qatorlar["talaba1"]["kelmadi"], 1)
        self.assertEqual(qatorlar["talaba2"]["kunlar"], ["keldi", None])

    def test_davomat_boshqa_oy_aralashmaydi(self):
        self.davomat_yoz(self.talaba, 3, Davomat.Holat.KELDI)
        Davomat.objects.create(
            sana=date(2026, 8, 20), guruh=self.guruh, talaba=self.talaba,
            holat=Davomat.Holat.KELDI,
        )
        javob = self.mijoz(self.admin).get(
            f"/api/crm/guruhlar/{self.guruh.id}/davomat/?oy=2026-09"
        )
        self.assertEqual(javob.data["sanalar"], [date(2026, 9, 3)])

    def test_davomat_talaba_uchun_403(self):
        self.assertEqual(
            self.mijoz(self.talaba).get(
                f"/api/crm/guruhlar/{self.guruh.id}/davomat/"
            ).status_code,
            403,
        )

    def test_natijalar_davomat_foizi(self):
        self.davomat_yoz(self.talaba, 3, Davomat.Holat.KELDI)
        self.davomat_yoz(self.talaba, 4, Davomat.Holat.KELDI)
        self.davomat_yoz(self.talaba, 5, Davomat.Holat.KELMADI)

        javob = self.mijoz(self.admin).get(
            f"/api/crm/guruhlar/{self.guruh.id}/natijalar/"
        )
        self.assertEqual(javob.status_code, 200)
        qator = next(x for x in javob.data["talabalar"] if x["ism"] == "talaba1")
        self.assertEqual(qator["keldi"], 2)
        self.assertEqual(qator["kelmadi"], 1)
        self.assertEqual(qator["davomat_foizi"], 67)

    def test_natijalar_writing_ortachasi(self):
        for band in (6.0, 7.0):
            WritingTekshiruv.objects.create(
                talaba=self.talaba, matn="sinov", overall_band=band
            )
        javob = self.mijoz(self.admin).get(
            f"/api/crm/guruhlar/{self.guruh.id}/natijalar/"
        )
        qator = next(x for x in javob.data["talabalar"] if x["ism"] == "talaba1")
        self.assertEqual(qator["writing_band"], 6.5)
        self.assertEqual(qator["writing_soni"], 2)

    def test_natijalar_malumotsiz_talaba(self):
        """Hech narsa yechmagan talaba ham ro'yxatda bo'lishi kerak —
        aks holda admin uni topa olmaydi."""
        javob = self.mijoz(self.admin).get(
            f"/api/crm/guruhlar/{self.guruh.id}/natijalar/"
        )
        self.assertEqual(len(javob.data["talabalar"]), 2)
        qator = javob.data["talabalar"][0]
        self.assertIsNone(qator["writing_band"])
        self.assertIsNone(qator["davomat_foizi"])
        self.assertEqual(qator["mashq_soni"], 0)


class BayroqTest(ApiAsos):
    """`settings.CRM_YOQILGAN` — CRM prodga chiqmasligi uchun qulf
    (2026-09-15, Shuhrat: "to'liq ulaymiz, lekin prodga chiqmaydigan
    qilib").

    Prodda `DEBUG=False`, ya'ni bayroq standart bo'yicha o'chiq va
    `/api/crm/...` umuman yo'q bo'lib ko'rinadi.
    """

    @override_settings(CRM_YOQILGAN=False)
    def test_bayroq_ochiq_bolsa_404(self):
        """404, ATAYLAB 403 emas: 403 "bu yerda nimadir bor" degani,
        404 esa bo'limning borligini ham oshkor qilmaydi."""
        mijoz = self.mijoz(self.owner)
        for yol in ("/api/crm/hisoblar/", "/api/crm/guruhlar/", "/api/crm/hisobot/"):
            self.assertEqual(mijoz.get(yol).status_code, 404, yol)

    @override_settings(CRM_YOQILGAN=False)
    def test_bayroq_ochiq_bolsa_yozish_ham_404(self):
        javob = self.mijoz(self.owner).post(
            "/api/crm/filiallar/", {"nomi": "Yangi"}, format="json"
        )
        self.assertEqual(javob.status_code, 404)

    def test_bayroq_yoqiq_bolsa_ishlaydi(self):
        # `CrmAsos` da @override_settings(CRM_YOQILGAN=True) turibdi
        self.assertEqual(self.mijoz(self.owner).get("/api/crm/hisoblar/").status_code, 200)


class EslatmaTest(ApiAsos):
    """Eslatmalar — LMS'da ham, CRM'da ham bo'lmagan yagona narsa edi
    (2026-09-15, SoffCRM'dagi "ESLATMALAR" tabi)."""

    def test_guruhga_eslatma_qoshiladi(self):
        javob = self.mijoz(self.admin).post(
            "/api/crm/eslatmalar/",
            {"guruh_id": self.guruh.id, "matn": "Dars vaqtini ko'chirishni so'radi"},
            format="json",
        )
        self.assertEqual(javob.status_code, 201)
        self.assertEqual(javob.data["kim"], "crm_admin")

        royxat = self.mijoz(self.admin).get(f"/api/crm/eslatmalar/?guruh={self.guruh.id}")
        self.assertEqual(len(royxat.data), 1)

    def test_talabaga_eslatma_qoshiladi(self):
        javob = self.mijoz(self.admin).post(
            "/api/crm/eslatmalar/",
            {"talaba_id": self.talaba.id, "matn": "Onasi 15-sentabrda to'layman dedi"},
            format="json",
        )
        self.assertEqual(javob.status_code, 201)
        royxat = self.mijoz(self.admin).get(f"/api/crm/eslatmalar/?talaba={self.talaba.id}")
        self.assertEqual(len(royxat.data), 1)

    def test_egasiz_eslatma_rad_etiladi(self):
        """Eslatma nimagadir tegishli bo'lmasa, uni hech kim qayta topa
        olmaydi — shuning uchun kamida bittasi majburiy."""
        javob = self.mijoz(self.admin).post(
            "/api/crm/eslatmalar/", {"matn": "shunchaki"}, format="json"
        )
        self.assertEqual(javob.status_code, 400)

    def test_bosh_matn_rad_etiladi(self):
        javob = self.mijoz(self.admin).post(
            "/api/crm/eslatmalar/",
            {"guruh_id": self.guruh.id, "matn": "   "},
            format="json",
        )
        self.assertEqual(javob.status_code, 400)

    def test_filtrsiz_royxat_rad_etiladi(self):
        """Barcha eslatmalarni birdaniga berish ma'nosiz — ular doim
        aniq guruh yoki talaba kontekstida o'qiladi."""
        self.assertEqual(
            self.mijoz(self.admin).get("/api/crm/eslatmalar/").status_code, 400
        )

    def test_boshqaning_eslatmasini_admin_ochira_olmaydi(self):
        yaratildi = self.mijoz(self.owner).post(
            "/api/crm/eslatmalar/",
            {"guruh_id": self.guruh.id, "matn": "owner yozgan"},
            format="json",
        ).data

        javob = self.mijoz(self.admin).delete(f"/api/crm/eslatmalar/{yaratildi['id']}/")
        self.assertEqual(javob.status_code, 403)

        # Owner — o'zinikini o'chira oladi
        javob = self.mijoz(self.owner).delete(f"/api/crm/eslatmalar/{yaratildi['id']}/")
        self.assertEqual(javob.status_code, 200)

    def test_owner_boshqaning_eslatmasini_ochira_oladi(self):
        yaratildi = self.mijoz(self.admin).post(
            "/api/crm/eslatmalar/",
            {"guruh_id": self.guruh.id, "matn": "admin yozgan"},
            format="json",
        ).data
        javob = self.mijoz(self.owner).delete(f"/api/crm/eslatmalar/{yaratildi['id']}/")
        self.assertEqual(javob.status_code, 200)

    def test_talabaga_403(self):
        self.assertEqual(
            self.mijoz(self.talaba).get(
                f"/api/crm/eslatmalar/?guruh={self.guruh.id}"
            ).status_code,
            403,
        )


class KeyingiTolovTest(ApiAsos):
    def test_tolanmagan_oy_korsatiladi(self):
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        mijoz = self.mijoz(self.admin)
        with bugun_qilib(date(2026, 9, 30)):
            mijoz.get("/api/crm/hisoblar/")
            javob = mijoz.get(f"/api/crm/talaba/{self.talaba.id}/")
        self.assertEqual(javob.data["guruhlar"][0]["keyingi_tolov"], SENTABR)

    def test_hammasi_tolangan_bolsa_keyingi_oy(self):
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        mijoz = self.mijoz(self.admin)
        with bugun_qilib(date(2026, 9, 30)):
            mijoz.get("/api/crm/hisoblar/")
            mijoz.post(
                "/api/crm/tolov/",
                {"talaba_id": self.talaba.id, "guruh_id": self.guruh.id,
                 "oy": "2026-09", "summa": "660000"},
                format="json",
            )
            javob = mijoz.get(f"/api/crm/talaba/{self.talaba.id}/")
        self.assertEqual(javob.data["guruhlar"][0]["keyingi_tolov"], date(2026, 10, 1))


class TalabaNatijalariTest(ApiAsos):
    """Talaba kartasidagi umumiy o'quv natijasi (SoffCRM: "Baho")."""

    def test_natijalar_kartada_chiqadi(self):
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        WritingTekshiruv.objects.create(talaba=self.talaba, matn="x", overall_band=6.5)
        Davomat.objects.create(
            sana=date(2026, 9, 3), guruh=self.guruh, talaba=self.talaba,
            holat=Davomat.Holat.KELDI,
        )
        Davomat.objects.create(
            sana=date(2026, 9, 4), guruh=self.guruh, talaba=self.talaba,
            holat=Davomat.Holat.KELMADI,
        )

        javob = self.mijoz(self.admin).get(f"/api/crm/talaba/{self.talaba.id}/")
        natijalar = javob.data["natijalar"]
        self.assertEqual(natijalar["writing_band"], 6.5)
        self.assertEqual(natijalar["davomat_foizi"], 50)
        self.assertEqual(natijalar["keldi"], 1)

    def test_malumotsiz_talabada_none(self):
        """0% deb ko'rsatish "yomon natija" degan yolg'on taassurot
        berardi — ma'lumot yo'q bo'lsa `None`."""
        self.azolik_qosh()
        javob = self.mijoz(self.admin).get(f"/api/crm/talaba/{self.talaba.id}/")
        natijalar = javob.data["natijalar"]
        self.assertIsNone(natijalar["writing_band"])
        self.assertIsNone(natijalar["davomat_foizi"])
        self.assertEqual(natijalar["mashq_soni"], 0)


class TirikNomTest(ApiAsos):
    """Nom SNAPSHOT'dan emas, LMS'dan olinishi kerak.

    2026-09-16, Shuhrat topdi: "ma'lumotlarni asosiy saytdan olmayabdi".
    `Hisob`/`Tolov` da `talaba_ism` snapshot saqlanadi (talaba
    o'chirilsa yozuv o'qiladigan bo'lib qolishi uchun), lekin
    ro'yxatlarda AYNAN SHU snapshot ko'rsatilardi — admin saytda ismni
    tuzatsa, CRM eskisini ko'rsatib turardi.
    """

    def setUp(self):
        super().setUp()
        self.talaba.first_name = "Aziz"
        self.talaba.last_name = "Qodirov"
        self.talaba.save()
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        with bugun_qilib(date(2026, 9, 30)):
            mantiq.hisoblarni_generatsiya_qil()

    def test_lmsda_ism_ozgarsa_crm_yangisini_korsatadi(self):
        mijoz = self.mijoz(self.admin)
        with bugun_qilib(date(2026, 9, 30)):
            self.assertEqual(mijoz.get("/api/crm/hisoblar/").data[0]["talaba"], "Aziz Qodirov")

        User.objects.filter(pk=self.talaba.pk).update(last_name="Yangi")
        with bugun_qilib(date(2026, 9, 30)):
            javob = mijoz.get("/api/crm/hisoblar/")

        self.assertEqual(javob.data[0]["talaba"], "Aziz Yangi")
        # Snapshot ataylab ESKILIGICHA qoladi — u o'chirilgan talaba uchun
        self.assertEqual(Hisob.objects.get().talaba_ism, "Aziz Qodirov")

    def test_lmsda_guruh_nomi_ozgarsa_crm_yangisini_korsatadi(self):
        Guruh.objects.filter(pk=self.guruh.pk).update(name="Yangi guruh nomi")
        with bugun_qilib(date(2026, 9, 30)):
            javob = self.mijoz(self.admin).get("/api/crm/hisoblar/")
        self.assertEqual(javob.data[0]["guruh"], "Yangi guruh nomi")

    def test_qidiruv_yangi_ism_boyicha_ishlaydi(self):
        User.objects.filter(pk=self.talaba.pk).update(last_name="Yangi")
        with bugun_qilib(date(2026, 9, 30)):
            javob = self.mijoz(self.admin).get("/api/crm/hisoblar/?q=Yangi")
        self.assertEqual(len(javob.data), 1)

    def test_talaba_ochirilsa_snapshot_ishlatiladi(self):
        """FK NULL bo'lgach yagona manba — snapshot."""
        self.talaba.delete()
        with bugun_qilib(date(2026, 9, 30)):
            javob = self.mijoz(self.admin).get("/api/crm/hisoblar/")
        self.assertEqual(javob.data[0]["talaba"], "Aziz Qodirov")
        self.assertIsNone(javob.data[0]["talaba_id"])


class BonusTest(ApiAsos):
    """`bonus` TZ qamrovida (§2) — chegirma bilan bir xil ishlaydi,
    lekin hisobotda ALOHIDA ustunda (2026-09-16)."""

    def setUp(self):
        super().setUp()
        self.azolik_qosh(boshlanish=date(2026, 9, 1))
        with bugun_qilib(date(2026, 9, 30)):
            mantiq.hisoblarni_generatsiya_qil()

    def test_bonus_qarzni_yopadi(self):
        mijoz = self.mijoz(self.admin)
        mijoz.post(
            "/api/crm/tolov/",
            {"talaba_id": self.talaba.id, "guruh_id": self.guruh.id,
             "oy": "2026-09", "summa": "660000", "turi": "bonus"},
            format="json",
        )
        self.assertEqual(Hisob.objects.get().holat, Hisob.Holat.TOLANDI)
        self.assertEqual(mantiq.balans(self.talaba), Decimal("0"))

    def test_bonus_kassaga_pul_qoshmaydi(self):
        """Hisobotda "olingan pul" 0 bo'lishi kerak — bonus pul emas."""
        mijoz = self.mijoz(self.admin)
        mijoz.post(
            "/api/crm/tolov/",
            {"talaba_id": self.talaba.id, "guruh_id": self.guruh.id,
             "oy": "2026-09", "summa": "660000", "turi": "bonus"},
            format="json",
        )
        with bugun_qilib(date(2026, 9, 30)):
            hisobot = mijoz.get("/api/crm/hisobot/?oy=2026-09").data
        self.assertEqual(hisobot["jami"]["olingan"], Decimal("0"))
        self.assertEqual(hisobot["jami"]["bonus"], Decimal("660000"))
        self.assertEqual(hisobot["jami"]["chegirma"], Decimal("0"))
        self.assertEqual(hisobot["jami"]["qarz"], Decimal("0"))
