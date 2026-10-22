"""Qarzdorlik — bosh sahifa va Talabalar -> "Qarzdorlar" BITTA qoida bilan
(2026-09-28, Shuhrat): qarzdor — balansi manfiy o'quvchi; filial tanlansa,
balans faqat o'sha filialning hisob va to'lovlaridan.

Avval bosh sahifa har oy hisobini alohida qarardi: bir oyga ortiqcha
to'langan pul boshqa oy qarzini yopmasdi, arxivlangan / guruhdan chiqqan
qarzdor esa ro'yxatda chiqmasdi — prod'da "4 qarzdor", ro'yxat bo'sh edi.
"""

from datetime import date
from decimal import Decimal

from academics.models import Guruh, GuruhAzoligi
from crm.models import AzolikMoliya, Filial, GuruhMoliya, Hisob, Tolov
from crm.test_api import ApiAsos
from crm.tests import AVGUST, NARX, SENTABR, bugun_qilib


class QarzdorlikTest(ApiAsos):
    def setUp(self):
        super().setUp()
        # Sana sentabrda qotadi: aks holda har GET so'rovi bugungi oy uchun
        # qo'shimcha hisob ochib, qarzni ikki barobar qilardi.
        self.enterContext(bugun_qilib(date(2026, 9, 30)))
        self.azolik_qosh()

    def hisob(self, oy, guruh=None, filial=None, summa=NARX):
        guruh = guruh or self.guruh
        return Hisob.objects.create(
            talaba=self.talaba, talaba_ism="T", guruh=guruh, guruh_nomi=guruh.name,
            filial=filial or self.filial, oy=oy, summa=summa, holat=Hisob.Holat.QARZDOR,
        )

    def tolov(self, hisob, summa):
        from crm import mantiq

        Tolov.objects.create(talaba=self.talaba, talaba_ism="T", guruh=hisob.guruh, guruh_nomi="G", hisob=hisob,
                             sana=date(2026, 9, 5), summa=summa, turi=Tolov.Turi.TOLOV)
        mantiq.hisobni_yangila(hisob)

    def ikkalasi(self, filial=None):
        """(bosh sahifa: qarzdorlar, qolgan qarz) va Talabalar -> Qarzdorlar ro'yxati."""
        m = self.mijoz(self.owner)
        f = f"filial={filial}" if filial else ""
        k = m.get(f"/api/crm/korsatkichlar/?{f}").data
        royxat = m.get(f"/api/crm/talabalar/?qarzdor=1&{f}").data
        # Ikkala joy har doim bir xil sonni va summani ko'rsatadi.
        self.assertEqual(k["qarzdorlar"], len(royxat))
        self.assertEqual(k["jami_qarz"], -sum((x["balans"] for x in royxat), Decimal(0)))
        return k["qarzdorlar"], k["jami_qarz"], royxat

    def test_tepadagi_karta_joriy_oy_pastdagi_jami(self):
        """Video-TZ 2026-10-16: "Qolgan qarz" — faqat joriy (sentabr) oy,
        "jami_qarz" — barcha oylar; `?qarzli=1` — kim, qaysi oy uchun qarzdor."""
        self.hisob(AVGUST)
        self.tolov(self.hisob(SENTABR), NARX - 100000)
        m = self.mijoz(self.owner)
        k = m.get("/api/crm/korsatkichlar/").data
        self.assertEqual(k["qolgan_qarz"], 100000)
        self.assertEqual(k["jami_qarz"], NARX + 100000)
        d = m.get("/api/crm/hisoblar/?qarzli=1").data
        self.assertEqual(sorted(str(x["oy"]) for x in d), [str(AVGUST), str(SENTABR)])
        self.assertEqual(sum((x["qoldiq"] for x in d), Decimal(0)), NARX + 100000)
        # Excel ham shu filtr bilan (avval `holat=qarzli` bo'sh varaq berardi);
        # to'langan iyul hisobi tushmaydi.
        self.tolov(self.hisob(date(2026, 7, 1)), NARX)
        from io import BytesIO

        from openpyxl import load_workbook

        kitob = load_workbook(BytesIO(m.get("/api/crm/eksport/?qarzli=1").content))
        self.assertEqual(kitob["Qarzdorlar"].max_row - 1, 2)

    def test_oy_qarzi_ortiqcha_tolov_boshqa_qarzni_yemaydi(self):
        """Video-TZ IMG_2210: bir guruhda A 100 000 ortiqcha to'lagan, B to'lamagan —
        "Shu oy qarzi" B ning to'liq qarzi, ro'yxat yig'indisi bilan teng."""
        self.tolov(self.hisob(SENTABR), NARX + 100000)
        ikkinchi = self.talaba
        from accounts.models import User

        self.talaba = User.objects.create(username="sinov_b", role=ikkinchi.role, markaz=self.markaz)
        self.hisob(SENTABR)
        m = self.mijoz(self.owner)
        hisobot = m.get(f"/api/crm/hisobot/?oy={SENTABR:%Y-%m}").data
        self.assertEqual(hisobot["jami"]["qarz"], NARX)
        # Guruh qatori ham shu qoida bilan — jami qator ustunlar yig'indisiga teng.
        self.assertEqual([q["qarz"] for q in hisobot["qatorlar"]], [NARX])
        self.assertEqual(m.get("/api/crm/korsatkichlar/").data["qolgan_qarz"], NARX)
        d = m.get(f"/api/crm/hisoblar/?qarzli=1&oy={SENTABR:%Y-%m}").data
        self.assertEqual(sum((x["qoldiq"] for x in d), Decimal(0)), NARX)

    def test_boshqa_oyga_ortiqcha_tolangan_pul_qarz_emas(self):
        """Avgustga 2 oylik pul — sentabr hisobi "ochiq", lekin balans 0: qarzdor emas."""
        avgust = self.hisob(AVGUST)
        self.hisob(SENTABR)
        self.tolov(avgust, NARX * 2)
        soni, summa, _ = self.ikkalasi()
        self.assertEqual((soni, summa), (0, 0))

    def test_tekin_oquvchi_qarzlilar_royxatida_yoq(self):
        """Matn-TZ #465: to'liq chegirma oyga bog'lanmagan — sentabr hisobi
        "qarzdor", balans 0. "Qarzdor + Qisman" ro'yxatida va Excel'da yo'q."""
        self.hisob(SENTABR)
        Tolov.objects.create(talaba=self.talaba, talaba_ism="T", guruh=self.guruh, guruh_nomi="G",
                             sana=date(2026, 9, 5), summa=NARX, turi=Tolov.Turi.CHEGIRMA)
        m = self.mijoz(self.owner)
        self.assertEqual(m.get("/api/crm/hisoblar/?qarzli=1").data, [])
        # Holat filtri o'zgarmagan: barcha holatlarda hisob ko'rinadi.
        self.assertEqual(len(m.get("/api/crm/hisoblar/").data), 1)
        from io import BytesIO

        from openpyxl import load_workbook

        kitob = load_workbook(BytesIO(m.get("/api/crm/eksport/?qarzli=1").content))
        self.assertEqual(kitob["Qarzdorlar"].max_row - 1, 0)

    def test_ortiqcha_tolagan_balansi_musbat(self):
        self.tolov(self.hisob(SENTABR), NARX + 100000)
        d = self.mijoz(self.owner).get("/api/crm/talabalar/").data
        self.assertEqual([x["balans"] for x in d if x["id"] == self.talaba.id], [Decimal("100000")])

    def test_arxivlangan_qarzdor_royxatda_chiqadi(self):
        self.hisob(SENTABR)
        GuruhAzoligi.objects.filter(talaba=self.talaba).delete()
        self.talaba.is_active = False
        self.talaba.save()
        soni, summa, royxat = self.ikkalasi()
        self.assertEqual((soni, summa), (1, NARX))
        self.assertEqual(royxat[0]["id"], self.talaba.id)
        self.assertFalse(royxat[0]["faol"])

    def test_guruhsiz_qarzdor_filial_tanlanganda_ham_chiqadi(self):
        self.hisob(SENTABR)
        GuruhAzoligi.objects.filter(talaba=self.talaba).delete()
        self.assertEqual(self.ikkalasi()[0], 1)
        self.assertEqual(self.ikkalasi(filial=self.filial.id)[0], 1)

    def test_filial_tanlansa_balans_shu_filial_boyicha(self):
        """A filialda 100 000 ortiqcha to'lagan, B filialda 660 000 qarz:
        A — qarzdor emas, B — qarzdor 660 000, umumiy — 560 000."""
        fb = Filial.objects.create(markaz=self.markaz, nomi="B")
        gb = Guruh.objects.create(name="B guruh", markaz=self.markaz, daraja=self.daraja)
        GuruhMoliya.objects.create(guruh=gb, filial=fb, boshlanish_sana=date(2026, 1, 1))
        AzolikMoliya.objects.create(azolik=GuruhAzoligi.objects.create(guruh=gb, talaba=self.talaba),
                                    boshlanish_sana=SENTABR)
        self.tolov(self.hisob(SENTABR), NARX + 100000)
        self.hisob(SENTABR, guruh=gb, filial=fb)
        self.assertEqual(self.ikkalasi(filial=self.filial.id)[:2], (0, 0))
        self.assertEqual(self.ikkalasi(filial=fb.id)[:2], (1, NARX))
        self.assertEqual(self.ikkalasi()[:2], (1, NARX - 100000))
        # Ro'yxatdagi balans ustuni ham tanlangan filial bo'yicha.
        d = self.mijoz(self.owner).get(f"/api/crm/talabalar/?filial={self.filial.id}").data
        self.assertEqual([x["balans"] for x in d if x["id"] == self.talaba.id], [Decimal("100000")])
