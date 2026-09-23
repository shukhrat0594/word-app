"""CRM — video-TZ'ni 5 soniyalik kadrlar bilan qayta tekshirishda
topilgan qo'shimchalar (2026-09-23).

Lid doskalari, dars mavzulari, "hammasi keldi", Excel eksport/import,
guruhdagi ommaviy faollashtirish va sobiq a'zolar, o'quvchi tarixi,
xodimlar davomati, harakatlar tarixi va hisobotlar (ketganlar, lidlar,
to'lovlar, bitiruvchilar).

Konvensiya `crm/views.py` va `crm/boshqaruv.py` bilan bir xil.
"""

from collections import Counter
from datetime import date

from django.db import transaction
from django.db.models import Count, Q, Sum
from django.http import Http404
from django.shortcuts import get_object_or_404
from django.utils import timezone
from openpyxl import Workbook
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from academics.models import Davomat, Guruh, GuruhAzoligi
from accounts.models import User
from audit.models import FaoliyatYozuvi

from . import mantiq
from .boshqaruv import (
    _ism,
    _telefon,
    _xodimlar_qs,
    guruh_dars_sanalari,
    talaba_yarat,
)
from .eksport import _varaq_yoz, javob_qil
from .models import (
    AzolikMoliya,
    DarsMavzusi,
    GuruhdanChiqish,
    Lid,
    LidBolim,
    LidDoska,
    TalabaProfil,
    Tolov,
    XodimDavomat,
)
from .permissions import CrmView
from .filial import filial_q, guruh_q, xodim_korinadimi
from .ruxsatlar import ruxsatlar
from .views import _oy, _ruxsatsiz, _sana, _xato

NOL = 0


def _davr(request):
    """`?dan=YYYY-MM-DD&gacha=...` — standart: joriy oy boshidan bugungacha."""
    bugun = timezone.localdate()
    dan = _sana(request.query_params.get("dan"), "dan", majburiy=False) or mantiq.oy_boshi(bugun)
    gacha = _sana(request.query_params.get("gacha"), "gacha", majburiy=False) or bugun
    return dan, gacha


def _excel(nomi, sarlavhalar, qatorlar, fayl_nomi, pul_ustunlari=()):
    kitob = Workbook()
    ws = kitob.active
    ws.title = nomi[:31]
    _varaq_yoz(ws, sarlavhalar, qatorlar, pul_ustunlari=pul_ustunlari)
    return javob_qil(kitob, f"{fayl_nomi}_{timezone.localdate():%Y-%m-%d}.xlsx")


# ── Lid doskalari ────────────────────────────────────────────────────


def _doska_dict(d, soni=None):
    return {"id": d.id, "nomi": d.nomi, "tartib": d.tartib, "soni": soni}


class LidDoskalarView(CrmView):
    """SoffCRM "Bo'lim yaratish": lidlar doskasi. Yangi doska bitta
    "NEW LEADS" ustuni bilan ochiladi — bo'sh doskaga lid qo'shib
    bo'lmasdi."""

    bolim = "lidlar"

    def get(self, request):
        u = request.user
        doskalar = list(LidDoska.objects.annotate(
            _soni=Count("ustunlar__lidlar", filter=Q(ustunlar__lidlar__arxiv=False,
                                                     ustunlar__lidlar__qora_royxat=False)
                        & filial_q(u, "ustunlar__lidlar__filial"))
        ))
        # Doskasiz ("Umumiy") — video-TZ'dan oldingi ustunlar va ustunsiz lidlar.
        umumiy = Lid.objects.filter(arxiv=False, qora_royxat=False).filter(filial_q(u, "filial")).filter(
            Q(bolim__isnull=True) | Q(bolim__doska__isnull=True)
        ).count()
        return Response({
            "doskalar": [_doska_dict(d, d._soni) for d in doskalar],
            "umumiy": umumiy,
            "qora_royxat": Lid.objects.filter(qora_royxat=True).filter(filial_q(u, "filial")).count(),
        })

    def post(self, request):
        if "lidlar.bolim" not in ruxsatlar(request.user):
            return _xato("Bo'lim yaratishga ruxsat yo'q", kod=403)
        nomi = (request.data.get("nomi") or "").strip()
        if not nomi:
            return _xato("Bo'lim nomi bo'sh bo'lmasin")
        oxirgi = LidDoska.objects.order_by("-tartib").values_list("tartib", flat=True).first() or 0
        with transaction.atomic():
            doska = LidDoska.objects.create(nomi=nomi[:100], tartib=oxirgi + 1)
            LidBolim.objects.create(nomi="NEW LEADS", doska=doska, tartib=0)
        return Response(_doska_dict(doska, 0), status=201)


class LidDoskaDetailView(CrmView):
    bolim = "lidlar"

    def patch(self, request, pk):
        if xato := _ruxsatsiz(request, "lidlar.bolim"):
            return xato
        doska = get_object_or_404(LidDoska, pk=pk)
        nomi = (request.data.get("nomi") or "").strip()
        if not nomi:
            return _xato("Bo'lim nomi bo'sh bo'lmasin")
        doska.nomi = nomi[:100]
        doska.save(update_fields=["nomi"])
        return Response(_doska_dict(doska))

    def delete(self, request, pk):
        if "lidlar.bolim" not in ruxsatlar(request.user):
            return _xato("Ruxsat yo'q", kod=403)
        doska = get_object_or_404(LidDoska, pk=pk)
        # Lidlar yo'qolmaydi: ustunlar o'chadi, lidlar "Umumiy"ga tushadi.
        Lid.objects.filter(bolim__doska=doska).update(bolim=None)
        doska.delete()
        return Response(status=204)


# ── Dars mavzulari va "hammasi keldi" ────────────────────────────────


class DarsMavzulariView(CrmView):
    """Davomat jadvalidagi "Mavzular" qatori: `{sana, mavzu}` (bo'sh —
    o'chirish)."""

    pk_turi = "guruh"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"

    def get(self, request, pk):
        guruh = get_object_or_404(Guruh, pk=pk)
        try:
            oy = _oy(request.query_params.get("oy") or timezone.localdate().strftime("%Y-%m"))
        except ValueError as e:
            return _xato(str(e))
        return Response({
            str(m.sana): m.mavzu
            for m in DarsMavzusi.objects.filter(guruh=guruh, sana__range=(oy, mantiq.oy_oxiri(oy)))
        })

    def post(self, request, pk):
        if "guruhlar.davomat" not in ruxsatlar(request.user):
            return _xato("Ruxsat yo'q", kod=403)
        guruh = get_object_or_404(Guruh, pk=pk)
        try:
            sana = _sana(request.data.get("sana"), "sana")
        except ValueError as e:
            return _xato(str(e))
        mavzu = (request.data.get("mavzu") or "").strip()[:200]
        if not mavzu:
            DarsMavzusi.objects.filter(guruh=guruh, sana=sana).delete()
        else:
            DarsMavzusi.objects.update_or_create(guruh=guruh, sana=sana, defaults={"mavzu": mavzu, "kim": request.user})
        return Response({"sana": sana, "mavzu": mavzu})


class DavomatHammasiView(CrmView):
    """Bir kun uchun hammani "keldi" qilish (SoffCRM ustun boshidagi
    belgi). Allaqachon belgilangan kataklarga TEGILMAYDI — faqat
    bo'shlari to'ldiriladi, qo'shilishidan oldingi kunlar o'tkaziladi."""

    pk_turi = "guruh"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"

    def post(self, request, pk):
        if "guruhlar.davomat" not in ruxsatlar(request.user):
            return _xato("Ruxsat yo'q", kod=403)
        guruh = get_object_or_404(Guruh, pk=pk)
        try:
            sana = _sana(request.data.get("sana"), "sana")
        except ValueError as e:
            return _xato(str(e))
        if sana > timezone.localdate():
            return _xato("Kelajakdagi darsga davomat qo'yib bo'lmaydi")
        bor = set(Davomat.objects.filter(guruh=guruh, sana=sana).values_list("talaba_id", flat=True))
        yangi = []
        for a in GuruhAzoligi.objects.filter(guruh=guruh).select_related("moliya"):
            am = getattr(a, "moliya", None)
            if a.talaba_id in bor or (am and am.boshlanish_sana and sana < am.boshlanish_sana):
                continue
            yangi.append(Davomat(guruh=guruh, talaba_id=a.talaba_id, sana=sana,
                                 holat=Davomat.Holat.KELDI, belgilagan=request.user))
        Davomat.objects.bulk_create(yangi)
        return Response({"belgilandi": len(yangi)})


# ── Excel eksportlar ─────────────────────────────────────────────────


class DavomatEksportView(CrmView):
    pk_turi = "guruh"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"

    def get(self, request, pk):
        guruh = get_object_or_404(Guruh, pk=pk)
        try:
            oy = _oy(request.query_params.get("oy") or timezone.localdate().strftime("%Y-%m"))
        except ValueError as e:
            return _xato(str(e))
        yozuvlar = Davomat.objects.filter(guruh=guruh, sana__range=(oy, mantiq.oy_oxiri(oy))).select_related("crm_izoh")
        sanalar = sorted(guruh_dars_sanalari(guruh, oy) | {y.sana for y in yozuvlar})
        katak = {}
        for y in yozuvlar:
            izoh = getattr(y, "crm_izoh", None)
            katak[(y.talaba_id, y.sana)] = "S" if (izoh and izoh.sababli) else ("+" if y.holat == "keldi" else "-")
        qatorlar = []
        for a in sorted(GuruhAzoligi.objects.filter(guruh=guruh).select_related("talaba"),
                        key=lambda a: (_ism(a.talaba) or "").lower()):
            kunlar = [katak.get((a.talaba_id, s), "") for s in sanalar]
            qatorlar.append([_ism(a.talaba), *kunlar, kunlar.count("+"), kunlar.count("-"), kunlar.count("S")])
        return _excel(
            "Davomat", ["O'quvchi", *[s.strftime("%d.%m") for s in sanalar], "Keldi", "Kelmadi", "Sababli"],
            qatorlar, f"davomat_{guruh.name}_{oy:%Y-%m}",
        )


class GuruhlarEksportView(CrmView):
    bolim = "guruhlar"

    def get(self, request):
        from .views import _guruh_dict

        qs = (
            Guruh.objects.filter(faol=not request.query_params.get("arxiv")).filter(guruh_q(request.user))
            .select_related("daraja", "oqituvchi", "moliya", "moliya__filial", "daraja__crm_narxi")
            .prefetch_related("crm_jadval", "crm_oqituvchilar__oqituvchi")
            .annotate(_soni=Count("talabalar", distinct=True))
            .order_by("name")
        )
        if request.query_params.get("filial"):
            qs = qs.filter(moliya__filial_id=request.query_params["filial"])
        kun = ["Du", "Se", "Chor", "Pay", "Ju", "Sha", "Yak"]
        qatorlar = []
        for g in qs:
            d = _guruh_dict(g, talaba_soni=g._soni)
            qatorlar.append([
                g.id, g.name, d["daraja"]["nomi"] if d["daraja"] else "", d["oqituvchi"] or "",
                ", ".join(d["yordamchilar"]),
                ", ".join(f"{kun[j['hafta_kuni']]} {j['boshlanish_vaqti']}-{j['tugash_vaqti']}"
                          f"{' ' + j['xona'] if j['xona'] else ''}" for j in d["jadval"]),
                g._soni, d["filial"]["nomi"] if d["filial"] else "",
                d["narx"] or 0,
                d["boshlanish_sana"].strftime("%d.%m.%Y") if d["boshlanish_sana"] else "",
                d["tugash_sana"].strftime("%d.%m.%Y") if d["tugash_sana"] else "",
                "Faol" if g.faol else "Arxiv",
            ])
        return _excel(
            "Guruhlar", ["ID", "Guruh", "Kurs", "O'qituvchi", "Support ustoz", "Dars jadvali", "O'quvchilar",
                         "Filial", "Narx", "Ochilgan", "Yakunlanadi", "Status"],
            qatorlar, "guruhlar", pul_ustunlari=(9,),
        )


class TalabalarEksportView(CrmView):
    bolim = "talabalar"

    def get(self, request):
        from .views import TalabalarView

        if "talabalar.excel" not in ruxsatlar(request.user):
            return _xato("Excel eksportga ruxsat yo'q", kod=403)
        # Ro'yxat AYNAN ekrandagi filtrlar bilan — o'sha view qayta ishlatiladi.
        royxat = TalabalarView().get(request).data
        profillar = {
            p.user_id: p for p in TalabaProfil.objects.filter(user_id__in=[x["id"] for x in royxat])
        }
        userlar = User.objects.in_bulk([x["id"] for x in royxat])
        qatorlar = []
        for x in royxat:
            u = userlar.get(x["id"])
            p = profillar.get(x["id"])
            qatorlar.append([
                x["id"], x["ism"], x["telefon"] or "", u.ota_ona_telefon if u else "",
                u.username if u else "", ", ".join(g["guruh"] for g in x["guruhlar"]),
                ", ".join(g["holat"] for g in x["guruhlar"]), p.maktab if p else "",
                u.manba if u else "", u.izoh if u else "", "Ha" if x.get("qora_royxat") else "",
                float(x["balans"] or 0),
            ])
        return _excel(
            "O'quvchilar", ["ID", "Ism familiya", "Telefon", "Ota-ona telefoni", "Login", "Guruhlar", "Holat",
                            "Maktab", "Manba", "Izoh", "Qora ro'yxat", "Balans"],
            qatorlar, "oquvchilar", pul_ustunlari=(12,),
        )


class XodimlarEksportView(CrmView):
    bolim = "xodimlar"

    def get(self, request):
        from .boshqaruv import _xodim_dict

        oylik = "xodimlar.oylik" in ruxsatlar(request.user)
        qatorlar = []
        for u in _xodimlar_qs().filter(is_active=not request.query_params.get("arxiv")).order_by("first_name"):
            if not xodim_korinadimi(request.user, u):
                continue
            x = _xodim_dict(u, oylik)
            qatorlar.append([
                x["id"], x["ism"], x["username"], x["telefon"] or "", x["lavozim"], x["rol"] or "",
                x["filial"] or "", x["ishga_olingan_sana"].strftime("%d.%m.%Y") if x["ishga_olingan_sana"] else "",
                float(x["oylik"]) if x["oylik"] is not None else "", float(x["foiz_ulushi"]) if x["foiz_ulushi"] is not None else "",
            ])
        return _excel(
            "Xodimlar", ["ID", "Ism familiya", "Login", "Telefon", "Lavozim", "Rol", "Filial",
                         "Ishga olingan", "Oylik", "Foiz ulushi"],
            qatorlar, "xodimlar", pul_ustunlari=(9,),
        )


# ── O'quvchilarni Excel orqali qo'shish (guruhsiz) ───────────────────


class TalabalarImportView(CrmView):
    """SoffCRM "EXCEL ORQALI QO'SHISH": A=ism, B=telefon, C=ota-ona
    telefoni, D=tug'ilgan sana (ixtiyoriy). Telefoni mos talaba bo'lsa —
    o'tkazib yuboriladi (dublikat ochilmaydi)."""

    bolim = "talabalar"
    parser_classes = [MultiPartParser, FormParser]

    def post(self, request):
        import openpyxl

        if "talabalar.qoshish" not in ruxsatlar(request.user):
            return _xato("Ruxsat yo'q", kod=403)
        fayl = request.FILES.get("excel_fayl")
        if fayl is None:
            return _xato("excel_fayl majburiy")
        try:
            varaq = openpyxl.load_workbook(fayl, read_only=True, data_only=True).active
            qatorlar = [r for i, r in enumerate(varaq.iter_rows(values_only=True)) if i > 0 and r and any(r)]
        except Exception:
            return _xato("Excel fayl noto'g'ri formatda")
        natija, xatolar = [], []
        for n, r in enumerate(qatorlar, start=2):
            ism = str(r[0] or "").strip() if len(r) > 0 else ""
            telefon = _telefon(r[1]) if len(r) > 1 else ""
            if not ism:
                xatolar.append({"qator": n, "xato": "ism bo'sh"})
                continue
            if telefon and User.objects.filter(role=User.Role.STUDENT, telefon=telefon).exists():
                xatolar.append({"qator": n, "xato": f"{telefon} — bunday o'quvchi bor"})
                continue
            tugilgan = r[3] if len(r) > 3 else None
            if hasattr(tugilgan, "date"):
                tugilgan = tugilgan.date()
            try:
                with transaction.atomic():
                    user, parol, xato = talaba_yarat({
                        "ism": ism, "telefon": telefon,
                        "ota_ona_telefon": _telefon(r[2]) if len(r) > 2 else "",
                        **({"tugilgan_sana": tugilgan.isoformat()} if isinstance(tugilgan, date) else {}),
                    }, request.user)
                    if xato:
                        raise ValueError(xato)
                natija.append({"ism": _ism(user), "username": user.username, "parol": parol})
            except ValueError as e:
                xatolar.append({"qator": n, "xato": str(e)})
        return Response({"qoshildi": natija, "xatolar": xatolar}, status=201)


# ── Guruh: ommaviy faollashtirish, sobiq a'zolar ─────────────────────


class GuruhFaollashtirishView(CrmView):
    """"O'QUVCHILARNI FAOLLASHTIRISH" — guruhdagi sinovdagilarni faolga."""

    pk_turi = "guruh"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"

    def post(self, request, pk):
        if xato := _ruxsatsiz(request, "guruhlar.talaba_qoshish"):
            return xato
        guruh = get_object_or_404(Guruh, pk=pk)
        idlar = request.data.get("talaba_idlar")
        qs = AzolikMoliya.objects.filter(azolik__guruh=guruh, holat=AzolikMoliya.Holat.SINOV)
        if idlar:
            qs = qs.filter(azolik__talaba_id__in=idlar)
        soni = 0
        for am in qs:
            am.holat = AzolikMoliya.Holat.FAOL
            am.save(update_fields=["holat"])
            soni += 1
        return Response({"faollashtirildi": soni})


class GuruhSobiqlariView(CrmView):
    """"Arxivdagi o'quvchilarni ko'rish" — guruhdan chiqqanlar."""

    pk_turi = "guruh"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "guruhlar"

    def get(self, request, pk):
        guruh = get_object_or_404(Guruh, pk=pk)
        return Response([
            {
                "talaba_id": c.talaba_id, "talaba": c.talaba_ism, "boshlagan_sana": c.boshlagan_sana,
                "sana": c.sana, "sabab": c.sabab, "kim": _ism(c.kim) if c.kim_id else None,
                "balans": mantiq.balans(c.talaba, guruh=guruh) if c.talaba_id else None,
            }
            for c in GuruhdanChiqish.objects.filter(guruh=guruh).select_related("talaba", "kim")
        ])


# ── O'quvchi tarixi ──────────────────────────────────────────────────


class TalabaTarixiView(CrmView):
    """"O'QUVCHI TARIX" tabi: audit yozuvlari (profil, parol, guruhga
    qo'shish/chiqarish) + to'lovlar + guruhdan chiqishlar — bitta
    vaqt chizig'ida."""

    pk_turi = "talaba"  # filial cheklovi: `crm.filial.obyekt_tekshir`
    bolim = "talabalar"

    def get(self, request, pk):
        talaba = get_object_or_404(User, pk=pk)
        voqealar = []
        for f in FaoliyatYozuvi.objects.filter(
            obyekt_id=talaba.pk, obyekt_turi__in=["Talaba", "Foydalanuvchi"]
        ).select_related("foydalanuvchi")[:200]:
            voqealar.append({
                "vaqt": f.vaqt, "turi": f.get_harakat_display(),
                "matn": ", ".join(f.ozgarishlar.keys()) if isinstance(f.ozgarishlar, dict) else "",
                "kim": _ism(f.foydalanuvchi) if f.foydalanuvchi_id else None,
            })
        for t in Tolov.objects.filter(talaba=talaba).select_related("kim_kiritdi")[:200]:
            voqealar.append({
                "vaqt": t.created_at, "turi": t.get_turi_display(),
                "matn": f"{t.guruh_nomi}: {t.summa:,.0f} so'm".replace(",", " ") + (f" — {t.izoh}" if t.izoh else ""),
                "kim": _ism(t.kim_kiritdi) if t.kim_kiritdi_id else None,
            })
        for c in GuruhdanChiqish.objects.filter(talaba=talaba).select_related("kim"):
            voqealar.append({
                "vaqt": c.created_at, "turi": "Guruhdan chiqdi",
                "matn": f"{c.guruh_nomi} ({c.sana:%d.%m.%Y})" + (f" — {c.sabab}" if c.sabab else ""),
                "kim": _ism(c.kim) if c.kim_id else None,
            })
        for a in GuruhAzoligi.objects.filter(talaba=talaba).select_related("guruh"):
            voqealar.append({"vaqt": a.created_at, "turi": "Guruhga qo'shildi", "matn": a.guruh.name, "kim": None})
        voqealar.sort(key=lambda v: v["vaqt"], reverse=True)
        return Response(voqealar)


# ── Xodimlar davomati ────────────────────────────────────────────────


class XodimDavomatView(CrmView):
    """Oy bo'yicha xodimlar davomati jadvali (ham belgilash, ham hisobot)."""

    bolim = "xodimlar"

    def get(self, request):
        try:
            oy = _oy(request.query_params.get("oy") or timezone.localdate().strftime("%Y-%m"))
        except ValueError as e:
            return _xato(str(e))
        oxiri = mantiq.oy_oxiri(oy)
        kunlar = [date(oy.year, oy.month, k) for k in range(1, oxiri.day + 1)]
        yozuvlar = {
            (y.xodim_id, y.sana): {"holat": y.holat, "izoh": y.izoh}
            for y in XodimDavomat.objects.filter(sana__range=(oy, oxiri))
        }
        xodimlar = []
        for u in _xodimlar_qs().filter(is_active=True).order_by("first_name", "username"):
            if not xodim_korinadimi(request.user, u):
                continue
            qator = [yozuvlar.get((u.id, k)) for k in kunlar]
            sanoq = Counter(q["holat"] for q in qator if q)
            xodimlar.append({"id": u.id, "ism": _ism(u), "kunlar": qator, "jami": dict(sanoq)})
        return Response({"oy": oy, "kunlar": kunlar, "xodimlar": xodimlar})

    def post(self, request):
        if xato := _ruxsatsiz(request, "xodimlar.tahrirlash"):
            return xato
        xodim = get_object_or_404(_xodimlar_qs(), pk=request.data.get("xodim_id"))
        if not xodim_korinadimi(request.user, xodim):
            raise Http404
        try:
            sana = _sana(request.data.get("sana"), "sana")
        except ValueError as e:
            return _xato(str(e))
        if sana > timezone.localdate():
            return _xato("Kelajak kuniga davomat qo'yib bo'lmaydi")
        holat = request.data.get("holat")
        if not holat:
            XodimDavomat.objects.filter(xodim=xodim, sana=sana).delete()
            return Response({"holat": None})
        if holat not in dict(XodimDavomat.Holat.choices):
            return _xato("Holat noto'g'ri")
        XodimDavomat.objects.update_or_create(
            xodim=xodim, sana=sana,
            defaults={"holat": holat, "izoh": (request.data.get("izoh") or "").strip()[:300], "kim": request.user},
        )
        return Response({"holat": holat})


# ── Harakatlar tarixi ────────────────────────────────────────────────


class HarakatlarTarixiView(CrmView):
    """Sozlamalar -> "Harakatlar tarixi": CRM'dagi o'zgarishlar audit
    jurnali (kim, qachon, nima). Faqat o'qish."""

    bolim = "sozlamalar"

    def get(self, request):
        qs = FaoliyatYozuvi.objects.filter(
            Q(obyekt_turi__startswith="CRM") | Q(obyekt_turi__in=["Guruh", "Talaba", "Foydalanuvchi"])
        ).select_related("foydalanuvchi").order_by("-vaqt")
        p = request.query_params
        if p.get("turi"):
            qs = qs.filter(obyekt_turi=p["turi"])
        if p.get("kim"):
            qs = qs.filter(foydalanuvchi_id=p["kim"])
        if p.get("q"):
            qs = qs.filter(obyekt_nomi__icontains=p["q"].strip())
        try:
            dan, gacha = _davr(request) if (p.get("dan") or p.get("gacha")) else (None, None)
        except ValueError as e:
            return _xato(str(e))
        if dan:
            qs = qs.filter(vaqt__date__range=(dan, gacha))
        return Response([
            {
                "id": f.id, "vaqt": f.vaqt, "harakat": f.get_harakat_display(),
                "turi": f.obyekt_turi, "nomi": f.obyekt_nomi, "ozgarishlar": f.ozgarishlar,
                "kim": _ism(f.foydalanuvchi) if f.foydalanuvchi_id else None,
            }
            for f in qs[:300]
        ])


# ── Hisobotlar ───────────────────────────────────────────────────────


class KetganlarHisobotiView(CrmView):
    """"Ketgan o'quvchilar hisoboti" — davr ichida guruhdan chiqqanlar."""

    bolim = "hisobotlar"

    def get(self, request):
        try:
            dan, gacha = _davr(request)
        except ValueError as e:
            return _xato(str(e))
        qs = GuruhdanChiqish.objects.filter(sana__range=(dan, gacha)).filter(
            filial_q(request.user, "filial")).select_related("kim", "filial")
        if request.query_params.get("filial"):
            qs = qs.filter(filial_id=request.query_params["filial"])
        return Response({
            "dan": dan, "gacha": gacha,
            "royxat": [
                {
                    "talaba_id": c.talaba_id, "talaba": c.talaba_ism, "guruh": c.guruh_nomi,
                    "filial": c.filial.nomi if c.filial_id else None, "boshlagan_sana": c.boshlagan_sana,
                    "sana": c.sana, "sabab": c.sabab, "kim": _ism(c.kim) if c.kim_id else None,
                }
                for c in qs
            ],
        })


class LidlarHisobotiView(CrmView):
    """"Lidlar hisoboti" — davrda kelgan lidlar manba va holat bo'yicha,
    nechtasi o'quvchiga aylangani (konversiya)."""

    bolim = "hisobotlar"

    def get(self, request):
        try:
            dan, gacha = _davr(request)
        except ValueError as e:
            return _xato(str(e))
        qs = Lid.objects.filter(created_at__date__range=(dan, gacha)).filter(filial_q(request.user, "filial"))
        if request.query_params.get("filial"):
            qs = qs.filter(Q(filial_id=request.query_params["filial"]) | Q(filial__isnull=True))
        manbalar = []
        for x in qs.values("manba").annotate(
            jami=Count("id"), qoshildi=Count("id", filter=Q(talaba__isnull=False)),
        ).order_by("-jami"):
            manbalar.append({
                "manba": x["manba"] or "—", "jami": x["jami"], "qoshildi": x["qoshildi"],
                "konversiya": round(x["qoshildi"] / x["jami"] * 100, 1) if x["jami"] else 0,
            })
        holatlar = dict(Lid.Holat.choices)
        return Response({
            "dan": dan, "gacha": gacha, "jami": qs.count(),
            "qoshildi": qs.filter(talaba__isnull=False).count(),
            "manbalar": manbalar,
            "holatlar": [
                {"holat": x["holat"], "nomi": holatlar.get(x["holat"], x["holat"]), "soni": x["soni"]}
                for x in qs.values("holat").annotate(soni=Count("id")).order_by("-soni")
            ],
        })


class TolovlarHisobotiView(CrmView):
    """"To'lovlar hisoboti" — davrdagi to'lovlar to'lov usuli va kun
    bo'yicha, qaytarishlar alohida."""

    bolim = "hisobotlar"

    def get(self, request):
        try:
            dan, gacha = _davr(request)
        except ValueError as e:
            return _xato(str(e))
        qs = Tolov.objects.filter(sana__range=(dan, gacha)).filter(guruh_q(request.user, "guruh__"))
        if request.query_params.get("filial"):
            qs = qs.filter(guruh__moliya__filial_id=request.query_params["filial"])
        usullar = dict(Tolov.Usul.choices)
        tolov = qs.filter(turi=Tolov.Turi.TOLOV)
        return Response({
            "dan": dan, "gacha": gacha,
            "jami": tolov.aggregate(s=Sum("summa"))["s"] or NOL,
            "qaytarish": qs.filter(turi=Tolov.Turi.QAYTARISH).aggregate(s=Sum("summa"))["s"] or NOL,
            "chegirma_bonus": qs.filter(turi__in=[Tolov.Turi.CHEGIRMA, Tolov.Turi.BONUS]).aggregate(s=Sum("summa"))["s"] or NOL,
            "usullar": [
                {"usul": x["usul"], "nomi": usullar.get(x["usul"], x["usul"] or "—"), "summa": x["s"], "soni": x["n"]}
                for x in tolov.values("usul").annotate(s=Sum("summa"), n=Count("id")).order_by("-s")
            ],
            "kunlar": [
                {"sana": x["sana"], "summa": x["s"], "soni": x["n"]}
                for x in tolov.values("sana").annotate(s=Sum("summa"), n=Count("id")).order_by("sana")
            ],
        })


class BitiruvchilarHisobotiView(CrmView):
    """"Bitiruvchilar hisoboti" — davrda tugagan guruhlar va ularning
    o'quvchilari."""

    bolim = "hisobotlar"

    def get(self, request):
        try:
            dan, gacha = _davr(request)
        except ValueError as e:
            return _xato(str(e))
        guruhlar = Guruh.objects.filter(moliya__tugash_sana__range=(dan, gacha)).filter(
            guruh_q(request.user)).select_related(
            "moliya", "oqituvchi", "daraja"
        )
        natija = []
        for g in guruhlar:
            talabalar = [_ism(a.talaba) for a in GuruhAzoligi.objects.filter(guruh=g).select_related("talaba")]
            natija.append({
                "guruh_id": g.id, "guruh": g.name, "kurs": g.daraja.nomi if g.daraja_id else None,
                "oqituvchi": _ism(g.oqituvchi) if g.oqituvchi_id else None,
                "tugash_sana": g.moliya.tugash_sana, "talabalar": talabalar,
            })
        return Response({"dan": dan, "gacha": gacha, "guruhlar": natija})


class HisobotlarDavrView(CrmView):
    """Yuqoridagilarni Excel'ga — `?turi=ketganlar|lidlar|tolovlar`."""

    bolim = "hisobotlar"

    def get(self, request):
        turi = request.query_params.get("turi")
        if turi == "ketganlar":
            d = KetganlarHisobotiView().get(request).data
            return _excel("Ketganlar", ["O'quvchi", "Guruh", "Filial", "Boshlagan", "Chiqqan", "Sabab", "Kim"], [
                [x["talaba"], x["guruh"], x["filial"] or "",
                 x["boshlagan_sana"].strftime("%d.%m.%Y") if x["boshlagan_sana"] else "",
                 x["sana"].strftime("%d.%m.%Y"), x["sabab"], x["kim"] or ""]
                for x in d.get("royxat", [])
            ], "ketgan_oquvchilar")
        if turi == "lidlar":
            d = LidlarHisobotiView().get(request).data
            return _excel("Lidlar", ["Manba", "Jami", "O'quvchi bo'ldi", "Konversiya %"], [
                [x["manba"], x["jami"], x["qoshildi"], x["konversiya"]] for x in d.get("manbalar", [])
            ], "lidlar_hisoboti")
        if turi == "tolovlar":
            d = TolovlarHisobotiView().get(request).data
            return _excel("To'lovlar", ["Sana", "Summa", "Soni"], [
                [x["sana"].strftime("%d.%m.%Y"), float(x["summa"]), x["soni"]] for x in d.get("kunlar", [])
            ], "tolovlar_hisoboti", pul_ustunlari=(2,))
        return _xato("turi: ketganlar | lidlar | tolovlar")

