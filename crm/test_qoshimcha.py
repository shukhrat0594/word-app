"""Video-TZ qayta tekshiruvi (5 soniyalik kadrlar) qo'shimchalari testlari."""

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook, load_workbook

from academics.models import Davomat
from accounts.models import User
from crm import mantiq
from crm.models import AzolikMoliya, GuruhdanChiqish, Hisob, Lid, LidBolim, Tolov, XodimDavomat
from crm.test_api import ApiAsos
from crm.tests import NARX, SENTABR, bugun_qilib


class LidDoskaTest(ApiAsos):
    def test_doska_yangi_ustun_bilan_ochiladi(self):
        m = self.mijoz(self.admin)
        doska = m.post("/api/crm/lid-doskalar/", {"nomi": "LEADS uzb"}, format="json").data
        ustunlar = m.get(f"/api/crm/lid-bolimlar/?doska={doska['id']}").data
        self.assertEqual([u["nomi"] for u in ustunlar], ["NEW LEADS"])
        m.post("/api/crm/lidlar/", {"ism": "A", "telefon": "901000001", "bolim_id": ustunlar[0]["id"],
                                    "harorat": "issiq"}, format="json")
        self.assertEqual(len(m.get(f"/api/crm/lidlar/?doska={doska['id']}").data), 1)
        self.assertEqual(len(m.get("/api/crm/lidlar/?doska=0").data), 0)
        # Doska o'chsa — lid yo'qolmaydi, "Umumiy"ga tushadi
        m.delete(f"/api/crm/lid-doskalar/{doska['id']}/")
        self.assertEqual(Lid.objects.get().bolim_id, None)
        self.assertEqual(Lid.objects.get().harorat, "issiq")
        self.assertFalse(LidBolim.objects.exists())


class GuruhQoshimchaTest(ApiAsos):
    def test_hammasi_keldi_faqat_boshlarini(self):
        self.azolik_qosh(boshlanish=SENTABR)
        ikkinchi = User.objects.create_user(username="t2", password="x", role=User.Role.STUDENT)
        self.azolik_qosh(talaba=ikkinchi, boshlanish=SENTABR)
        Davomat.objects.create(guruh=self.guruh, talaba=ikkinchi, sana=date(2026, 9, 3), holat="kelmadi")
        r = self.mijoz(self.admin).post(f"/api/crm/guruhlar/{self.guruh.id}/davomat/hammasi/",
                                        {"sana": "2026-09-03"}, format="json")
        self.assertEqual(r.data["belgilandi"], 1)
        self.assertEqual(Davomat.objects.get(talaba=ikkinchi).holat, "kelmadi")  # tegilmadi

    def test_mavzu_va_davomat_excel(self):
        m = self.mijoz(self.admin)
        self.azolik_qosh(boshlanish=SENTABR)
        m.post(f"/api/crm/guruhlar/{self.guruh.id}/mavzular/", {"sana": "2026-09-03", "mavzu": "Present simple"},
               format="json")
        self.assertEqual(m.get(f"/api/crm/guruhlar/{self.guruh.id}/mavzular/?oy=2026-09").data,
                         {"2026-09-03": "Present simple"})
        r = m.get(f"/api/crm/guruhlar/{self.guruh.id}/davomat/eksport/?oy=2026-09")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(load_workbook(BytesIO(r.content)).active.cell(2, 1).value, "talaba1")

    def test_chiqish_yoziladi_va_hisobotda(self):
        self.azolik_qosh(boshlanish=SENTABR)
        with bugun_qilib(date(2026, 9, 12)):
            self.mijoz(self.admin).delete(
                f"/api/crm/guruhlar/{self.guruh.id}/talabalar/?talaba={self.talaba.id}&sana=2026-09-12&sabab=Ko'chdi"
            )
        c = GuruhdanChiqish.objects.get()
        self.assertEqual((c.sabab, c.boshlagan_sana), ("Ko'chdi", SENTABR))
        r = self.mijoz(self.admin).get("/api/crm/hisobot/ketganlar/?dan=2026-09-01&gacha=2026-09-30")
        self.assertEqual([x["talaba"] for x in r.data["royxat"]], ["talaba1"])
        sobiq = self.mijoz(self.admin).get(f"/api/crm/guruhlar/{self.guruh.id}/sobiqlar/").data
        self.assertEqual(len(sobiq), 1)

    def test_faollashtirish(self):
        am = self.azolik_qosh(boshlanish=SENTABR, holat="sinov")
        self.mijoz(self.admin).post(f"/api/crm/guruhlar/{self.guruh.id}/faollashtirish/", {}, format="json")
        am.refresh_from_db()
        self.assertEqual(am.holat, "faol")

    def test_chegirma_foizda_va_doimiy(self):
        am = self.azolik_qosh(boshlanish=SENTABR)
        # Sentabr "joriy oy" bo'lsin: o'tgan oyga chegirmani faqat owner beradi.
        with bugun_qilib(date(2026, 9, 15)):
            r = self.mijoz(self.admin).post(f"/api/crm/guruhlar/{self.guruh.id}/chegirmalar/", {
                "azolik_moliya_id": am.id, "foiz": "50", "boshlanish_oy": "2026-09", "oylar_soni": 0,
            }, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertTrue(r.data["doimiy"])
        self.assertEqual(mantiq.amaldagi_narx(am, date(2027, 5, 1)), NARX / 2)


class TalabaQoshimchaTest(ApiAsos):
    def test_excel_import_dublikatsiz(self):
        User.objects.filter(pk=self.talaba.pk).update(telefon="+998901112233")
        kitob = Workbook()
        ws = kitob.active
        ws.append(["Ism", "Telefon", "Ota-ona"])
        ws.append(["Yangi Bola", "901234599", "901234598"])
        ws.append(["Takror", "+998901112233", ""])
        fayl = BytesIO()
        kitob.save(fayl)
        fayl.seek(0)
        fayl.name = "t.xlsx"
        r = self.mijoz(self.admin).post("/api/crm/talabalar/import/", {"excel_fayl": fayl}, format="multipart")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual([x["ism"] for x in r.data["qoshildi"]], ["Yangi Bola"])
        self.assertEqual(len(r.data["xatolar"]), 1)

    def test_eksport_va_tarix(self):
        self.azolik_qosh(boshlanish=SENTABR)
        m = self.mijoz(self.admin)
        r = m.get("/api/crm/talabalar/eksport/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("talaba1", [c.value for c in load_workbook(BytesIO(r.content)).active["B"]])
        tarix = m.get(f"/api/crm/talaba/{self.talaba.id}/tarix/").data
        self.assertTrue(any(v["turi"] == "Guruhga qo'shildi" for v in tarix))

    def test_filtr_guruh_boyicha(self):
        self.azolik_qosh(boshlanish=SENTABR)
        r = self.mijoz(self.admin).get(f"/api/crm/talabalar/?guruh={self.guruh.id}")
        self.assertEqual([x["ism"] for x in r.data], ["talaba1"])

    def test_hisob_izohi(self):
        self.azolik_qosh(boshlanish=SENTABR)
        with bugun_qilib(date(2026, 9, 5)):
            mantiq.hisoblarni_generatsiya_qil()
        h = Hisob.objects.get()
        r = self.mijoz(self.owner).patch(f"/api/crm/hisoblar/{h.id}/", {"summa": "180000", "izoh": "4 ta dars"},
                                         format="json")
        self.assertEqual(r.status_code, 200, r.data)
        h.refresh_from_db()
        self.assertEqual((h.summa, h.izoh), (Decimal("180000"), "4 ta dars"))


class XodimVaHisobotTest(ApiAsos):
    def test_xodim_davomati(self):
        m = self.mijoz(self.admin)
        r = m.post("/api/crm/xodimlar/davomat/", {"xodim_id": self.oqituvchi.id, "sana": "2026-09-03",
                                                  "holat": "kechikdi"}, format="json")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertEqual(XodimDavomat.objects.get().holat, "kechikdi")
        jadval = m.get("/api/crm/xodimlar/davomat/?oy=2026-09").data
        qator = next(x for x in jadval["xodimlar"] if x["id"] == self.oqituvchi.id)
        self.assertEqual(qator["jami"], {"kechikdi": 1})

    def test_lidlar_va_tolovlar_hisoboti(self):
        m = self.mijoz(self.admin)
        m.post("/api/crm/lidlar/", {"ism": "A", "telefon": "901000011", "manba": "Instagram"}, format="json")
        r = m.get("/api/crm/hisobot/lidlar/")
        self.assertEqual(r.data["manbalar"][0]["manba"], "Instagram")
        self.assertEqual(m.get("/api/crm/hisobot/tolovlar/").status_code, 200)
        self.assertEqual(m.get("/api/crm/hisobot/eksport/?turi=lidlar").status_code, 200)
        self.assertEqual(m.get("/api/crm/harakatlar/").status_code, 200)
        self.assertEqual(m.get("/api/crm/guruhlar/eksport/").status_code, 200)
        self.assertEqual(m.get("/api/crm/xodimlar/eksport/").status_code, 200)

    def test_tolovlar_hisoboti_birinchi_va_qayta(self):
        """Video-TZ 2026-10-12: sentyabrda birinchi to'lagan o'quvchining
        oktyabrdagi to'lovi — qayta; oktyabrda ilk to'lagan — birinchi."""
        from crm.models import Tolov
        yangi = User.objects.create_user(username="sinov_yangi", password="x", role=self.talaba.role)
        for talaba, sana in ((self.talaba, date(2026, 9, 5)), (self.talaba, date(2026, 10, 3)),
                             (yangi, date(2026, 10, 3)), (yangi, date(2026, 10, 3))):
            Tolov.objects.create(talaba=talaba, talaba_ism="SINOV", guruh=self.guruh, guruh_nomi="y",
                                 sana=sana, summa=NARX, turi=Tolov.Turi.TOLOV)
        m = self.mijoz(self.owner)
        d = m.get("/api/crm/hisobot/tolovlar/?dan=2026-10-01&gacha=2026-10-31").data
        self.assertEqual((d["birinchi_soni"], d["qayta_soni"]), (1, 2))
        self.assertEqual((d["birinchi"], d["qayta"]), (NARX, 2 * NARX))
        kun = d["kunlar"][0]
        self.assertEqual((kun["birinchi_soni"], kun["qayta_soni"], kun["soni"]), (1, 2, 3))
        # Har bir to'lov alohida: o'quvchi, usul, summa, birinchi/qayta va balans.
        self.assertEqual([(x["talaba_id"], x["qayta"]) for x in d["royxat"]],
                         [(self.talaba.id, True), (yangi.id, False), (yangi.id, True)])
        self.assertEqual(d["royxat"][1]["summa"], NARX)
        self.assertEqual(d["royxat"][1]["balans"], 2 * NARX)
        r = m.get("/api/crm/hisobot/eksport/?turi=tolovlar&dan=2026-10-01&gacha=2026-10-31")
        self.assertEqual(r.status_code, 200)

    def test_lidlar_voronkasi(self):
        Lid.objects.create(ism="A", telefon="1", yigilayotgan_guruh=self.guruh)
        Lid.objects.create(ism="B", telefon="2", holat=Lid.Holat.SINOV)
        Lid.objects.create(ism="C", telefon="3", holat=Lid.Holat.KELMADI)
        Lid.objects.create(ism="D", telefon="4", holat=Lid.Holat.QOSHILDI, talaba=self.talaba)
        self.azolik_qosh()  # faol — qoldi
        for _ in range(2):
            Tolov.objects.create(talaba=self.talaba, talaba_ism="t", guruh=self.guruh, guruh_nomi="g",
                                 sana=date(2026, 9, 3), summa=NARX, turi=Tolov.Turi.TOLOV)
        # Guruhga sinov a'zolik bilan qo'shilgan — hali qaror yo'q
        sinovda = User.objects.create_user(username="t2", password="x", role=User.Role.STUDENT)
        Lid.objects.create(ism="E", telefon="5", holat=Lid.Holat.QOSHILDI, talaba=sinovda)
        self.azolik_qosh(talaba=sinovda, holat=AzolikMoliya.Holat.SINOV)
        # Sinovdan keyin lidlarga qaytarilgan — `talaba` bog'liq qoladi, lekin qolmadi
        qaytgan = User.objects.create_user(username="t3", password="x", role=User.Role.STUDENT)
        Lid.objects.create(ism="F", telefon="6", holat=Lid.Holat.YANGI, talaba=qaytgan)
        r = self.mijoz(self.admin).get("/api/crm/hisobot/lidlar/")
        self.assertEqual({x["kalit"]: x["soni"] for x in r.data["voronka"]}, {
            "trafik": 6, "naborga": 1, "sinov_chaqirildi": 5, "sinov_keldi": 4, "sinov_kelmadi": 1,
            "yoqdi": 1, "yoqmadi": 1, "tolov_1": 1, "tolov_2": 1,
        })

    def test_eslatma_tahrirlash_faqat_ozinikini(self):
        m = self.mijoz(self.admin)
        e = m.post("/api/crm/eslatmalar/", {"talaba_id": self.talaba.id, "matn": "a"}, format="json").data
        self.assertEqual(m.patch(f"/api/crm/eslatmalar/{e['id']}/", {"matn": "b",
                                 "eslatish_vaqti": "2026-09-28T14:00"}, format="json").status_code, 200)
        self.assertEqual(self.mijoz(self.owner).patch(f"/api/crm/eslatmalar/{e['id']}/", {"matn": "c"},
                                                      format="json").status_code, 403)




class IkkiSoniyaliTekshiruvTest(ApiAsos):
    """2 soniyalik kadrlardan chiqqan qo'shimchalar."""

    def test_yigilayotgan_guruh_va_qabul_soni(self):
        m = self.mijoz(self.admin)
        lid = m.post("/api/crm/lidlar/", {"ism": "Navbat", "telefon": "901000021"}, format="json").data
        r = m.patch(f"/api/crm/lidlar/{lid['id']}/", {"yigilayotgan_guruh_id": self.guruh.id}, format="json")
        self.assertEqual(r.data["yigilayotgan_guruh"], self.guruh.name)
        self.assertEqual(m.get("/api/crm/korsatkichlar/").data["yangi_guruhga_qabul"], 1)
        m.post("/api/crm/lidlar/guruhga/", {"lid_idlar": [lid["id"]], "guruh_id": self.guruh.id}, format="json")
        self.assertIsNone(Lid.objects.get(pk=lid["id"]).yigilayotgan_guruh_id)

    def test_bosh_sahifa_kartochkalari_filtri(self):
        """"Bugungi lidlar" / "Yangi guruhga qabul" bosilganda ro'yxat
        kartochkadagi son bilan bir xil (video-TZ 2026-10-12)."""
        m = self.mijoz(self.admin)
        bugungi = m.post("/api/crm/lidlar/", {"ism": "Bugun", "telefon": "901000031"}, format="json").data
        eski = m.post("/api/crm/lidlar/", {"ism": "Eski", "telefon": "901000032"}, format="json").data
        Lid.objects.filter(pk=eski["id"]).update(created_at=timezone.now() - timedelta(days=3))
        m.patch(f"/api/crm/lidlar/{eski['id']}/", {"yigilayotgan_guruh_id": self.guruh.id}, format="json")
        k = m.get("/api/crm/korsatkichlar/").data
        bugun = [x["id"] for x in m.get("/api/crm/lidlar/?bugun=1").data]
        self.assertEqual(bugun, [bugungi["id"]])
        self.assertEqual(len(bugun), k["yangi_lidlar_bugun"])
        navbat = [x["id"] for x in m.get("/api/crm/lidlar/?yigilayotgan=1").data]
        self.assertEqual(navbat, [eski["id"]])
        self.assertEqual(len(navbat), k["yangi_guruhga_qabul"])

    def test_taqvimda_davomat_va_sanoq(self):
        self.azolik_qosh(boshlanish=SENTABR)
        Davomat.objects.create(guruh=self.guruh, talaba=self.talaba, sana=date(2026, 9, 3), holat="keldi")
        g = self.mijoz(self.admin).get(f"/api/crm/talaba/{self.talaba.id}/?oy=2026-09").data["guruhlar"][0]
        kun = next(k for k in g["darslar_taqvimi"] if k["sana"] == date(2026, 9, 3))
        self.assertEqual(kun["davomat"], "keldi")
        self.assertEqual(g["taqvim_sanogi"]["keldi"], 1)

    def test_royxatda_keyingi_tolov_va_setka(self):
        self.azolik_qosh(boshlanish=SENTABR)
        with bugun_qilib(date(2026, 9, 5)):
            mantiq.hisoblarni_generatsiya_qil()
        x = next(t for t in self.mijoz(self.admin).get("/api/crm/talabalar/").data if t["id"] == self.talaba.id)
        self.assertEqual(x["keyingi_tolov"], SENTABR)
        self.assertIn("baho", x)
        dars = self.mijoz(self.admin).get("/api/crm/jadval/").data["darslar"][0]
        self.assertEqual((dars["jami"], dars["faol"], dars["kurs"]), (1, 1, "IELTS"))
